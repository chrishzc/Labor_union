"""Typed client-registry order-accounting XLSX export."""

from __future__ import annotations

from collections.abc import Sequence
from dataclasses import dataclass
from datetime import date
from decimal import Decimal
from io import BytesIO
from typing import Literal, Protocol

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from domains.case_import.beclass_correction import VALID_MULTI_BIRTH_COUNTS
from subsystems.client_profile.registry_query import (
    ClientRegistrySortBy,
    ClientRegistrySortOrder,
)


XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

HEADERS = (
    "項次", "案件編號", "姓名", "身分資格", "服務時間",
    "預產期／預計服務開始月份", "預計服務日期", "雙胞胎",
    "目前訂單狀態", "匯入虛擬帳號", "內建虛擬帳號", "服務天數", "每日服務時數", "總服務時數",
    "是否需要下廚", "樓層費", "每小時服務單價", "客戶應付總額", "訂金金額",
    "第一期金額", "第二期金額", "已收總額", "客戶未收餘額", "補助返還金額",
    "預計服務開始日", "預計服務結束日", "實際服務開始日", "實際服務結束日",
    "訂金到期日", "第一期到期日", "第二期到期日", "訂金入帳日",
    "補助返還到期日", "補助返還狀態", "服務人員付款到期日",
    "補助申請年度", "補助申請月份",
)


@dataclass(frozen=True, slots=True)
class OrderAccountingExportQuery:
    query: str | None = None
    multi_birth_count: Literal["單胞胎", "雙胞胎"] | None = None
    order_status: str | None = None
    requires_cooking: bool | None = None
    sort_by: ClientRegistrySortBy | None = None
    sort_order: ClientRegistrySortOrder | None = None


@dataclass(frozen=True, slots=True)
class OrderAccountingExportRow:
    seq_num: int | None
    case_no: str
    name: str | None
    identity_status: str | None
    service_time: str | None
    due_month: str | None
    hcm_service_start_date: str | None
    is_twins: bool
    order_status: str | None
    imported_virtual_accounts: tuple[str, ...]
    built_in_virtual_account: str | None
    service_days: int | None
    service_hours_per_day: int | float | Decimal | None
    service_hours: int | float | Decimal | None
    requires_cooking: bool | None
    floor_fee_ntd: int | float | Decimal | None
    service_unit_price_ntd: int | None
    customer_payable_total_ntd: int | None
    deposit_amount_ntd: int | None
    first_payment_amount_ntd: int | None
    second_payment_amount_ntd: int | None
    received_total_ntd: int | None
    customer_balance_ntd: int | None
    subsidy_return_amount_ntd: int | None
    planned_start_date: date | None
    planned_end_date: date | None
    actual_start_date: date | None
    actual_end_date: date | None
    deposit_due_date: date | None
    first_payment_due_date: date | None
    second_payment_due_date: date | None
    deposit_settled_on: date | None
    subsidy_return_due_date: date | None
    subsidy_return_status: str | None
    staff_payment_due_date: date | None
    claim_application_year: int | None
    claim_application_month: int | None


class OrderAccountingExportRepository(Protocol):
    def query_order_accounting_rows(
        self, selection: OrderAccountingExportQuery
    ) -> tuple[OrderAccountingExportRow, ...]: ...


class ClientRegistryOrderAccountingExportApplication:
    def __init__(self, repository: OrderAccountingExportRepository) -> None:
        self._repository = repository

    def export(self, selection: OrderAccountingExportQuery) -> bytes:
        normalized = _normalize_query(selection)
        return build_order_accounting_workbook(
            self._repository.query_order_accounting_rows(normalized)
        )


def build_order_accounting_workbook(
    rows: Sequence[OrderAccountingExportRow],
) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "訂單帳務"
    worksheet.append(HEADERS)
    for row in rows:
        worksheet.append((
            row.seq_num,
            row.case_no,
            row.name,
            row.identity_status,
            row.service_time,
            row.due_month,
            row.hcm_service_start_date,
            "是" if row.is_twins else "否",
            row.order_status,
            "\n".join(row.imported_virtual_accounts) or None,
            row.built_in_virtual_account,
            row.service_days,
            row.service_hours_per_day,
            row.service_hours,
            _yes_no(row.requires_cooking),
            row.floor_fee_ntd,
            row.service_unit_price_ntd,
            row.customer_payable_total_ntd,
            row.deposit_amount_ntd,
            row.first_payment_amount_ntd,
            row.second_payment_amount_ntd,
            row.received_total_ntd,
            row.customer_balance_ntd,
            row.subsidy_return_amount_ntd,
            row.planned_start_date,
            row.planned_end_date,
            row.actual_start_date,
            row.actual_end_date,
            row.deposit_due_date,
            row.first_payment_due_date,
            row.second_payment_due_date,
            row.deposit_settled_on,
            row.subsidy_return_due_date,
            row.subsidy_return_status,
            row.staff_payment_due_date,
            row.claim_application_year,
            row.claim_application_month,
        ))

    header_fill = PatternFill(fill_type="solid", fgColor="FCE7D6")
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="7C2D12")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.sheet_view.showGridLines = False
    worksheet.row_dimensions[1].height = 32

    text_columns = (2, 10, 11)
    date_columns = tuple(range(25, 34)) + (35,)
    money_columns = tuple(range(16, 25))
    decimal_columns = (13, 14)
    for row_number in range(2, worksheet.max_row + 1):
        for column in text_columns:
            worksheet.cell(row=row_number, column=column).number_format = "@"
        for column in date_columns:
            worksheet.cell(row=row_number, column=column).number_format = "yyyy-mm-dd"
        for column in money_columns:
            worksheet.cell(row=row_number, column=column).number_format = "#,##0"
        for column in decimal_columns:
            worksheet.cell(row=row_number, column=column).number_format = "#,##0.0"
        worksheet.cell(row=row_number, column=10).alignment = Alignment(wrap_text=True)

    widths = (
        9, 16, 14, 14, 18, 24, 18, 10, 18, 20, 20, 11, 14, 14, 14,
        12, 16, 16, 14, 14, 14, 14, 16, 16,
    ) + (16,) * 13
    for column, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(column)].width = width

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _normalize_query(selection: OrderAccountingExportQuery) -> OrderAccountingExportQuery:
    query = _optional_text(selection.query, 100)
    multi_birth_count = _optional_text(selection.multi_birth_count, 20)
    if multi_birth_count is not None and multi_birth_count not in VALID_MULTI_BIRTH_COUNTS:
        raise ValueError("client_registry_multi_birth_count_invalid")
    order_status = _optional_text(selection.order_status, 50)
    if selection.requires_cooking is not None and not isinstance(selection.requires_cooking, bool):
        raise ValueError("client_registry_requires_cooking_invalid")
    if selection.sort_by is not None and selection.sort_by not in {
        "case_no", "customer_name", "service_days", "expected_start_date",
    }:
        raise ValueError("client_registry_sort_by_invalid")
    if selection.sort_order is not None and selection.sort_order not in {"asc", "desc"}:
        raise ValueError("client_registry_sort_order_invalid")
    if selection.sort_by is None and selection.sort_order is not None:
        raise ValueError("client_registry_sort_order_without_field")
    return OrderAccountingExportQuery(
        query=query,
        multi_birth_count=multi_birth_count,  # type: ignore[arg-type]
        order_status=order_status,
        requires_cooking=selection.requires_cooking,
        sort_by=selection.sort_by,
        sort_order=selection.sort_order or ("asc" if selection.sort_by is not None else None),
    )


def _optional_text(value: object, maximum: int) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    if len(text) > maximum:
        raise ValueError("client_registry_query_invalid")
    return text


def _yes_no(value: bool | None) -> str | None:
    if value is None:
        return None
    return "是" if value else "否"


__all__ = [
    "ClientRegistryOrderAccountingExportApplication",
    "HEADERS",
    "OrderAccountingExportQuery",
    "OrderAccountingExportRepository",
    "OrderAccountingExportRow",
    "XLSX_MEDIA_TYPE",
    "build_order_accounting_workbook",
]
