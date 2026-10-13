"""Case-centered, read-only inspection workbook with separate detail grains."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from typing import NamedTuple, Protocol

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from subsystems.client_profile.order_accounting_export import (
    OrderAccountingExportQuery, OrderAccountingExportRow,
    _normalize_query, build_order_accounting_workbook,
)


class CaseRecord(NamedTuple):
    case_no: str
    customer_name: str | None
    client_id: int
    staff_id: int | None
    staff_name: str | None
    status: str
    lifecycle_version: int
    client_profile_version: int
    finance_version: int | None
    terms_version: int | None
    service_type: str | None
    start_date: date | None
    end_date: date | None
    actual_start_date: date | None
    actual_end_date: date | None
    service_start_time: str | None
    service_end_time: str | None
    service_end_day_offset: int | None
    custom_rest_dates: str | None
    service_days: int | None
    service_hours_per_day: Decimal | None
    requires_cooking: bool | None
    floor_fee: Decimal | None
    client_hourly_rate_ntd: int | None
    created_at: datetime | None
    updated_at: datetime | None


class AssignmentRecord(NamedTuple):
    case_no: str
    customer_name: str | None
    assignment_id: int
    staff_id: int
    staff_name: str | None
    generation_id: int | None
    assignment_sequence: int
    assigned_start_date: date | None
    assigned_end_date: date | None
    original_assigned_start_date: date | None
    original_assigned_end_date: date | None
    planned_hours: Decimal | None
    actual_hours: Decimal | None
    hourly_rate: Decimal | None
    floor_fee_allocated: Decimal | None
    status: str
    replaced_assignment_id: int | None


class ScheduleRecord(NamedTuple):
    case_no: str
    customer_name: str | None
    schedule_id: int
    assignment_id: int | None
    staff_id: int
    staff_name: str | None
    generation_id: int | None
    work_date: date
    is_work_day: bool
    is_double_pay: bool
    effective_marker: bool | None


class ClientObligationRecord(NamedTuple):
    case_no: str
    customer_name: str | None
    obligation_identity: str
    obligation_type: str
    direction: str
    source_obligation_identity: str | None
    amount_due_ntd: int
    due_date: date | None
    status: str
    current_event_id: int | None
    projection_version: int


class LedgerRecord(NamedTuple):
    case_no: str
    customer_name: str | None
    ledger_entry_id: int
    finance_import_row_id: int | None
    entry_type: str
    amount_ntd: int
    occurred_on: date
    reconciliation_reference: str | None
    reversal_of_entry_id: int | None


class AllocationRecord(NamedTuple):
    case_no: str
    customer_name: str | None
    allocation_id: int
    ledger_entry_id: int
    obligation_identity: str
    obligation_type: str | None
    direction: str | None
    amount_ntd: int
    allocation_ordinal: int
    occurred_on: date


class StaffObligationRecord(NamedTuple):
    case_no: str
    customer_name: str | None
    obligation_identity: str
    assignment_id: int | None
    staff_id: int
    staff_name: str | None
    obligation_kind: str
    direction: str
    source_obligation_identity: str | None
    amount_due_ntd: int
    due_date: date | None
    status: str
    current_event_id: int | None
    payroll_version: int
    payout_history_exists: bool


class StaffTransferRecord(NamedTuple):
    case_no: str
    customer_name: str | None
    staff_id: int
    staff_name: str | None
    assignment_id: int | None
    settlement_detail_id: int
    transfer_id: int
    allocation_id: int
    allocated_amount: Decimal
    component_type: str
    allocation_review_status: str | None
    transaction_type: str
    transaction_status: str
    occurred_at: datetime
    external_reference: str | None
    reversal_of_transfer_id: int | None
    reversal_of_allocation_id: int | None


@dataclass(frozen=True, slots=True)
class CaseInspectionData:
    overview: tuple[OrderAccountingExportRow, ...]
    cases: tuple[CaseRecord, ...] = ()
    assignments: tuple[AssignmentRecord, ...] = ()
    schedules: tuple[ScheduleRecord, ...] = ()
    client_obligations: tuple[ClientObligationRecord, ...] = ()
    ledger: tuple[LedgerRecord, ...] = ()
    allocations: tuple[AllocationRecord, ...] = ()
    staff_obligations: tuple[StaffObligationRecord, ...] = ()
    staff_transfers: tuple[StaffTransferRecord, ...] = ()


class CaseInspectionRepository(Protocol):
    def query_case_inspection(self, selection: OrderAccountingExportQuery) -> CaseInspectionData: ...


class ClientRegistryCaseInspectionExportApplication:
    def __init__(self, repository: CaseInspectionRepository) -> None:
        self._repository = repository

    def export(self, selection: OrderAccountingExportQuery) -> bytes:
        return build_case_inspection_workbook(self._repository.query_case_inspection(_normalize_query(selection)))


def build_case_inspection_workbook(data: CaseInspectionData) -> bytes:
    # Reuse the canonical finance/subsidy summary, without recalculating owner facts.
    workbook = load_workbook(BytesIO(build_order_accounting_workbook(data.overview)))
    workbook.active.title = "案件總覽"
    sheets = (
        ("案件保存資料", CaseRecord, data.cases, (
            "案件編號", "客戶姓名", "客戶ID", "訂單月嫂ID", "訂單月嫂姓名", "訂單狀態",
            "訂單版本", "客戶版本", "客戶帳務版本", "付款條件版本", "服務方式", "預計開始日", "預計結束日",
            "實際開始日", "實際結束日", "每日開始時間", "每日結束時間", "跨日天數", "自訂休息日",
            "服務天數", "每日時數", "需要下廚", "樓層費", "客戶每小時單價", "建立時間", "更新時間")),
        ("派案明細", AssignmentRecord, data.assignments, (
            "案件編號", "客戶姓名", "派案ID", "月嫂ID", "月嫂姓名", "排程代次ID", "派案順序",
            "派案開始日", "派案結束日", "原派案開始日", "原派案結束日", "預計工時", "實際工時",
            "月嫂時薪", "分配樓層費", "派案狀態", "被替換派案ID")),
        ("每日排班", ScheduleRecord, data.schedules, (
            "案件編號", "客戶姓名", "排班ID", "派案ID", "月嫂ID", "月嫂姓名", "排程代次ID",
            "日期", "工作日", "雙倍薪資日", "生效標記")),
        ("客戶義務", ClientObligationRecord, data.client_obligations, (
            "案件編號", "客戶姓名", "義務識別碼", "義務類型", "方向", "來源義務識別碼", "應付金額",
            "到期日", "狀態", "目前事件ID", "投影版本")),
        ("客戶收付流水", LedgerRecord, data.ledger, (
            "案件編號", "客戶姓名", "流水ID", "銀行匯入列ID", "流水類型", "金額", "發生日", "對帳參照", "沖銷來源流水ID")),
        ("客戶款項分配", AllocationRecord, data.allocations, (
            "案件編號", "客戶姓名", "分配ID", "流水ID", "義務識別碼", "義務類型", "方向", "分配金額", "分配順序", "發生日")),
        ("月嫂應付", StaffObligationRecord, data.staff_obligations, (
            "案件編號", "客戶姓名", "義務識別碼", "派案ID", "月嫂ID", "月嫂姓名", "義務類型", "方向",
            "來源義務識別碼", "應付金額", "到期日", "狀態", "目前事件ID", "薪資版本", "已有付款歷史")),
        ("月嫂付款分配", StaffTransferRecord, data.staff_transfers, (
            "案件編號", "客戶姓名", "月嫂ID", "月嫂姓名", "派案ID", "結算明細ID", "轉帳ID", "分配ID",
            "本案分配金額", "款項組成", "分配審核狀態", "交易類型", "交易狀態", "發生時間", "外部參照",
            "沖銷來源轉帳ID", "沖銷來源分配ID")),
    )
    for title, row_type, rows, headers in sheets:
        worksheet = workbook.create_sheet(title)
        worksheet.append(headers)
        for record in rows:
            worksheet.append(tuple(
                str(value) if value is not None and field.endswith("_id") else value
                for field, value in zip(row_type._fields, record, strict=True)
            ))
        worksheet.freeze_panes = "C2"
        worksheet.auto_filter.ref = worksheet.dimensions
        worksheet.sheet_view.showGridLines = False
        worksheet.row_dimensions[1].height = 32
        for cell in worksheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(fill_type="solid", fgColor="233D58")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for column, field in enumerate(row_type._fields, 1):
            worksheet.column_dimensions[get_column_letter(column)].width = 24 if "identity" in field or "reference" in field else 18
            for cells in worksheet.iter_rows(min_row=2, min_col=column, max_col=column):
                cell = cells[0]
                if isinstance(cell.value, datetime):
                    cell.number_format = "yyyy-mm-dd hh:mm:ss"
                elif isinstance(cell.value, date):
                    cell.number_format = "yyyy-mm-dd"
                elif isinstance(cell.value, (int, float, Decimal)) and not isinstance(cell.value, bool):
                    cell.number_format = "#,##0.##"
                elif isinstance(cell.value, str):
                    cell.number_format = "@"
    # DB text is data, including values beginning with '='. Never create formulas.
    for worksheet in workbook:
        for cells in worksheet.iter_rows():
            for cell in cells:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
