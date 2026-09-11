"""
File: reconciliation_register_query.py
Description: 依既有補助公式建立成立訂單的季度、年度、營運年度統計及正式送件期間唯讀核銷資料。
"""

from __future__ import annotations

import json
from datetime import date, datetime, timedelta
from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO
from typing import Any, Callable

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

QUARTERLY_HEADERS = (
    "\u5e8f\u865f", "\u5e02\u5e9c\u8a02\u55ae\u865f\u78bc", "\u88dc\u52a9\u8cc7\u683c", "\u670d\u52d9\u958b\u59cb", "\u670d\u52d9\u7d50\u675f",
    "\u88dc\u52a9\u6642\u6578", "\u88dc\u52a9\u5929\u6578", "\u670d\u52d9\u5929\u6578", "\u88dc\u52a9\u6b3e\u91d1\u984d", "\u55ae\u50f9",
    "\u96c7\u4e3b", "\u670d\u52d9\u4eba\u54e1", "\u8eab\u5206\u8b49\u5b57\u865f", "\u5730\u5740", "\u7c3d\u9818",
)
ANNUAL_HEADERS = (
    "\u5e8f\u865f", "\u5e02\u5e9c\u8a02\u55ae\u865f\u78bc", "\u88dc\u52a9\u8cc7\u683c", "\u670d\u52d9\u958b\u59cb", "\u670d\u52d9\u7d50\u675f",
    "\u670d\u52d9\u5929\u6578", "\u88dc\u52a9\u6b3e\u91d1\u984d", "\u55ae\u50f9", "\u96c7\u4e3b", "\u670d\u52d9\u4eba\u54e1",
)

GENERAL_CITIZEN = "\u4e00\u822c\u5e02\u6c11"
SUBSIDIZED_CITIZEN = "\u88dc\u52a9\u5e02\u6c11"
IDENTITY_CARD_KEY = "\u8eab\u5206\u8b49\u5b57\u865f"
CLAIMED_BATCH_STATUSES = ("submitted", "approved", "partially_paid", "paid")
ESTABLISHED_ORDER_STATUSES = ("訂單成立", "服務中", "訂單完成")
OPERATIONS_REPORT_ORDER_STATUSES = ESTABLISHED_ORDER_STATUSES + (
    "歷史訂單－未服務",
    "歷史訂單－服務中",
    "歷史訂單－服務完成",
    "歷史訂單－帳務完成",
)
COMPLETED_ORDER_STATUSES = (
    "訂單完成",
    "歷史訂單－服務完成",
    "歷史訂單－帳務完成",
)
RECONCILIATION_QUARTER_LABELS = ("第一季", "第二季", "第三季", "第四季")


def _as_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        try:
            return datetime.strptime(value[:10], "%Y-%m-%d").date()
        except ValueError:
            return None
    return None


def _validate_year_and_quarter(application_year: int, quarter: int) -> None:
    if not isinstance(application_year, int) or application_year < 1912:
        raise ValueError("application_year must be a Gregorian year")
    if quarter not in (1, 2, 3, 4):
        raise ValueError("quarter must be 1, 2, 3, or 4")


def _decode_legacy_key(key: object) -> str:
    text = str(key)
    for encoding in ("latin1", "cp1252"):
        try:
            decoded = text.encode(encoding).decode("utf-8")
        except (UnicodeDecodeError, UnicodeEncodeError):
            continue
        if IDENTITY_CARD_KEY in decoded:
            return decoded
    return text


def extract_employer_identity_card(survey_details: object) -> str:
    """Read the employer identity card from current or legacy-encoded survey JSON."""
    if isinstance(survey_details, str):
        try:
            survey_details = json.loads(survey_details)
        except json.JSONDecodeError:
            return ""
    if not isinstance(survey_details, dict):
        return ""
    for key, value in survey_details.items():
        if IDENTITY_CARD_KEY in _decode_legacy_key(key):
            return str(value or "").strip()
    return ""


def _fetch_completed_cases(connection_factory: Callable[[], Any]) -> list[dict]:
    conn = connection_factory()
    try:
        with conn.cursor() as cursor:
            cursor.execute(
                """
                SELECT o.case_no, c.identity_status, o.actual_start_date,
                       o.actual_end_date, o.service_days, o.service_hours_per_day,
                       c.name AS employer_name, c.address AS employer_address,
                       s.name AS staff_name, br.survey_details
                FROM orders o
                JOIN clients c ON c.id = o.client_id
                LEFT JOIN staff s ON s.id = o.staff_id
                LEFT JOIN beclass_records br ON (br.query_no = o.case_no OR br.bound_case_no = o.case_no)
                WHERE o.actual_end_date IS NOT NULL
                  AND c.identity_status IN (%s, %s)
                ORDER BY o.case_no
                """,
                (GENERAL_CITIZEN, SUBSIDIZED_CITIZEN),
            )
            return cursor.fetchall()
    finally:
        conn.close()


def _fetch_established_cases(
    period_start: date,
    period_end: date,
    connection_factory: Callable[[], Any],
    order_statuses: tuple[str, ...] = ESTABLISHED_ORDER_STATUSES,
) -> list[dict]:
    status_placeholders = ", ".join("%s" for _ in order_statuses)
    conn = connection_factory()
    try:
        with conn.cursor() as cursor:
            cursor.execute(
                f"""
                SELECT o.case_no, o.status AS order_status, c.identity_status,
                       COALESCE(o.actual_start_date, o.start_date) AS actual_start_date,
                       COALESCE(o.actual_end_date, o.end_date) AS actual_end_date,
                       o.service_days, o.service_hours_per_day,
                       c.name AS employer_name, c.address AS employer_address,
                       s.name AS staff_name, br.survey_details
                FROM orders o
                JOIN clients c ON c.id = o.client_id
                LEFT JOIN staff s ON s.id = o.staff_id
                LEFT JOIN beclass_records br
                    ON (br.query_no = o.case_no OR br.bound_case_no = o.case_no)
                WHERE o.status IN ({status_placeholders})
                  AND c.identity_status IN (%s, %s)
                  AND COALESCE(o.actual_end_date, o.end_date) >= %s
                  AND COALESCE(o.actual_end_date, o.end_date) < %s
                ORDER BY o.case_no
                """,
                (
                    *order_statuses,
                    GENERAL_CITIZEN,
                    SUBSIDIZED_CITIZEN,
                    period_start,
                    period_end,
                ),
            )
            return cursor.fetchall()
    finally:
        conn.close()


def _fetch_claim_submission_period_cases(
    period_start: date,
    period_end: date,
    connection_factory: Callable[[], Any],
) -> list[dict]:
    conn = connection_factory()
    try:
        with conn.cursor() as cursor:
            cursor.execute(
                """
                SELECT o.case_no, c.identity_status,
                       COALESCE(o.actual_start_date, a.assigned_start_date) AS actual_start_date,
                       COALESCE(o.actual_end_date, a.assigned_end_date) AS actual_end_date,
                       o.service_days, o.service_hours_per_day,
                       c.name AS employer_name, c.address AS employer_address,
                       s.name AS staff_name, br.survey_details,
                       item.claimed_hours, item.unit_price,
                       item.requested_amount, batch.application_year,
                       batch.quarter, batch.status AS claim_status,
                       batch.submitted_at AS claim_submitted_at
                FROM subsidy_claim_batch_items item
                JOIN subsidy_claim_batches batch ON batch.id = item.batch_id
                JOIN orders o ON o.case_no = item.case_no
                JOIN clients c ON c.id = o.client_id
                JOIN case_staff_assignments a ON a.id = item.assignment_id
                    AND a.case_no = item.case_no
                LEFT JOIN staff s ON s.id = item.staff_id
                LEFT JOIN beclass_records br
                    ON (br.query_no = o.case_no OR br.bound_case_no = o.case_no)
                WHERE batch.status IN (%s, %s, %s, %s)
                  AND batch.submitted_at >= %s
                  AND batch.submitted_at < %s
                  AND batch.revision = (
                      SELECT MAX(latest.revision)
                      FROM subsidy_claim_batches latest
                      WHERE latest.application_year = batch.application_year
                        AND latest.quarter = batch.quarter
                        AND latest.status IN (%s, %s, %s, %s)
                  )
                  AND c.identity_status IN (%s, %s)
                ORDER BY o.case_no, item.id
                """,
                (
                    *CLAIMED_BATCH_STATUSES,
                    period_start,
                    period_end + timedelta(days=1),
                    *CLAIMED_BATCH_STATUSES,
                    GENERAL_CITIZEN,
                    SUBSIDIZED_CITIZEN,
                ),
            )
            return cursor.fetchall()
    finally:
        conn.close()


def _fetch_claim_batch_cases(
    application_year: int,
    quarter: int | None,
    connection_factory: Callable[[], Any],
) -> list[dict]:
    conn = connection_factory()
    try:
        with conn.cursor() as cursor:
            cursor.execute(
                """
                SELECT o.case_no, c.identity_status,
                       COALESCE(o.actual_start_date, a.assigned_start_date) AS actual_start_date,
                       COALESCE(o.actual_end_date, a.assigned_end_date) AS actual_end_date,
                       o.service_days, o.service_hours_per_day,
                       c.name AS employer_name, c.address AS employer_address,
                       s.name AS staff_name, br.survey_details,
                       item.claimed_hours, item.unit_price,
                       item.requested_amount, batch.application_year,
                       batch.quarter, batch.status AS claim_status,
                       batch.submitted_at AS claim_submitted_at
                FROM subsidy_claim_batch_items item
                JOIN subsidy_claim_batches batch ON batch.id = item.batch_id
                JOIN orders o ON o.case_no = item.case_no
                JOIN clients c ON c.id = o.client_id
                JOIN case_staff_assignments a ON a.id = item.assignment_id
                    AND a.case_no = item.case_no
                LEFT JOIN staff s ON s.id = item.staff_id
                LEFT JOIN beclass_records br
                    ON (br.query_no = o.case_no OR br.bound_case_no = o.case_no)
                WHERE batch.status IN (%s, %s, %s, %s)
                  AND batch.application_year = %s
                  AND (%s IS NULL OR batch.quarter = %s)
                  AND batch.revision = (
                      SELECT MAX(latest.revision)
                      FROM subsidy_claim_batches latest
                      WHERE latest.application_year = batch.application_year
                        AND latest.quarter = batch.quarter
                        AND latest.status IN (%s, %s, %s, %s)
                  )
                  AND c.identity_status IN (%s, %s)
                ORDER BY batch.quarter, o.case_no, item.id
                """,
                (
                    *CLAIMED_BATCH_STATUSES,
                    application_year,
                    quarter,
                    quarter,
                    *CLAIMED_BATCH_STATUSES,
                    GENERAL_CITIZEN,
                    SUBSIDIZED_CITIZEN,
                ),
            )
            return cursor.fetchall()
    finally:
        conn.close()


def _subsidy_terms(eligibility: str, total_service_hours: Decimal) -> tuple[Decimal, Decimal]:
    if eligibility == GENERAL_CITIZEN:
        return min(Decimal("40"), total_service_hours), Decimal("300")
    if eligibility == SUBSIDIZED_CITIZEN:
        return min(Decimal("120"), total_service_hours), Decimal("350")
    return Decimal("0"), Decimal("0")


def _to_register_row(source: dict) -> dict | None:
    actual_start = _as_date(source.get("actual_start_date"))
    actual_end = _as_date(source.get("actual_end_date"))
    daily_hours = Decimal(str(source.get("service_hours_per_day") or 0))
    service_days = Decimal(str(source.get("service_days") or 0))
    subsidy_hours, unit_price = _subsidy_terms(
        source.get("identity_status"), service_days * daily_hours,
    )
    if not actual_start or not actual_end or subsidy_hours <= 0 or daily_hours <= 0:
        return None

    return {
        "市府訂單號碼": str(source["case_no"]),
        "補助資格": source["identity_status"],
        "服務開始": actual_start,
        "服務結束": actual_end,
        "補助時數": subsidy_hours,
        "補助天數": (subsidy_hours / daily_hours).quantize(
            Decimal("0.01"), rounding=ROUND_HALF_UP
        ),
        "服務天數": source.get("service_days") or 0,
        "補助款金額": subsidy_hours * unit_price,
        "單價": unit_price,
        "雇主": source.get("employer_name") or "",
        "服務人員": source.get("staff_name") or "",
        "身分證字號": extract_employer_identity_card(source.get("survey_details")),
        "地址": source.get("employer_address") or "",
        "簽領": "",
        "訂單狀態": source.get("order_status") or "",
    }


def _to_claim_register_row(source: dict) -> dict | None:
    actual_start = _as_date(source.get("actual_start_date"))
    actual_end = _as_date(source.get("actual_end_date"))
    submitted_at = _as_date(source.get("claim_submitted_at"))
    claimed_hours = Decimal(str(source.get("claimed_hours") or 0))
    daily_hours = Decimal(str(source.get("service_hours_per_day") or 0))
    unit_price = Decimal(str(source.get("unit_price") or 0))
    requested_amount = Decimal(str(source.get("requested_amount") or 0))
    if (
        actual_start is None
        or actual_end is None
        or submitted_at is None
        or claimed_hours <= 0
        or daily_hours <= 0
        or unit_price <= 0
        or requested_amount < 0
    ):
        return None
    subsidy_days = (claimed_hours / daily_hours).quantize(
        Decimal("0.01"), rounding=ROUND_HALF_UP
    )
    return {
        "市府訂單號碼": str(source["case_no"]),
        "補助資格": source["identity_status"],
        "服務開始": actual_start,
        "服務結束": actual_end,
        "補助時數": claimed_hours,
        "補助天數": subsidy_days,
        "服務天數": int(source.get("service_days") or subsidy_days),
        "補助款金額": requested_amount,
        "單價": unit_price,
        "雇主": source.get("employer_name") or "",
        "服務人員": source.get("staff_name") or "",
        "身分證字號": extract_employer_identity_card(source.get("survey_details")),
        "地址": source.get("employer_address") or "",
        "簽領": "",
        "核銷月份": f"{submitted_at.year:04d}-{submitted_at.month:02d}",
        "核銷狀態": _claim_status_label(source.get("claim_status")),
    }


def _claim_status_label(value: object) -> str:
    return {
        "submitted": "已送件",
        "approved": "已核准",
        "partially_paid": "部分撥款",
        "paid": "已撥款",
    }.get(str(value or ""), "已送件")


def _partition_rows(rows: list[dict]) -> tuple[list[dict], list[dict]]:
    general = []
    subsidized = []
    for row in rows:
        (subsidized if row["補助資格"] == SUBSIDIZED_CITIZEN else general).append(row)
    return general, subsidized


def _with_serials(rows: list[dict]) -> list[dict]:
    return [{"序號": index, **row} for index, row in enumerate(rows, start=1)]


def _populate_subsidy_sheet(
    worksheet,
    headers: tuple[str, ...],
    general_rows: list[dict],
    subsidized_rows: list[dict],
) -> None:
    yellow_fill = PatternFill(fill_type="solid", fgColor="FFFF00")

    def append_section(rows: list[dict], label: str) -> None:
        worksheet.append((label,))
        worksheet.append(headers)
        header_row = worksheet.max_row
        for cell in worksheet[header_row]:
            cell.fill = yellow_fill
            cell.font = Font(bold=True)
        for row in rows:
            worksheet.append(tuple(row.get(header, "") for header in headers))
            if "補助天數" in headers:
                worksheet.cell(worksheet.max_row, headers.index("補助天數") + 1).number_format = "0.00"

    append_section(general_rows, GENERAL_CITIZEN)
    if subsidized_rows:
        worksheet.append(())
        append_section(subsidized_rows, SUBSIDIZED_CITIZEN)

    worksheet.freeze_panes = "A3"
    for column in ("D", "E"):
        worksheet.column_dimensions[column].width = 14
    for column in ("B", "K", "L", "M", "N", "O"):
        worksheet.column_dimensions[column].width = 18


def _build_workbook(headers: tuple[str, ...], general_rows: list[dict], subsidized_rows: list[dict], title: str) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = title
    _populate_subsidy_sheet(worksheet, headers, general_rows, subsidized_rows)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _established_order_rows(
    application_year: int,
    quarter: int | None,
    connection_factory: Callable[[], Any],
    order_statuses: tuple[str, ...] = ESTABLISHED_ORDER_STATUSES,
) -> tuple[list[dict], list[dict]]:
    start_month = (quarter - 1) * 3 + 1 if quarter is not None else 1
    period_start = date(application_year, start_month, 1)
    if quarter is None or quarter == 4:
        period_end = date(application_year + 1, 1, 1)
    else:
        period_end = date(application_year, start_month + 3, 1)
    rows = []
    for source in _fetch_established_cases(
        period_start,
        period_end,
        connection_factory,
        order_statuses,
    ):
        row = _to_register_row(source)
        if row is None:
            continue
        service_end = row["服務結束"]
        if not period_start <= service_end < period_end:
            continue
        rows.append(row)
    rows.sort(key=lambda row: row["\u5e02\u5e9c\u8a02\u55ae\u865f\u78bc"])
    general, subsidized = _partition_rows(rows)
    return _with_serials(general), _with_serials(subsidized)


def build_year_to_date_subsidy_rows(
    application_year: int,
    cutoff_date: date,
    connection_factory: Callable[[], Any],
) -> dict:
    """Return owner-calculated rows completed from January 1 through cutoff_date."""
    if not isinstance(application_year, int) or application_year < 1912:
        raise ValueError("application_year must be a Gregorian year")
    if not isinstance(cutoff_date, date) or cutoff_date < date(application_year, 1, 1):
        raise ValueError("cutoff_date must not precede application_year")
    rows = []
    for source in _fetch_completed_cases(connection_factory):
        row = _to_register_row(source)
        if row is None:
            continue
        completion = row["服務結束"]
        if completion.year == application_year and completion <= cutoff_date:
            rows.append(row)
    rows.sort(key=lambda row: row["市府訂單號碼"])
    general, subsidized = _partition_rows(rows)
    return {
        "general_citizen_rows": _with_serials(general),
        "subsidized_citizen_rows": _with_serials(subsidized),
    }


def build_claim_submission_period_subsidy_rows(
    period_start: date,
    period_end: date,
    connection_factory: Callable[[], Any],
) -> dict:
    """Return formal claim items submitted inside the selected inclusive period."""
    if (
        not isinstance(period_start, date)
        or not isinstance(period_end, date)
        or period_start > period_end
    ):
        raise ValueError("period_start must not follow period_end")
    rows = [
        row
        for source in _fetch_claim_submission_period_cases(
            period_start, period_end, connection_factory
        )
        if (row := _to_claim_register_row(source)) is not None
    ]
    rows.sort(key=lambda row: row["市府訂單號碼"])
    general, subsidized = _partition_rows(rows)
    return {
        "general_citizen_rows": _with_serials(general),
        "subsidized_citizen_rows": _with_serials(subsidized),
    }


def build_quarterly_subsidy_register(
    application_year: int,
    quarter: int,
    connection_factory: Callable[[], Any],
) -> dict:
    """Build a register for established orders ending in the selected quarter."""
    _validate_year_and_quarter(application_year, quarter)
    general_rows, subsidized_rows = _established_order_rows(
        application_year, quarter, connection_factory
    )
    return {
        "general_citizen_rows": general_rows,
        "subsidized_citizen_rows": subsidized_rows,
        "xlsx_bytes": _build_workbook(QUARTERLY_HEADERS, general_rows, subsidized_rows, "\u5206\u5b63\u6838\u92b7"),
    }


def build_annual_subsidy_summary(
    application_year: int,
    connection_factory: Callable[[], Any],
) -> dict:
    """Build an annual summary for established orders ending in that year."""
    if not isinstance(application_year, int) or application_year < 1912:
        raise ValueError("application_year must be a Gregorian year")
    general_rows, subsidized_rows = _established_order_rows(
        application_year, None, connection_factory
    )
    return {
        "general_citizen_rows": general_rows,
        "subsidized_citizen_rows": subsidized_rows,
        "xlsx_bytes": _build_workbook(ANNUAL_HEADERS, general_rows, subsidized_rows, "\u5e74\u5ea6\u7e3d\u8868"),
    }


def build_operations_report_annual_subsidy_rows(
    report_year: int,
    connection_factory: Callable[[], Any],
) -> dict:
    """Return the operations-report annual candidate in its dedicated row format."""
    if not isinstance(report_year, int) or report_year < 1912:
        raise ValueError("report_year must be a Gregorian year")
    general_rows, subsidized_rows = _established_order_rows(
        report_year,
        None,
        connection_factory,
        OPERATIONS_REPORT_ORDER_STATUSES,
    )

    def operations_row(row: dict) -> dict:
        result = dict(row)
        reconciliation_year, reconciliation_period = _operations_reconciliation_period(
            result["服務結束"]
        )
        if reconciliation_year != report_year:
            raise RuntimeError("operations_report_reconciliation_year_mismatch")
        result["核銷月份"] = reconciliation_period
        result["核銷狀態"] = (
            "結案" if result.get("訂單狀態") in COMPLETED_ORDER_STATUSES else ""
        )
        return result

    return {
        "general_citizen_rows": [operations_row(row) for row in general_rows],
        "subsidized_citizen_rows": [operations_row(row) for row in subsidized_rows],
    }


def _operations_reconciliation_period(service_end: date) -> tuple[int, str]:
    quarter_index = (service_end.month - 1) // 3
    return service_end.year, RECONCILIATION_QUARTER_LABELS[quarter_index]


def build_combined_subsidy_register(
    application_year: int,
    quarter: int,
    connection_factory: Callable[[], Any],
) -> dict:
    """Build a combined workbook with both quarterly and annual reconciliation sheets."""
    _validate_year_and_quarter(application_year, quarter)
    q_general, q_subsidized = _established_order_rows(
        application_year, quarter, connection_factory
    )
    a_general, a_subsidized = _established_order_rows(
        application_year, None, connection_factory
    )
    workbook = Workbook()
    ws_quarterly = workbook.active
    ws_quarterly.title = "季核銷"
    _populate_subsidy_sheet(ws_quarterly, QUARTERLY_HEADERS, q_general, q_subsidized)
    ws_annual = workbook.create_sheet(title="年度總表")
    _populate_subsidy_sheet(ws_annual, ANNUAL_HEADERS, a_general, a_subsidized)
    output = BytesIO()
    workbook.save(output)
    return {
        "application_year": application_year,
        "quarter": quarter,
        "quarterly_general_rows": q_general,
        "quarterly_subsidized_rows": q_subsidized,
        "annual_general_rows": a_general,
        "annual_subsidized_rows": a_subsidized,
        "xlsx_bytes": output.getvalue(),
    }
