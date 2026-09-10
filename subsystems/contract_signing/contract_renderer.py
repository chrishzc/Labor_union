"""
File: contract_renderer.py
Description: 定義契約 PDF renderer port，並保留核准 XLSX 模板的安全 literal 填值相容能力。
"""

from __future__ import annotations

from dataclasses import dataclass
import hashlib
from io import BytesIO
import json
from pathlib import Path
import re
from typing import Mapping, Protocol, runtime_checkable

from openpyxl import load_workbook

from subsystems.contract_signing.template_catalog import mapping_is_applicable


PDF_MEDIA_TYPE = "application/pdf"
_SHA256 = re.compile(r"^[0-9a-f]{64}$")
_FORMULA_PREFIXES = ("=", "+", "-", "@")
_APPROVED_TEMPLATE_IDS = frozenset(
    {"contract_client_copy", "contract_staff_service"}
)
_MAPPING_REQUIREDNESS = frozenset({"required", "conditional", "optional"})
_EXTERNAL_WORKBOOK_REFERENCE = re.compile(r"\[\d+\]")


def external_formula_cells(content: bytes) -> tuple[str, ...]:
    """Detect unresolved external workbook links before an office engine runs."""
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=False)
    try:
        return tuple(f"{sheet.title}!{cell.coordinate}" for sheet in workbook for row in sheet
                     for cell in row if cell.data_type == 'f' and _EXTERNAL_WORKBOOK_REFERENCE.search(str(cell.value)))
    finally:
        workbook.close()


class ContractRendererError(RuntimeError):
    """Closed renderer failure that never includes executable paths or raw stderr."""

    def __init__(self, code: str, message: str, *, retryable: bool = False) -> None:
        super().__init__(message)
        self.code = code
        self.retryable = retryable


@dataclass(frozen=True, slots=True)
class RenderedContract:
    content: bytes
    mime_type: str
    filename: str
    sha256: str
    renderer_identity: str

    @classmethod
    def from_pdf_bytes(
        cls,
        *,
        content: bytes,
        filename: str,
        renderer_identity: str,
    ) -> "RenderedContract":
        if not content:
            raise ContractRendererError(
                "contract_pdf_renderer_output_empty",
                "契約 PDF renderer 未產生內容。",
            )
        if not content.startswith(b"%PDF-") or not content.rstrip().endswith(b"%%EOF"):
            raise ContractRendererError(
                "contract_pdf_renderer_output_invalid",
                "契約 PDF renderer 產生無效文件。",
            )
        canonical_filename = Path(filename).name
        if canonical_filename != filename or not canonical_filename.lower().endswith(".pdf"):
            raise ContractRendererError(
                "contract_pdf_renderer_filename_invalid",
                "契約 PDF renderer 產生無效檔名。",
            )
        canonical_identity = renderer_identity.strip()
        if not canonical_identity:
            raise ContractRendererError(
                "contract_pdf_renderer_identity_invalid",
                "契約 PDF renderer 缺少版本身分。",
            )
        digest = hashlib.sha256(content).hexdigest()
        if _SHA256.fullmatch(digest) is None:
            raise ContractRendererError(
                "contract_pdf_renderer_digest_invalid",
                "契約 PDF renderer 無法建立完整性摘要。",
            )
        return cls(
            content=content,
            mime_type=PDF_MEDIA_TYPE,
            filename=canonical_filename,
            sha256=digest,
            renderer_identity=canonical_identity,
        )


@runtime_checkable
class ContractRenderer(Protocol):
    def render(
        self,
        *,
        template_path: Path,
        mapping_path: Path,
        facts: Mapping[str, object],
    ) -> RenderedContract: ...


def render_contract_template(
    *, template_path: Path, mapping_path: Path, facts: Mapping[str, object]
) -> bytes:
    """Return the historical XLSX artifact while treating external facts as literals."""

    try:
        mapping = json.loads(mapping_path.read_text(encoding="utf-8"))
    except (OSError, UnicodeError, json.JSONDecodeError) as error:
        raise ContractRendererError(
            "contract_pdf_mapping_invalid",
            "契約 PDF 欄位映射無法讀取。",
        ) from error
    if not isinstance(mapping, dict) or not isinstance(
        mapping.get("param_mappings"), dict
    ):
        raise ContractRendererError(
            "contract_pdf_mapping_invalid",
            "契約 PDF 欄位映射格式無效。",
        )
    approved_template = mapping.get("id") in _APPROVED_TEMPLATE_IDS
    # Old external-link caches include unrelated historical order data. They
    # must not be copied into newly archived contract sources.
    workbook = load_workbook(template_path, keep_links=False)
    worksheet = workbook.active
    # Approved static contract wording is versioned with the mapping digest.
    # It replaces the old workbook's missing catalogue links, not owner facts.
    static_cells = mapping.get('static_cells', {})
    if not isinstance(static_cells, dict):
        raise ContractRendererError('contract_pdf_mapping_invalid', '契約固定內容格式無效。')
    for cell, value in static_cells.items():
        if not isinstance(cell, str) or not re.fullmatch(r'[A-Z]{1,3}[1-9][0-9]{0,6}', cell) or not isinstance(value, str) or cell in mapping['param_mappings']:
            raise ContractRendererError('contract_pdf_mapping_invalid', '契約固定內容格式無效。')
        worksheet[cell] = value
        worksheet[cell].data_type = 's'
    row_heights = mapping.get('row_heights', {})
    if not isinstance(row_heights, dict):
        raise ContractRendererError('contract_pdf_mapping_invalid', '契約列高設定無效。')
    for row, height in row_heights.items():
        if not str(row).isdigit() or not 1 <= int(row) <= 1048576 or isinstance(height, bool) or not isinstance(height, (int, float)) or not 0 < height <= 409:
            raise ContractRendererError('contract_pdf_mapping_invalid', '契約列高設定無效。')
        worksheet.row_dimensions[int(row)].height = height
    for cell, descriptor in mapping["param_mappings"].items():
        if not isinstance(descriptor, dict):
            raise ContractRendererError(
                "contract_pdf_mapping_invalid",
                "契約 PDF 欄位映射格式無效。",
            )
        key = descriptor.get("db_key")
        status = descriptor.get("status")
        requiredness = descriptor.get("requiredness")
        if status == "not_applicable":
            # Real legal templates may still contain legacy placeholder text.
            # Clear the declared cell so staff never need to edit the value.
            if descriptor.get("requiredness") in {"optional", "conditional"}:
                worksheet[cell] = None
                continue
            raise ContractRendererError(
                "contract_pdf_required_mapping_unresolved",
                "契約 PDF 的不適用欄位不能是 required。",
            )
        if status in {"pending", "unresolved"}:
            if requiredness == "optional" or (
                requiredness == "conditional"
                and not mapping_is_applicable(descriptor, facts)
            ):
                worksheet[cell] = None
                continue
            raise ContractRendererError(
                "contract_pdf_required_mapping_unresolved",
                "契約 PDF 仍有欄位缺少核准的 typed owner source。",
            )
        if not isinstance(key, str) or not key:
            raise ContractRendererError(
                "contract_pdf_required_mapping_unresolved",
                "契約 PDF 仍有欄位缺少核准的 typed owner source。",
            )
        if approved_template and requiredness not in _MAPPING_REQUIREDNESS:
            raise ContractRendererError(
                "contract_pdf_required_mapping_unresolved",
                "契約 PDF 欄位缺少核准的 requiredness 與 typed owner source。",
            )
        if key not in facts or facts.get(key) is None:
            if requiredness == "required":
                raise ContractRendererError(
                    "contract_pdf_required_mapping_missing",
                    "契約 PDF 欄位缺少核准的 typed owner source。",
                )
            worksheet[cell] = None
            continue
        value = facts[key]
        target = worksheet[cell]
        target.value = value
        if isinstance(value, str) and value.lstrip().startswith(_FORMULA_PREFIXES):
            target.data_type = "s"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
