"""Authenticated bank archive downloads and bounded case accounting commands."""
from dataclasses import asdict
from io import BytesIO
from typing import Annotated

from fastapi import APIRouter, Depends, Header, Response
from openpyxl import Workbook
from pymysql import MySQLError

from api.dependencies.admin_auth import require_admin
from api.dependencies.finance_import import get_accounting_workspace, get_statement_archive_application
from api.error_contracts import typed_http_error
from api.schemas.base import BaseResponse
from api.schemas.finance_import import (
    AccountingWarningBody, CaseAccountingView, ManualStatementBody,
    StatementCommandReceiptView, StatementDuplicateBody,
)
from domains.finance_import.statement_archive import STATEMENT_COLUMNS
from subsystems.finance_import.accounting_workspace import AccountingConflict, AccountingWarning
from subsystems.finance_import.statement_archive import StatementCommandConflict

router = APIRouter()


def _error(error):
    if isinstance(error, (AccountingConflict, StatementCommandConflict)):
        return typed_http_error(409, "conflict", "accounting_source_changed", "資料已變更，請重新整理後再確認。", "finance-workspace")
    if isinstance(error, LookupError):
        return typed_http_error(404, "not_found", "accounting_case_not_found", "找不到案件。", "finance-workspace")
    if isinstance(error, MySQLError):
        return typed_http_error(503, "unavailable", "accounting_unavailable", "帳務資料暫時無法取得，請稍後重試。", "finance-workspace")
    return typed_http_error(422, "validation", "invalid_statement_command", "請檢查交易日、單一收支金額、銷帳編號或備註及指定的明細。", "finance-workspace")


def _actor(principal):
    return f"admin_user_id:{principal.id}" if principal.id is not None else principal.username


def _case_view(case):
    payload = asdict(case)
    for field in ("client_items", "staff_items", "transactions", "warnings"):
        payload[field] = list(payload[field])
    return CaseAccountingView(**payload)


@router.get("/statements/export")
def export_statements(principal=Depends(require_admin), application=Depends(get_statement_archive_application)):
    del principal
    try:
        rows = application.export_rows()
        workbook = Workbook()
        workbook.remove(workbook.active)
        banks = sorted({row.bank for row in rows}) or ["manual"]
        labels = {"sinopac": "永豐", "taishin": "台新", "legacy": "歷史格式", "manual": "人工新增"}
        for bank in banks:
            selected = [row for row in rows if row.bank == bank]
            columns = list(STATEMENT_COLUMNS) if bank in {"sinopac", "manual"} else []
            for row in selected:
                for name, _ in row.fields:
                    if name not in columns:
                        columns.append(name)
            metadata = []
            for name in ("明細識別", "資料來源", "關聯案件"):
                while name in columns:
                    name = "_" + name
                metadata.append(name)
            sheet = workbook.create_sheet(labels.get(bank, bank) if rows else "無成功明細（0筆）")
            sheet.append([*columns, *metadata])
            for cell in sheet[1]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
            for row in selected:
                fields = dict(row.fields)
                values = [fields.get(column) for column in columns] + [row.identity, row.source, ",".join(row.cases)]
                sheet.append(values)
                # Bank text is untrusted spreadsheet content, never a formula.
                for cell in sheet[sheet.max_row]:
                    if isinstance(cell.value, str):
                        cell.data_type = "s"
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
        content = BytesIO()
        workbook.save(content)
        return Response(content.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={
            "Content-Disposition": 'attachment; filename="bank-statements.xlsx"',
            "X-Statement-Count": str(len(rows)), "Cache-Control": "no-store",
        })
    except (ValueError, MySQLError) as error:
        raise _error(error) from None


@router.post("/statements/manual", response_model=BaseResponse[StatementCommandReceiptView])
def create_manual_statement(
    body: ManualStatementBody,
    idempotency_key: Annotated[str, Header(min_length=1, max_length=191)],
    principal=Depends(require_admin), application=Depends(get_statement_archive_application),
):
    try:
        result = application.create_manual(body.fields, idempotency_key, _actor(principal), body.duplicate_of)
        return BaseResponse(data=StatementCommandReceiptView(**asdict(result)), message="重複明細已排除" if result.excluded_count else "人工明細已保存")
    except (ValueError, MySQLError) as error:
        raise _error(error) from None


@router.post("/statements/duplicates", response_model=BaseResponse[StatementCommandReceiptView])
def exclude_statement_duplicate(
    body: StatementDuplicateBody,
    idempotency_key: Annotated[str, Header(min_length=1, max_length=191)],
    principal=Depends(require_admin), application=Depends(get_statement_archive_application),
):
    try:
        result = application.exclude_duplicate(body.candidate_identity, body.kept_identity, idempotency_key, _actor(principal))
        return BaseResponse(data=StatementCommandReceiptView(**asdict(result)), message="重複明細已排除，保留既存明細")
    except (ValueError, MySQLError) as error:
        raise _error(error) from None


@router.get("/cases/{case_no}/accounting", response_model=BaseResponse[CaseAccountingView])
def case_accounting(case_no: str, principal=Depends(require_admin), application=Depends(get_accounting_workspace)):
    del principal
    try:
        return BaseResponse(data=_case_view(application.query(case_no)), message="成功取得案件帳務")
    except (LookupError, ValueError, MySQLError) as error:
        raise _error(error) from None


@router.get("/accounting/warnings", response_model=BaseResponse[list[CaseAccountingView]])
def accounting_warnings(principal=Depends(require_admin), application=Depends(get_accounting_workspace)):
    del principal
    try:
        return BaseResponse(data=[_case_view(case) for case in application.warnings()], message="成功取得帳務警示")
    except (LookupError, ValueError, MySQLError) as error:
        raise _error(error) from None


@router.post("/cases/{case_no}/accounting/dismiss", response_model=BaseResponse[CaseAccountingView])
def dismiss_accounting_warning(body: AccountingWarningBody, case_no: str, principal=Depends(require_admin), application=Depends(get_accounting_workspace)):
    try:
        result = application.dismiss(case_no, AccountingWarning(**body.model_dump()), _actor(principal))
        return BaseResponse(data=_case_view(result), message="帳務警示已解除")
    except (LookupError, ValueError, MySQLError) as error:
        raise _error(error) from None
