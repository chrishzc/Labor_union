"""Extract bounded raster drawings with their Excel anchor and displayed extent."""
from base64 import b64encode


def worksheet_images(sheet) -> tuple[dict[str, object], ...]:
    def column_width(index):
        from openpyxl.utils import get_column_letter
        return (sheet.column_dimensions[get_column_letter(index)].width or 8.43) * 7

    def row_height(index):
        return (sheet.row_dimensions[index].height or sheet.sheet_format.defaultRowHeight or 15) * 4 / 3

    result = []
    for image in sheet._images:
        if image.width * image.height > 20000000:
            raise ValueError('image pixel capacity exceeded')
        if image.format not in {'png', 'jpeg', 'jpg'}:
            raise ValueError('unsupported raster image')
        anchor = image.anchor
        marker = getattr(anchor, '_from', None)
        if marker is None:
            raise ValueError('unsupported image anchor')
        row, col = marker.row + 1, marker.col + 1
        owner_row, owner_col = row, col
        for merged in sheet.merged_cells.ranges:
            if merged.min_row <= row <= merged.max_row and merged.min_col <= col <= merged.max_col:
                owner_row, owner_col = merged.min_row, merged.min_col
                break
        left = marker.colOff / 9525 + sum(column_width(c) for c in range(owner_col, col))
        top = marker.rowOff / 9525 + sum(row_height(r) for r in range(owner_row, row))
        end = getattr(anchor, 'to', None)
        if end is not None:
            width = sum(column_width(c) for c in range(col, end.col + 1)) + (end.colOff - marker.colOff) / 9525
            height = sum(row_height(r) for r in range(row, end.row + 1)) + (end.rowOff - marker.rowOff) / 9525
        else:
            extent = getattr(anchor, 'ext', None)
            width = extent.cx / 9525 if extent else image.width
            height = extent.cy / 9525 if extent else image.height
        data = image._data()
        if len(data) > 8 * 1024 * 1024:
            raise ValueError('image too large')
        mime = 'image/png' if image.format == 'png' else 'image/jpeg'
        result.append({'row': owner_row, 'col': owner_col, 'left': left, 'top': top,
                       'width': max(1, width), 'height': max(1, height),
                       'data_url': f'data:{mime};base64,' + b64encode(data).decode('ascii')})
    return tuple(result)
