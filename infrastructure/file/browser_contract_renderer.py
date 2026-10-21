"""Render approved XLSX contract content to PDF with headless Chromium."""

from __future__ import annotations

from collections.abc import Callable
from datetime import date, datetime
from decimal import Decimal
from html import escape
from io import BytesIO
import os
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, range_boundaries

from subsystems.contract_signing.contract_renderer import (
    ContractRendererError,
    RenderedContract,
    external_formula_cells,
    render_contract_template,
)
from subsystems.contract_signing.template_catalog import CONTRACT_PDF_PRESENTATION_VERSION

_RENDERER_IDENTITY = CONTRACT_PDF_PRESENTATION_VERSION
_MAX_PDF_BYTES = 20 * 1024 * 1024


class BrowserContractRenderer:
    def __init__(self, pdf_generator: Callable[[str], bytes] | None = None) -> None:
        self._pdf_generator = pdf_generator or _render_html_pdf

    def render(self, *, template_path: Path, mapping_path: Path, facts) -> RenderedContract:
        content = render_contract_template(
            template_path=template_path,
            mapping_path=mapping_path,
            facts=facts,
        )
        return self.render_workbook(content=content, filename=template_path.name)

    def render_workbook(self, *, content: bytes, filename: str) -> RenderedContract:
        source_name = Path(filename).name
        if source_name != filename or not source_name.lower().endswith(".xlsx"):
            raise ContractRendererError(
                "contract_pdf_renderer_source_invalid", "契約 PDF 來源格式無效。"
            )
        try:
            if external_formula_cells(content):
                raise ContractRendererError(
                    "contract_pdf_external_reference_unresolved",
                    "契約模板仍引用缺少的舊版表格內容，暫不能產生可簽署 PDF。",
                )
            html = _workbook_html(content)
            pdf = self._pdf_generator(html)
        except ContractRendererError:
            raise
        except Exception:
            raise ContractRendererError(
                "contract_pdf_renderer_unavailable",
                "契約 PDF 瀏覽器 renderer 無法使用。",
                retryable=True,
            ) from None
        if len(pdf) > _MAX_PDF_BYTES:
            raise ContractRendererError(
                "contract_pdf_renderer_output_too_large", "契約 PDF 超過大小限制。"
            )
        return RenderedContract.from_pdf_bytes(
            content=pdf,
            filename=f"{Path(source_name).stem}.pdf",
            renderer_identity=_RENDERER_IDENTITY,
        )


def _render_html_pdf(html: str) -> bytes:
    try:
        from playwright.sync_api import Error as PlaywrightError
        from playwright.sync_api import sync_playwright

        with sync_playwright() as runtime:
            executable = os.getenv("CONTRACT_PDF_BROWSER_EXECUTABLE", "").strip()
            launch: dict[str, Any] = {"headless": True}
            if executable:
                launch["executable_path"] = executable
            try:
                browser = runtime.chromium.launch(**launch)
            except PlaywrightError:
                if executable or os.name != "nt":
                    raise
                browser = runtime.chromium.launch(channel="msedge", headless=True)
            try:
                page = browser.new_page()
                page.set_content(html, wait_until="load")
                page.emulate_media(media="print")
                return page.pdf(
                    format="A4",
                    print_background=True,
                    prefer_css_page_size=True,
                    margin={
                        "top": "8mm",
                        "right": "8mm",
                        "bottom": "8mm",
                        "left": "8mm",
                    },
                )
            finally:
                browser.close()
    except Exception:
        raise ContractRendererError(
            "contract_pdf_renderer_unavailable",
            "契約 PDF 瀏覽器 renderer 無法啟動。",
            retryable=True,
        ) from None


def render_contract_print_html(content: bytes, *, title: str, missing_fields: bool = False) -> str:
    """Produce the full approved document without starting a server browser."""
    html = _workbook_html(content, browser_print=True)
    if missing_fields:
        html = html.replace(
            '<aside class="print-controls">',
            '<aside class="print-controls"><p>部分契約欄位尚未提供，已留白；請核對目前資料。</p>',
            1,
        )
    return html.replace("<head>", f"<head><title>{escape(title)}</title>", 1)


def _workbook_html(content: bytes, *, browser_print: bool = False) -> str:
    workbook = None
    try:
        workbook = load_workbook(BytesIO(content), data_only=True)
        worksheet = workbook.active
        min_col, min_row, max_col, max_row = _print_bounds(worksheet)
        merged = _merged_cells(worksheet, min_row, max_row, min_col, max_col)
        covered = {
            (row, column)
            for start, (rowspan, colspan) in merged.items()
            for row in range(start[0], start[0] + rowspan)
            for column in range(start[1], start[1] + colspan)
            if (row, column) != start
        }
        total_width = sum(Decimal(_column_width(worksheet, column)) for column in range(min_col, max_col + 1))
        # Scale fonts and row heights together with columns into A4's 194 mm content width.
        scale = min(1.0, (194 * 72 / 25.4) / float(total_width))
        columns = "".join(
            f'<col style="width:{Decimal(_column_width(worksheet, column)) / total_width * 100:.4f}%">'
            for column in range(min_col, max_col + 1)
        )
        rows: list[str] = []
        for row in range(min_row, max_row + 1):
            height = worksheet.row_dimensions[row].height
            cells: list[str] = []
            for column in range(min_col, max_col + 1):
                if (row, column) in covered:
                    continue
                cell = worksheet.cell(row, column)
                rowspan, colspan = merged.get((row, column), (1, 1))
                background = _aligned_fill(
                    worksheet, row, column, colspan, min_col, max_col
                )
                # Excel lets labels flow into empty neighbours. Give that text
                # real layout space so later cell backgrounds cannot paint over it.
                if cell.value is not None and rowspan == 1:
                    while column + colspan <= max_col:
                        neighbour = worksheet.cell(row, column + colspan)
                        if (
                            neighbour.value is not None
                            or (row, column + colspan) in merged
                            or (row, column + colspan) in covered
                            or _aligned_fill(worksheet, row, column + colspan, 1, min_col, max_col) != background
                            or any(
                                getattr(getattr(candidate.border, side), "style", None)
                                for candidate in (cell, neighbour)
                                for side in ("top", "right", "bottom", "left")
                            )
                        ):
                            break
                        covered.add((row, column + colspan))
                        colspan += 1
                span = (f' rowspan="{rowspan}"' if rowspan > 1 else "") + (
                    f' colspan="{colspan}"' if colspan > 1 else ""
                )
                cells.append(
                    f'<td{span} style="{_cell_style(cell, colspan, background, scale, browser_print=browser_print)}">{_cell_text(cell.value)}</td>'
                )
            row_style = f' style="height:{height * scale:.2f}pt"' if height else ""
            rows.append(f"<tr{row_style}>{''.join(cells)}</tr>")
        print_style = (
            "html{background:#eee}body{width:194mm;margin:0 auto;padding:12px 0}"
            "article{width:194mm;background:white}"
            "table{print-color-adjust:exact;-webkit-print-color-adjust:exact}"
            "td{line-height:1.2}"
            ".print-controls{font:14px/1.6 sans-serif;margin:0 0 16px;padding:12px;background:white}"
            "@media print{html{background:white}body{width:194mm;margin:0;padding:0}"
            ".print-controls{display:none}}"
            if browser_print else ""
        )
        controls = (
            '<aside class="print-controls"><button type="button" id="print-contract">列印／另存 PDF</button>'
            '<p>請選 A4、直向、縮放 100%，關閉頁首及頁尾並開啟背景圖形。'
            '可在目的地選擇「另存為 PDF」。此文件使用目前資料，不代表已完成簽署。</p></aside>'
            if browser_print else ""
        )
        return (
            '<!doctype html><html lang="zh-Hant"><head><meta charset="utf-8"><style>'
            "@page{size:A4 portrait;margin:8mm}html,body{margin:0;padding:0}"
            "table{width:100%;border-collapse:collapse;table-layout:fixed}"
            "tr{break-inside:avoid}"
            "td{box-sizing:border-box;white-space:pre-wrap;overflow-wrap:anywhere;vertical-align:middle}"
            + print_style
            + "</style></head><body>"
            + controls
            + ("<article><table><colgroup>" if browser_print else "<table><colgroup>")
            + columns
            + "</colgroup>"
            + "".join(rows)
            + ("</table></article></body></html>" if browser_print else "</table></body></html>")
        )
    except ContractRendererError:
        raise
    except Exception:
        raise ContractRendererError(
            "contract_pdf_renderer_source_invalid", "契約 PDF 來源無法讀取。"
        ) from None
    finally:
        if workbook is not None:
            workbook.close()


def _print_bounds(worksheet) -> tuple[int, int, int, int]:
    default_area = f"A1:{worksheet.cell(worksheet.max_row, worksheet.max_column).coordinate}"
    area = str(worksheet.print_area or default_area)
    area = area.split("!", 1)[-1].replace("'", "").replace("$", "")
    if "," in area:
        raise ContractRendererError(
            "contract_pdf_renderer_source_invalid", "契約列印範圍無效。"
        )
    return range_boundaries(area)


def _merged_cells(worksheet, min_row, max_row, min_col, max_col):
    result = {}
    for merged_range in worksheet.merged_cells.ranges:
        if (
            merged_range.min_row < min_row
            or merged_range.max_row > max_row
            or merged_range.min_col < min_col
            or merged_range.max_col > max_col
        ):
            continue
        result[(merged_range.min_row, merged_range.min_col)] = (
            merged_range.max_row - merged_range.min_row + 1,
            merged_range.max_col - merged_range.min_col + 1,
        )
    return result


def _column_width(worksheet, column: int) -> str:
    width = worksheet.column_dimensions[get_column_letter(column)].width
    points = max(Decimal("4"), Decimal(str(width or "8.43")) * Decimal("5.25"))
    return format(points.quantize(Decimal("0.01")), "f")


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return escape(value.date().isoformat() if value.time().isoformat() == "00:00:00" else value.isoformat(sep=" "))
    if isinstance(value, date):
        return escape(value.isoformat())
    return escape(str(value))


def _rgb(color) -> str | None:
    value = getattr(color, "rgb", None)
    return f"#{value[2:]}" if isinstance(value, str) and len(value) == 8 else None


def _source_fill(cell) -> str | None:
    if cell.fill.fill_type != "solid":
        return None
    color = _rgb(cell.fill.fgColor)
    # The legacy orange registration block belongs to the pale-yellow payment panel.
    return "#FFFFCC" if color == "#FFC000" else color


def _aligned_fill(worksheet, row: int, column: int, colspan: int, min_col: int, max_col: int) -> str | None:
    current = _source_fill(worksheet.cell(row, column))
    if current != "#FFFF00":
        return current

    left = _nearest_non_highlight_fill(
        worksheet, row, column - 1, -1, min_col, max_col
    )
    right = _nearest_non_highlight_fill(
        worksheet, row, column + colspan, 1, min_col, max_col
    )
    if left == right:
        return left
    return left if left is not None else right


def _nearest_non_highlight_fill(worksheet, row: int, start: int, step: int, min_col: int, max_col: int) -> str | None:
    column = start
    while min_col <= column <= max_col:
        fill = _source_fill(worksheet.cell(row, column))
        if fill != "#FFFF00":
            return fill
        column += step
    return None


def _cell_style(cell, colspan: int, background: str | None, scale: float = 1.0, *, browser_print: bool = False) -> str:
    styles = ["padding:1px 2px"]
    if cell.font.name:
        fallback = ",PMingLiU,Noto Serif CJK TC,serif" if browser_print else ""
        styles.append(f"font-family:{escape(cell.font.name)}{fallback}")
    if cell.font.sz:
        styles.append(f"font-size:{cell.font.sz * scale:.2f}pt")
    if (
        colspan > 1
        and isinstance(cell.value, str)
        and len(cell.value) > 14
        and cell.value.isascii()
    ):
        styles.append("font-size:8pt")
    if cell.font.bold:
        styles.append("font-weight:700")
    if cell.font.italic:
        styles.append("font-style:italic")
    if color := _rgb(cell.font.color):
        styles.append(f"color:{color}")
    if cell.alignment.horizontal in {"left", "center", "right", "justify"}:
        styles.append(f"text-align:{cell.alignment.horizontal}")
    if cell.alignment.vertical in {"top", "center", "bottom"}:
        vertical = "middle" if cell.alignment.vertical == "center" else cell.alignment.vertical
        styles.append(f"vertical-align:{vertical}")
    if cell.alignment.wrap_text:
        styles.extend(("white-space:pre-wrap", "overflow-wrap:anywhere"))
        if not browser_print:
            styles.append("overflow:hidden")
    if background:
        styles.append(f"background:{background}")
    for side_name in ("top", "right", "bottom", "left"):
        side = getattr(cell.border, side_name)
        if side is not None and side.style:
            width = "2px" if side.style in {"medium", "thick", "double"} else "1px"
            styles.append(
                f"border-{side_name}:{width} solid {_rgb(side.color) or '#000'}"
            )
    return ";".join(styles)


__all__ = ["BrowserContractRenderer", "render_contract_print_html"]
