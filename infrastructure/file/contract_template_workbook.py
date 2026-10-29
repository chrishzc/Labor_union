"""Safe XLSX inspection; never execute formulas, links, macros or embedded objects."""
from __future__ import annotations

from io import BytesIO
from zipfile import BadZipFile, ZipFile
from xml.etree import ElementTree

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, coordinate_to_tuple

from domains.contract_signing.template_management import TemplateManagementError
from subsystems.contract_signing.template_management import WorkbookSheetView
from infrastructure.file.contract_template_images import worksheet_images

MAX_TEMPLATE_BYTES = 20 * 1024 * 1024


def _preflight_xml(archive: ZipFile, entry, cell_count: int) -> int:
    """Bound allocation before openpyxl constructs cells or complete style tables."""
    if entry.file_size > 8 * 1024 * 1024:
        raise TemplateManagementError('validation', 'contract_template_capacity_exceeded', 'Excel 內部資料超出目前支援的預覽範圍。')
    text = archive.read(entry).decode('utf-8')
    if '<!DOCTYPE' in text:
        raise TemplateManagementError('validation', 'contract_template_xml_invalid', '模板含不支援的 XML 宣告。')
    depth, styles = 0, 0
    with archive.open(entry) as stream:
        for event, node in ElementTree.iterparse(stream, events=('start', 'end')):
            if event == 'end':
                depth -= 1
                node.clear()
                continue
            depth += 1
            local = node.tag.rsplit('}', 1)[-1]
            if depth > 64:
                raise TemplateManagementError('validation', 'contract_template_capacity_exceeded', 'Excel 內部結構過於複雜，請另存簡化的模板。')
            if entry.filename.lower() == 'xl/styles.xml' and local in {'xf', 'font', 'fill', 'border', 'numFmt'}:
                styles += 1
                if styles > 5000:
                    raise TemplateManagementError('validation', 'contract_template_capacity_exceeded', '模板樣式數量超出預覽範圍，請移除重複樣式。')
            if entry.filename.lower().startswith('xl/worksheets/') and local == 'c':
                cell_count += 1
                row, col = coordinate_to_tuple(node.attrib.get('r', 'A1'))
                if row > 1000 or col > 50 or cell_count > 50000:
                    raise TemplateManagementError('validation', 'contract_template_capacity_exceeded', '工作表超出目前支援的 1,000 列、50 欄或 50,000 格，請移除多餘的空白格式。')
    return cell_count


def _color(value, default='') -> str:
    rgb = getattr(value, 'rgb', None)
    return '#' + rgb[-6:] if isinstance(rgb, str) and len(rgb) == 8 else default


class OpenpyxlContractTemplateWorkbook:
    def inspect(self, content: bytes, worksheet_name: str | None = None) -> tuple[tuple[str, ...], WorkbookSheetView]:
        if not content or len(content) > MAX_TEMPLATE_BYTES:
            raise TemplateManagementError('validation', 'contract_template_size_invalid', '模板檔案為空或超過目前支援的 20 MB。')
        try:
            with ZipFile(BytesIO(content)) as archive:
                entries = archive.infolist()
                if len(entries) > 1500 or sum(item.file_size for item in entries) > 64 * 1024 * 1024:
                    raise TemplateManagementError('validation', 'contract_template_capacity_exceeded', '模板內容超出目前可預覽的範圍。')
                cell_count = 0
                for entry in entries:
                    name = entry.filename.lower()
                    if entry.flag_bits & 1 or any(part in name for part in ('vbaproject', 'xl/externallinks/', 'xl/embeddings/', 'xl/activex/', 'xl/charts/')):
                        raise TemplateManagementError('validation', 'contract_template_unsupported_content', '模板含加密、巨集、外部連結、嵌入物件或圖表，現行產製器無法完整保留。')
                    if name.endswith(('.xml', '.rels')):
                        cell_count = _preflight_xml(archive, entry, cell_count)
                        if name == '[content_types].xml' and any(marker in archive.read(entry).lower() for marker in (b'macroenabled', b'vbaproject', b'oleobject', b'activex')):
                            raise TemplateManagementError('validation', 'contract_template_unsupported_content', '模板宣告含有巨集或可執行嵌入物件，請另存為一般 .xlsx。')
                    if name.startswith('xl/drawings/') and name.endswith('.xml'):
                        drawing = ElementTree.fromstring(archive.read(entry))
                        if any(node.tag.rsplit('}', 1)[-1] in {'sp', 'graphicFrame', 'grpSp', 'cxnSp'} for node in drawing.iter()):
                            raise TemplateManagementError('validation', 'contract_template_drawing_unsupported', '模板含目前無法保留的圖形，請先改為 PNG 或 JPEG 圖片。')
            workbook = load_workbook(BytesIO(content), keep_links=False, data_only=False)
        except TemplateManagementError:
            raise
        except (BadZipFile, OSError, ValueError, KeyError, TypeError, EOFError, SyntaxError) as error:
            raise TemplateManagementError('validation', 'contract_template_workbook_invalid', '無法讀取 Excel，請確認檔案完整並另存為 .xlsx。') from error
        try:
            names = tuple(workbook.sheetnames)
            if worksheet_name and worksheet_name not in names:
                raise TemplateManagementError('validation', 'contract_template_worksheet_invalid', '指定工作表不存在。')
            sheet = workbook[worksheet_name] if worksheet_name else workbook.active
            if sheet is None or not workbook.worksheets:
                raise TemplateManagementError('validation', 'contract_template_worksheet_invalid', '模板沒有可使用的工作表。')
            if sheet.max_row > 1000 or sheet.max_column > 50 or sheet.max_row * sheet.max_column > 50000:
                raise TemplateManagementError('validation', 'contract_template_capacity_exceeded', '工作表超出目前支援的 1,000 列及 50 欄，請移除多餘的空白格式。')
            cells, formulas = [], []
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.data_type == 'f':
                        formulas.append(cell.coordinate)
                        value = ''
                    else:
                        value = '' if cell.value is None else str(cell.value)
                    style = {'font': cell.font.name or '新細明體', 'size': cell.font.sz or 11, 'bold': bool(cell.font.bold),
                             'italic': bool(cell.font.italic), 'color': _color(cell.font.color, '#222222'),
                             'fill': _color(cell.fill.fgColor) if cell.fill.fill_type == 'solid' else '',
                             'align': cell.alignment.horizontal or 'left', 'vertical': cell.alignment.vertical or 'center',
                             'wrap': bool(cell.alignment.wrap_text), 'borders': {}}
                    for side in ('top', 'right', 'bottom', 'left'):
                        edge = getattr(cell.border, side)
                        if edge and edge.style:
                            style['borders'][side] = ('2px' if edge.style in {'medium', 'thick', 'double'} else '1px') + ' solid ' + _color(edge.color, '#222222')
                    cells.append({'ref': cell.coordinate, 'row': cell.row, 'col': cell.column, 'value': value, 'style': style})
            merges = tuple({'start_row': value.min_row, 'end_row': value.max_row, 'start_col': value.min_col, 'end_col': value.max_col}
                           for value in sheet.merged_cells.ranges)
            return names, WorkbookSheetView(sheet.title, sheet.max_row, sheet.max_column,
                {str(col): float(sheet.column_dimensions[get_column_letter(col)].width or 8.43) for col in range(1, sheet.max_column + 1)},
                {str(row): float(sheet.row_dimensions[row].height or sheet.sheet_format.defaultRowHeight or 15) for row in range(1, sheet.max_row + 1)},
                merges, tuple(cells), tuple(formulas), worksheet_images(sheet))
        except TemplateManagementError:
            raise
        except (ValueError, OSError) as error:
            raise TemplateManagementError('validation', 'contract_template_image_unsupported', '模板圖片無法保留，請使用內嵌 PNG 或 JPEG 圖片。') from error
        finally:
            workbook.close()
