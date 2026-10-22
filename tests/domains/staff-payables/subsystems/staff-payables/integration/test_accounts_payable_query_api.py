"""Partial staff payouts expose the same remaining amount in query, audit and XLSX."""

from datetime import date, datetime, timezone
from io import BytesIO
import hashlib
from pathlib import Path

import pytest

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from api.dependencies.accounts_payable_export import (
    AccountsPayableExportApplication,
    _ReadOnlySnapshotComposition,
    get_accounts_payable_export_application,
)
from api.dependencies.admin_auth import require_admin
from api.routes.finance_reports import router
from infrastructure.mysql.accounts_payable_export_sources import (
    MySqlClientRefundExportSource,
    MySqlStaffPayableExportSource,
    _CLIENT_REFUNDS_SQL,
    _STAFF_PAYABLE_CASE_ORDER_FACTS_SQL,
    _STAFF_PAYABLE_CASE_AUDIT_SQL,
    _STAFF_PAYABLES_SQL,
)
from subsystems.access.authentication_session import AdminPrincipal
from subsystems.staff_payables.accounts_payable_export import (
    AccountsPayableExportWorkflow,
    ArchivedWorkbook,
)

URL = "/api/v1/finance-reports/accounts-payable"


class _Connection:
    def __init__(self):
        self.row = {
            "obligation_identity": "service:CASE-P:assignment:7",
            "case_no": "CASE-P", "staff_id": 7, "recipient_name": "月嫂甲",
            "identity_card": "", "amount_due_ntd": 20_000,
            "export_amount_ntd": 1500, "due_date": date(2026, 3, 15),
            "order_due_date": date(2026, 3, 15), "payout_status": "partially_paid",
            "primary_account_count": 1, "bank_code": "012", "account_no": "TEST-ACCOUNT",
            "historical_confirmation_kind": None,
        }
        self.rows = ()
        self.commits = 0
        self.rollbacks = 0

    def cursor(self):
        return self

    def __enter__(self):
        return self

    def __exit__(self, *_):
        return False

    def execute(self, statement, parameters=None):
        if statement == _STAFF_PAYABLES_SQL:
            eligible = (self.row["due_date"] == parameters[0]
                        and self.row["payout_status"] in {"payable", "partially_paid"}
                        and self.row["export_amount_ntd"] > 0)
            self.rows = (dict(self.row),) if eligible else ()
        elif statement == _STAFF_PAYABLE_CASE_AUDIT_SQL:
            self.rows = (dict(self.row),) if parameters[0] == self.row["case_no"] else ()
        elif statement == _STAFF_PAYABLE_CASE_ORDER_FACTS_SQL:
            self.rows = ({"case_no": self.row["case_no"], "staff_payment_due_date": self.row["order_due_date"]},)
        elif statement == _CLIENT_REFUNDS_SQL or "FROM scheduling_aggregates" in statement or "FROM case_staff_assignments assignment" in statement:
            self.rows = ()
        else:
            assert statement.startswith(("SET TRANSACTION", "START TRANSACTION"))

    def fetchall(self):
        return self.rows

    def fetchone(self):
        return self.rows[0] if self.rows else None

    def rollback(self):
        self.rollbacks += 1

    def commit(self):
        self.commits += 1
        raise AssertionError("Accounts payable queries must not commit accounting facts")

    def close(self):
        pass


class _Archive:
    def __init__(self):
        self.contents = []

    def save(self, year, filename, content, digest):
        self.contents.append(content)
        return ArchivedWorkbook("/test-archive/" + filename, digest)


def _client():
    connection = _Connection()
    archive = _Archive()
    application = AccountsPayableExportApplication(AccountsPayableExportWorkflow(
        MySqlStaffPayableExportSource(connection), MySqlClientRefundExportSource(connection),
        archive, lambda: _ReadOnlySnapshotComposition(connection),
        lambda: datetime(2026, 3, 1, tzinfo=timezone.utc),
    ))
    app = FastAPI()
    app.include_router(router)
    app.dependency_overrides[require_admin] = lambda: AdminPrincipal(7, "finance", "Finance", "admin")
    app.dependency_overrides[get_accounts_payable_export_application] = lambda: application
    return TestClient(app), connection, archive


def test_partial_staff_balance_is_consistent_in_query_case_audit_and_download():
    client, connection, archive = _client()
    original = dict(connection.row)
    response = client.get(URL, params={"target_month": "2026-03"})
    assert response.status_code == 200
    data = response.json()["data"]
    assert data["row_count"] == 1
    assert data["total_amount_ntd"] == 1500
    assert data["rows"][0]["amount_ntd"] == 1500
    assert data["rows"][0]["obligation_identities"] == [original["obligation_identity"]]

    response = client.get(URL + "/cases/CASE-P", params={"target_month": "2026-03"})
    assert response.status_code == 200
    item = response.json()["data"]["items"][0]
    assert item["amount_due_ntd"] == 20_000
    assert item["balance_ntd"] == 1500
    assert item["disposition"] == "selected_month"

    response = client.get(URL + "/export", params={"target_month": "2026-03"})
    assert response.status_code == 200
    assert archive.contents == [response.content]
    workbook = load_workbook(BytesIO(response.content))
    assert workbook.active["F3"].value == 1500
    assert connection.row == original
    assert connection.commits == 0
    assert connection.rollbacks == 3


def test_requery_reads_reduced_balance_and_removes_paid_items_without_changing_month_rules():
    client, connection, _ = _client()
    response = client.get(URL, params={"target_month": "2026-04"})
    assert response.status_code == 200
    assert response.json()["data"]["row_count"] == 0

    connection.row["export_amount_ntd"] = 500
    response = client.get(URL, params={"target_month": "2026-03"})
    assert response.status_code == 200
    assert response.json()["data"]["total_amount_ntd"] == 500

    connection.row.update(export_amount_ntd=0, payout_status="completed")
    response = client.get(URL, params={"target_month": "2026-03"})
    assert response.status_code == 200
    assert response.json()["data"]["rows"] == []
    response = client.get(URL + "/cases/CASE-P", params={"target_month": "2026-03"})
    assert response.status_code == 200
    item = response.json()["data"]["items"][0]
    assert item["balance_ntd"] == 0
    assert item["disposition"] == "paid_or_settled"
    assert connection.row["amount_due_ntd"] == 20_000
    assert connection.commits == 0
    assert connection.rollbacks == 4


@pytest.mark.parametrize("configured_archive", [False, True])
def test_composed_export_download_matches_persistent_archive(tmp_path, monkeypatch, configured_archive):
    from api.dependencies import accounts_payable_export as composition
    from infrastructure.archive.accounts_payable import LocalAccountsPayableArchive

    repository_root = tmp_path / "app"
    archive_root = tmp_path / "data" / "accounts-payable-archive"
    monkeypatch.setattr(composition, "_REPOSITORY_ROOT", repository_root)
    if configured_archive:
        monkeypatch.setenv("ACCOUNTS_PAYABLE_ARCHIVE_ROOT", str(archive_root))
    else:
        monkeypatch.delenv("ACCOUNTS_PAYABLE_ARCHIVE_ROOT", raising=False)
        archive_root = repository_root / "downloads" / "accounts_payable_archive"
    connection = _Connection()
    original = dict(connection.row)
    monkeypatch.setattr(composition, "get_connection", lambda: connection)
    app = FastAPI()
    app.include_router(router)
    app.dependency_overrides[require_admin] = lambda: AdminPrincipal(7, "finance", "Finance", "admin")
    client = TestClient(app)

    response = client.get(URL + "/export", params={"target_month": "2026-03"})
    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    filename = response.headers["content-disposition"].split('filename="', 1)[1].rstrip('"')
    archived_file = archive_root / "2026" / filename
    assert archived_file.read_bytes() == response.content
    assert load_workbook(BytesIO(response.content)).active["F3"].value == 1500
    if configured_archive:
        assert not repository_root.exists()

    # A fresh adapter/request reads the same durable file, independently of the workflow instance.
    reopened = LocalAccountsPayableArchive(repository_root, archive_root=archive_root)
    records = reopened.list(2026)
    assert len(records) == 1
    assert Path(records[0].absolute_path) == archived_file
    assert records[0].sha256 == hashlib.sha256(response.content).hexdigest()
    archive_response = client.get(URL + "/archive", params={"year": 2026})
    assert archive_response.status_code == 200
    assert archive_response.json()["data"]["records"][0]["sha256"] == records[0].sha256
    with pytest.raises(FileExistsError):
        reopened.save(2026, filename, b"replacement", records[0].sha256)
    assert archived_file.read_bytes() == response.content
    assert connection.row == original
    assert connection.commits == 0


def test_composed_export_does_not_download_when_archive_is_unwritable(tmp_path, monkeypatch):
    from api.dependencies import accounts_payable_export as composition

    blocked_root = tmp_path / "not-a-directory"
    blocked_root.write_bytes(b"existing file")
    monkeypatch.setenv("ACCOUNTS_PAYABLE_ARCHIVE_ROOT", str(blocked_root))
    monkeypatch.setattr(composition, "get_connection", _Connection)
    app = FastAPI()
    app.include_router(router)
    app.dependency_overrides[require_admin] = lambda: AdminPrincipal(7, "finance", "Finance", "admin")

    response = TestClient(app).get(URL + "/export", params={"target_month": "2026-03"})

    assert response.status_code == 503
    assert "accounts_payable_archive_failed" in response.text
    assert "content-disposition" not in response.headers
    assert blocked_root.read_bytes() == b"existing file"
