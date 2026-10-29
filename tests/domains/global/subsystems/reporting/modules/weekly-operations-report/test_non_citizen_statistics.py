"""Independent approved-case periods and exclusion from the existing report."""
from dataclasses import replace
from datetime import date, datetime
from io import BytesIO
from zoneinfo import ZoneInfo

import pytest
from openpyxl import load_workbook

from subsystems.reporting.weekly_operations_report_query import (
    SubsidyFacts, WeeklyCaseFact, WeeklyOperationsReportQuery,
)
from subsystems.reporting.weekly_operations_report_export import export_weekly_operations_report


def fact(case_no, day, identity="非市民", status="洽談中", category="non_citizen"):
    return WeeklyCaseFact(1, case_no, datetime.fromisoformat(day), f"Synthetic {case_no}", identity,
                          None, None, status, 10, 8, None, None, case_category=category)


class Facts:
    def __init__(self):
        self.approved = [fact(f"A115{i:03}", day, status=status) for i, (day, status) in enumerate([
            ("2026-01-01", "待補件"), ("2026-03-31", "洽談中"),
            ("2026-04-01", "訂單成立"), ("2026-06-30", "歷史訂單－服務中"),
            ("2026-07-01", "歷史訂單－帳務完成"), ("2026-12-31", "訂單取消"),
        ], 1)] + [fact("A113001", "2024-02-29"), fact("A116001", "2027-01-01")]
        self.original = [fact("115000001", "2026-01-01", "補助市民", category="subsidy"),
                         fact("115000002", "2026-01-01"),
                         fact("115000003", "2026-01-01", "自訂非市民", category="non_citizen")]
    def list_case_years(self): return [2026]
    def list_non_citizen_case_years(self): return [2027, 2026, 2024]
    def list_case_statistics_facts(self, start, end): return self.original
    def list_non_citizen_case_statistics_facts(self, start, end): return self.approved
    def list_case_facts(self, start, end): return self.original + self.approved
    def list_service_facts(self, *args): return []
    def list_subsidy_facts(self, *args): return SubsidyFacts((), ())
    def list_weekly_metrics(self, *args): return []


def query():
    return WeeklyOperationsReportQuery(Facts(), lambda: datetime(2026, 10, 29, tzinfo=ZoneInfo("Asia/Taipei")))


def test_non_citizen_calendar_periods_and_full_annual_count():
    owner = query()
    annual = owner.query_case_statistics(2026, case_category="non_citizen")
    assert annual.summary.application_count == annual.annual_application_count == 6
    assert set(annual.summary.order_status_counts.values()) == {1}
    for quarter in range(1, 5):
        selected = owner.query_case_statistics(2026, quarter=quarter, case_category="non_citizen")
        assert selected.annual_application_count == 6
    assert sum(owner.query_case_statistics(2026, quarter=q, case_category="non_citizen").summary.application_count for q in range(1,5)) == 6
    assert owner.query_case_statistics(2026, quarter=2, month=4, case_category="non_citizen").summary.application_count == 1
    assert owner.query_case_statistics(2024, month=2, case_category="non_citizen").end_date == date(2024,2,29)
    assert owner.query_case_statistics(2027, month=1, case_category="non_citizen").summary.application_count == 1
    assert owner.query_case_statistics(2025, case_category="non_citizen").summary.application_count == 0


@pytest.mark.parametrize("kwargs", [{"quarter": 1, "month": 4}, {"case_category": "all"}])
def test_invalid_classification_or_period_is_rejected(kwargs):
    with pytest.raises(ValueError, match="case_statistics_period_invalid"):
        query().query_case_statistics(2026, **kwargs)


def test_existing_statistics_weekly_totals_and_xlsx_exclude_non_citizens():
    owner = query()
    statistics = owner.query_case_statistics(2026)
    assert statistics.summary.application_count == statistics.annual_application_count == 1
    report = owner.query(date(2026,1,1), date(2026,12,31))
    assert [row.case_no for row in report.case_rows] == ["115000001"]
    assert report.summary.application_count == report.annual_totals[0].application_count == 1
    workbook = load_workbook(BytesIO(export_weekly_operations_report(report)))
    assert len(workbook.sheetnames) == 3
    cells = {str(cell.value) for sheet in workbook for row in sheet for cell in row}
    assert "Synthetic 115000001" in cells and "Synthetic 115000002" not in cells and "Synthetic A115001" not in cells
