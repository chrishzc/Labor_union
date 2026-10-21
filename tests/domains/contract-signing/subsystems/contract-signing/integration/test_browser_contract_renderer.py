from datetime import date
from pathlib import Path
from io import BytesIO

import pytest

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Border, PatternFill, Side

from infrastructure.file.browser_contract_renderer import BrowserContractRenderer, render_contract_print_html
from subsystems.contract_signing.contract_renderer import render_contract_template
from subsystems.contract_signing.template_catalog import TEMPLATE_DIRECTORY, approved_template_mapping_path, load_approved_template


def test_browser_renderer_keeps_approved_workbook_content_and_returns_pdf(tmp_path: Path):
    template = tmp_path / "contract.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "案件"
    workbook.active["A2"] = "placeholder"
    workbook.active["A3"] = "placeholder"
    workbook.active["A2"].fill = PatternFill("solid", fgColor="FFFF00")
    workbook.active["B2"].fill = PatternFill("solid", fgColor="FCDCE9")
    workbook.active["A4"] = "完整的服務平台派選服務人員說明"
    workbook.active["A4"].fill = PatternFill("solid", fgColor="FFC000")
    workbook.active["B4"].fill = PatternFill("solid", fgColor="FFC000")
    workbook.active["A5"] = "訂金"
    workbook.active["A5"].border = Border(right=Side(style="thin"))
    workbook.active["A5"].fill = PatternFill("solid", fgColor="FFFFCC")
    workbook.active["B5"].fill = PatternFill("solid", fgColor="FFFFCC")
    workbook.active.merge_cells("A1:B1")
    workbook.active.print_area = "A1:B5"
    workbook.save(template)
    mapping = tmp_path / "contract.json"
    mapping.write_text(
        '{"param_mappings":{"A2":{"db_key":"case_no","requiredness":"required"},"A3":{"db_key":"contract_date","requiredness":"required"}}}',
        encoding="utf-8",
    )
    observed: dict[str, str] = {}

    def generate(html: str) -> bytes:
        observed["html"] = html
        return b"%PDF-1.7\nbrowser\n%%EOF\n"

    result = BrowserContractRenderer(generate).render(
        template_path=template,
        mapping_path=mapping,
        facts={"case_no": "CASE-001", "contract_date": date(2026, 9, 14)},
    )

    assert "案件" in observed["html"]
    assert "CASE-001" in observed["html"]
    assert 'colspan="2"' in observed["html"]
    assert "2026-09-14" in observed["html"]
    assert "T00:00:00" not in observed["html"]
    assert "background:#FCDCE9" in observed["html"]
    assert "background:#FFFF00" not in observed["html"]
    assert "background:#FFC000" not in observed["html"]
    assert "background:#FFFFCC" in observed["html"]
    # Unbordered labels receive empty neighbouring space; table dividers survive.
    from html.parser import HTMLParser

    class Cells(HTMLParser):
        def __init__(self):
            super().__init__()
            self.cells = []

        def handle_starttag(self, tag, attrs):
            if tag == "td":
                self.cells.append(dict(attrs))

    cells = Cells()
    cells.feed(observed["html"])
    assert cells.cells[-3].get("colspan") == "2"
    assert "colspan" not in cells.cells[-2]
    assert "border-right:1px solid" in cells.cells[-2]["style"]
    assert "white-space:pre-wrap" in observed["html"]
    assert result.mime_type == "application/pdf"
    assert result.renderer_identity == "chromium-template-v3"


@pytest.mark.parametrize("template_key", ["contract_client_copy", "contract_staff_service"])
def test_browser_print_contains_full_approved_wording_and_a4_layout_without_pdf_engine(template_key):
    template = load_approved_template(template_key)
    content = render_contract_template(
        template_path=TEMPLATE_DIRECTORY / template.template_filename,
        mapping_path=approved_template_mapping_path(template_key),
        facts={"case_no": "PRINT-TEST", "client_name": '<script>alert("x")</script>'},
    )
    html = render_contract_print_html(content, title="PRINT-TEST <契約>")
    assert "@page{size:A4 portrait;margin:8mm}" in html
    assert "width:194mm" in html
    assert "print-color-adjust:exact" in html
    assert "tr{break-inside:avoid}" in html
    assert "overflow:hidden" not in html
    assert '<button type="button" id="print-contract">' in html
    assert "<title>PRINT-TEST &lt;契約&gt;</title>" in html
    assert "<script>" not in html
    workbook = load_workbook(BytesIO(content), data_only=True)
    # Every populated cell in the legal print area must survive the HTML conversion.
    from html import escape
    for row in workbook.active:
        for cell in row:
            if cell.value is not None:
                assert escape(str(cell.value)) in html, cell.coordinate
    workbook.close()
