"""Client preparation reuses the current plan without creating commitments."""
from datetime import datetime, timezone
from io import BytesIO
import json
from types import SimpleNamespace
from openpyxl import Workbook, load_workbook
import pytest
from shared_kernel.identities import CorrelationId, IdempotencyKey
from subsystems.contract_signing import client_contract_application as module
from api.dependencies.contract_external_signing import ContractExternalSigningApplication
from shared_kernel.identities import ActorContext


class Cursor:
    def __init__(self, plan, existing):
        self.plan, self.existing, self.sql = plan, existing, []
        self.order = {"case_no": "CASE-1", "contract_identity": None}
        self.signing = {"signed_received": False, "reported": False}
    def __enter__(self): return self
    def __exit__(self, *args): pass
    def execute(self, sql, args): self.sql.append(sql)
    def fetchall(self): return self.plan
    def fetchone(self):
        if "SELECT case_no,contract_identity FROM orders" in self.sql[-1]: return self.order
        if self.sql[-1].startswith("SELECT EXISTS("): return self.signing
        return self.existing


class Connection:
    def __init__(self, cursor): self.reader, self.commits, self.rollbacks = cursor, 0, 0
    def cursor(self): return self.reader
    def begin(self): pass
    def commit(self): self.commits += 1
    def rollback(self): self.rollbacks += 1
    def close(self): pass


@pytest.mark.parametrize('replayed', [True, False])
def test_preparation_uses_stable_facts_and_never_requires_prior_commitment(monkeypatch, tmp_path, replayed):
    facts = {'case_no': 'CASE-1', 'service_days': 5}
    snapshot_facts = {
        **facts,
        "__pdf_presentation_version__": module.CONTRACT_PDF_PRESENTATION_VERSION,
    }
    digest = module._sha256(module._canonical_json(snapshot_facts).encode())
    cursor = Cursor([{'matching_plan_id': 17}], {'id': 8, 'facts_snapshot_sha256': digest} if replayed else None)
    connection = Connection(cursor)
    archived = []
    monkeypatch.setattr(module, 'load_approved_template', lambda key: SimpleNamespace(template_key=key, template_filename='client.xlsx', template_sha256='a'*64, mapping_sha256='b'*64))
    monkeypatch.setattr(module, 'render_contract_template', lambda **kwargs: b'xlsx')
    monkeypatch.setattr(module, '_insert_generated_document', lambda *args, **kwargs: 9)
    app = module.ClientContractSigningApplication(lambda: connection, archive_root=tmp_path,
        now=lambda: datetime(2026,9,10,tzinfo=timezone.utc), template_facts_loader=lambda *args: facts,
        archive_document=lambda content, **kwargs: archived.append(kwargs['storage_key']) or SimpleNamespace(storage_key=kwargs['storage_key']))
    command = module.PrepareExternalClientContractCommand('CASE-1','admin:1',IdempotencyKey('client:1'),CorrelationId('test:1'))
    assert app.prepare_external_document(command) == ((8,True) if replayed else (9,False))
    assert connection.commits == 1 and connection.rollbacks == 0
    assert not any('commitment' in sql for sql in cursor.sql)
    plan_sql = next(sql for sql in cursor.sql if "FROM caregiver_matching_plans plan" in sql)
    document_sql = next(sql for sql in cursor.sql if "SELECT d.id,d.facts_snapshot_sha256" in sql)
    assert "plan.status IN ('proposed','accepted')" in plan_sql
    assert "matching_response_events" not in plan_sql
    assert "JOIN media_assets asset ON asset.id=d.media_asset_id" in document_sql
    assert "asset.mime_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'" in document_sql
    assert len(archived) == (0 if replayed else 1)
    if not replayed:
        assert archived == [
            f"CASE-1/client/external-{module._sha256(b'xlsx')[:24]}.xlsx"
        ]
        assert ":" not in archived[0]


def test_client_pdf_preparation_reuses_pdf_for_current_xlsx_source():
    class Documents:
        def prepare_external_document(self, _command):
            return 7, True

    existing = SimpleNamespace(
        document_version_id=8,
        filename="client.pdf",
        mime_type="application/pdf",
        size_bytes=1234,
    )
    facade = SimpleNamespace(
        unsigned_repository=SimpleNamespace(
            load_current_pdf_for_source=lambda case_no, source_id: existing
        ),
        unsigned_persistence=SimpleNamespace(
            prepare_and_persist=lambda _command: pytest.fail("must reuse current PDF")
        ),
    )

    result = ContractExternalSigningApplication.prepare_client_unsigned(
        facade,
        "CASE-1",
        Documents(),
        ActorContext("admin:1"),
        IdempotencyKey("client:reuse"),
        CorrelationId("test:reuse"),
    )

    assert result == {
        "document_version_id": 8,
        "filename": "client.pdf",
        "mime_type": "application/pdf",
        "size_bytes": 1234,
        "replayed": True,
    }


def test_account_change_regenerates_source_with_new_d36_and_preserves_old_workbook(monkeypatch, tmp_path):
    template = tmp_path / "client.xlsx"
    Workbook().save(template)
    mapping = tmp_path / "client.json"
    mapping.write_text(json.dumps({"param_mappings": {
        "D36": {"db_key": "client_virtual_account", "status": "approved"},
    }}), encoding="utf-8")
    monkeypatch.setattr(module, "TEMPLATE_DIRECTORY", tmp_path)
    monkeypatch.setattr(module, "approved_template_mapping_path", lambda _key: mapping)
    monkeypatch.setattr(module, "load_approved_template", lambda key: SimpleNamespace(
        template_key=key, template_filename=template.name, template_sha256="a" * 64, mapping_sha256="b" * 64,
    ))
    cursor = Cursor([{"matching_plan_id": 17}], None)
    connection = Connection(cursor)
    facts = {"case_no": "CASE-1", "client_virtual_account": "99781699114001"}
    contents = []
    def archive(content, **kwargs):
        contents.append(content)
        return SimpleNamespace(storage_key=kwargs["storage_key"])
    def insert(_connection, _command, _facts, _template, _archive, *, snapshot):
        identity = 9 + len(contents)
        cursor.existing = {"id": identity, "facts_snapshot_sha256": snapshot}
        return identity
    monkeypatch.setattr(module, "_insert_generated_document", insert)
    app = module.ClientContractSigningApplication(lambda: connection, archive_root=tmp_path,
        now=lambda: datetime(2026, 10, 15, tzinfo=timezone.utc),
        template_facts_loader=lambda *_: facts, archive_document=archive)
    def prepare(key):
        return app.prepare_external_document(module.PrepareExternalClientContractCommand(
            "CASE-1", "admin:1", IdempotencyKey(key), CorrelationId("regenerate:test")))
    assert prepare("original:1") == (10, False)
    old_snapshot = cursor.existing["facts_snapshot_sha256"]
    facts["client_virtual_account"] = "99781699114999"
    assert prepare("regenerate:2") == (11, False)
    assert cursor.existing["facts_snapshot_sha256"] != old_snapshot
    assert load_workbook(BytesIO(contents[0])).active["D36"].value == "99781699114001"
    assert load_workbook(BytesIO(contents[1])).active["D36"].value == "99781699114999"
    assert prepare("unchanged:3") == (11, True)
    assert len(contents) == 2
    assert not any(sql.lstrip().startswith(("UPDATE", "DELETE")) for sql in cursor.sql)


@pytest.mark.parametrize("signed_kind", ["contract_identity", "signed_received", "reported"])
def test_signed_client_contract_can_be_generated_without_changing_signed_facts(monkeypatch, tmp_path, signed_kind):
    cursor = Cursor([{"matching_plan_id": 17}], None)
    if signed_kind == "contract_identity":
        cursor.order["contract_identity"] = "signed-contract:existing"
    else:
        cursor.signing[signed_kind] = True
    connection = Connection(cursor)
    monkeypatch.setattr(module, 'load_approved_template', lambda key: SimpleNamespace(
        template_key=key, template_filename='client.xlsx', template_sha256='a'*64, mapping_sha256='b'*64))
    monkeypatch.setattr(module, "render_contract_template", lambda **_: b"new-xlsx")
    monkeypatch.setattr(module, "_insert_generated_document", lambda *args, **kwargs: 9)
    app = module.ClientContractSigningApplication(lambda: connection, archive_root=tmp_path,
        now=lambda: datetime(2026, 10, 15, tzinfo=timezone.utc),
        template_facts_loader=lambda *_: {"case_no": "CASE-1", "service_days": 5},
        archive_document=lambda *_args, **_kwargs: SimpleNamespace(storage_key="new-client.xlsx"))
    command = module.PrepareExternalClientContractCommand("CASE-1", "admin:1", IdempotencyKey("signed:1"), CorrelationId("signed:test"))
    assert app.prepare_external_document(command) == (9, False)
    assert connection.commits == 1 and connection.rollbacks == 0
    assert "FOR UPDATE" in cursor.sql[0]
    assert not any(sql.lstrip().startswith(("UPDATE", "DELETE")) for sql in cursor.sql)
    assert cursor.order["contract_identity"] == ("signed-contract:existing" if signed_kind == "contract_identity" else None)


def test_new_client_source_is_used_for_new_pdf_instead_of_reusing_old_pdf():
    seen = []
    facade = SimpleNamespace(
        unsigned_repository=SimpleNamespace(load_current_pdf_for_source=lambda case, source: seen.append((case, source)) or None),
        unsigned_persistence=SimpleNamespace(prepare_and_persist=lambda command: seen.append(command.source_document_version_id) or SimpleNamespace(
            document_version_id=12, filename="new-client.pdf", mime_type="application/pdf", size_bytes=100, replayed=False)),
    )
    documents = SimpleNamespace(prepare_external_document=lambda _command: (11, False))
    result = ContractExternalSigningApplication.prepare_client_unsigned(facade, "CASE-1", documents, ActorContext("admin:1"), IdempotencyKey("new:pdf"), CorrelationId("new:pdf"))
    assert seen == [("CASE-1", 11), 11]
    assert result["document_version_id"] == 12 and result["replayed"] is False


def test_missing_plan_rejection_is_a_typed_conflict():
    from subsystems.contract_signing.external_signing_contracts import ExternalSigningTypedError
    def reject(_command):
        raise ValueError("contract_external_signing_accepted_plan_required")
    with pytest.raises(ExternalSigningTypedError) as caught:
        ContractExternalSigningApplication.prepare_client_unsigned(SimpleNamespace(), "CASE-1", SimpleNamespace(prepare_external_document=reject), ActorContext("admin:1"), IdempotencyKey("signed:pdf"), CorrelationId("signed:pdf"))
    assert caught.value.category == "conflict"
    assert caught.value.code == "external_signing_accepted_plan_required"
