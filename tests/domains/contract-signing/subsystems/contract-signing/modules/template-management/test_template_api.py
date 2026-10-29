"""Verify the real multipart, typed errors, persistence and recovery HTTP boundary."""
from io import BytesIO
import json

from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
from openpyxl import Workbook
import pytest

from api.dependencies.admin_auth import require_system_admin
from api.dependencies.contract_templates import get_contract_template_management
from api.routes.contract_templates import router
from infrastructure.file.contract_template_store import FileContractTemplateStore
from infrastructure.file.contract_template_workbook import OpenpyxlContractTemplateWorkbook
from subsystems.access.authentication_session import AdminPrincipal
from subsystems.contract_signing.template_management import ContractTemplateManagement, template_fields

KEY = 'contract_client_copy'
URL = '/api/v1/contract-templates/' + KEY


@pytest.fixture
def boundary(tmp_path):
    app = FastAPI()
    application = ContractTemplateManagement(FileContractTemplateStore(tmp_path / 'catalog'), OpenpyxlContractTemplateWorkbook())
    app.include_router(router)
    app.dependency_overrides[get_contract_template_management] = lambda: application
    app.dependency_overrides[require_system_admin] = lambda: AdminPrincipal(7, 'template-admin', '模板管理員', 'system_admin')
    return TestClient(app), application, app


def payload():
    book = Workbook()
    sheet = book.active
    assignments = {}
    for index, field in enumerate(template_fields(KEY), 1):
        sheet.cell(index, 1, field.label)
        sheet.cell(index, 2, '')
        assignments[f'B{index}'] = field.key
    stream = BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue(), assignments


def test_real_upload_save_activate_readback_and_receipt(boundary):
    client, application, _ = boundary
    content, assignments = payload()
    inspected = client.post(URL + '/inspect', files={'document': ('new.xlsx', content)}, data={'assignments': '{}'})
    assert inspected.status_code == 200, inspected.text
    assert inspected.json()['data']['assignments'] == {}
    assert not application.store.root.exists()
    headers = {'Idempotency-Key': 'save-http', 'X-Correlation-ID': 'http-test'}
    saved = client.post(URL + '/versions', files={'document': ('new.xlsx', content)}, data={'assignments': json.dumps(assignments), 'expected_version': '0'}, headers=headers)
    assert saved.status_code == 200, saved.text
    result = saved.json()['data']
    version_id = result['version_id']
    assert result['catalog']['active_version_id'] == 'builtin'
    assert result['catalog']['fields'][0]['example']
    replay = client.post(URL + '/versions', files={'document': ('new.xlsx', content)}, data={'assignments': json.dumps(assignments), 'expected_version': '0'}, headers=headers)
    assert replay.json()['data']['replayed']
    assert client.get(URL + '/commands/save-http').json()['data']['version_id'] == version_id
    download = client.get(URL + '/versions/' + version_id + '/download')
    assert download.content == content
    activated = client.post(URL + '/activate', json={'version_id': version_id, 'expected_version': 1}, headers={'Idempotency-Key': 'activate-http', 'X-Correlation-ID': 'http-test'})
    assert activated.status_code == 200, activated.text
    assert client.get(URL).json()['data']['active_version_id'] == version_id
    preview = client.post(URL + '/inspect', data={'source_version_id': version_id, 'assignments': json.dumps(assignments), 'sample': 'true'})
    assert preview.status_code == 200 and preview.json()['data']['sample_values']


def test_auth_boundary_prevents_write_and_invalid_mapping_is_typed(boundary):
    client, application, app = boundary
    content, _assignments = payload()
    invalid = client.post(URL + '/inspect', files={'document': ('new.xlsx', content)}, data={'assignments': '{"B1":"SQL.column"}'})
    assert invalid.status_code == 422
    assert invalid.json()['detail']['error']['code'] == 'contract_template_field_unknown'
    def deny():
        raise HTTPException(403, 'forbidden')
    app.dependency_overrides[require_system_admin] = deny
    denied = client.post(URL + '/versions', files={'document': ('new.xlsx', content)}, data={'assignments': '{}', 'expected_version': '0'}, headers={'Idempotency-Key': 'denied', 'X-Correlation-ID': 'test'})
    assert denied.status_code == 403
    assert not application.store.root.exists()


def test_builtin_source_save_replays_with_stable_cleaned_bytes(boundary):
    client, application, _ = boundary
    current = client.get(URL + '/versions/builtin')
    assert current.status_code == 200, current.text
    assignments = current.json()['data']['assignments']
    headers = {'Idempotency-Key': 'builtin-copy', 'X-Correlation-ID': 'test'}
    data = {'source_version_id': 'builtin', 'assignments': json.dumps(assignments), 'expected_version': '0'}
    first = client.post(URL + '/versions', data=data, headers=headers)
    assert first.status_code == 200, first.text
    second = client.post(URL + '/versions', data=data, headers=headers)
    assert second.status_code == 200 and second.json()['data']['replayed']
    assert application.query(KEY).catalog_version == 1


def test_receipt_binds_immutable_admin_identity_not_display_username(boundary):
    client, application, app = boundary
    content, assignments = payload()
    headers = {'Idempotency-Key': 'identity-check', 'X-Correlation-ID': 'test'}
    data = {'assignments': json.dumps(assignments), 'expected_version': '0'}
    first = client.post(URL + '/versions', files={'document': ('new.xlsx', content)}, data=data, headers=headers)
    assert first.status_code == 200
    assert first.json()['data']['catalog']['versions'][-1]['created_by'] == 'admin:7'
    app.dependency_overrides[require_system_admin] = lambda: AdminPrincipal(7, 'renamed-user', '另一個顯示名稱', 'system_admin')
    assert client.post(URL + '/versions', files={'document': ('new.xlsx', content)}, data=data, headers=headers).json()['data']['replayed']
    app.dependency_overrides[require_system_admin] = lambda: AdminPrincipal(8, 'renamed-user', '另一管理員', 'system_admin')
    mismatch = client.post(URL + '/versions', files={'document': ('new.xlsx', content)}, data=data, headers=headers)
    assert mismatch.status_code == 409
    assert mismatch.json()['detail']['error']['code'] == 'contract_template_idempotency_mismatch'
    stored = json.loads((application.store.root / KEY / 'catalog.json').read_text(encoding='utf-8'))
    assert stored['receipts']['identity-check']['actor'] == 'admin:7'
