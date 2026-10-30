"""Protect immutable pairs, required mappings, zero-write preview and old lineage."""
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
import json

from openpyxl import Workbook, load_workbook
import pytest

from domains.contract_signing.template_management import TemplateManagementError
from infrastructure.file.contract_template_store import FileContractTemplateStore
from infrastructure.file.contract_template_workbook import OpenpyxlContractTemplateWorkbook
from subsystems.contract_signing.contract_renderer import render_contract_template
from subsystems.contract_signing.template_management import ContractTemplateManagement, template_fields
from subsystems.contract_signing.template_catalog import load_approved_template

KEY = 'contract_client_copy'


@pytest.mark.parametrize('bank_code', ['807', '812'])
def test_virtual_account_sample_reads_configured_bank_code(monkeypatch, bank_code):
    from subsystems.contract_signing.template_management import sample_value

    monkeypatch.setenv('CLIENT_DEPOSIT_BANK_CODE', bank_code)
    field = next(item for item in template_fields(KEY) if item.key == 'client_virtual_account')
    assert sample_value(field) == f'({bank_code}) 示範虛擬帳號'


def source():
    book = Workbook()
    sheet = book.active
    sheet.title = '契約'
    assignments = {}
    for index, field in enumerate(template_fields(KEY), 1):
        sheet.cell(index, 1, field.label)
        sheet.cell(index, 2, '待套值')
        assignments[f'B{index}'] = field.key
    sheet.merge_cells('B1:D1')
    sheet.column_dimensions['B'].width = 24
    sheet.row_dimensions[1].height = 32
    sheet['F1'] = '原有條文'
    stream = BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue(), assignments


@pytest.fixture
def app(tmp_path):
    return ContractTemplateManagement(FileContractTemplateStore(tmp_path / 'catalog'), OpenpyxlContractTemplateWorkbook())


def test_preview_is_zero_write_and_preserves_layout(app, tmp_path):
    content, assignments = source()
    preview = app.inspect(KEY, content, '契約.xlsx', None, assignments, sample=True)
    assert preview.sheet.column_widths['2'] == 24
    assert preview.sheet.row_heights['1'] == 32
    assert {'start_row': 1, 'end_row': 1, 'start_col': 2, 'end_col': 4} in preview.sheet.merges
    assert preview.sample_values['B1']
    assert not (tmp_path / 'catalog').exists()
    assert app.inspect(KEY, content, '新模板.xlsx', None, {}).assignments == {}


@pytest.mark.parametrize('key', ['contract_client_copy', 'contract_staff_service'])
def test_actual_builtin_templates_can_be_opened_without_creating_state(app, key, tmp_path):
    result = app.read_version(key, 'builtin')
    assert result.sheet.rows >= 90
    assert result.sheet.merges
    assert not (tmp_path / 'catalog').exists()


def test_save_reopen_activate_keeps_baseline_and_exact_bytes(app):
    baseline = app.store.load_active(KEY)
    content, assignments = source()
    result = app.save(KEY, content, '新版.xlsx', '契約', assignments, 0, 'admin', 'save-1')
    assert result.catalog.active_version_id == 'builtin'
    assert result.catalog.catalog_version == 1
    reopened = FileContractTemplateStore(app.store.root)
    assert reopened.content(KEY, result.version_id)[0] == content
    app.activate(KEY, result.version_id, 1, 'admin', 'activate-1')
    active = reopened.load_active(KEY)
    assert active.version_id == result.version_id
    assert active.template_path.read_bytes() == content
    historical = reopened.load_exact(KEY, baseline.template_sha256, baseline.mapping_sha256)
    assert historical.version_id == 'builtin'
    assert historical.template_path.read_bytes() == baseline.template_path.read_bytes()
    # The selected worksheet, moved cell, style and static wording survive real production filling.
    output = render_contract_template(template_path=active.template_path, mapping_path=active.mapping_path,
                                      facts={key: '示範' for key in assignments.values()})
    book = load_workbook(BytesIO(output))
    assert book.active['B1'].value == '示範'
    assert book.active['F1'].value == '原有條文'
    assert book.active.column_dimensions['B'].width == 24
    book.close()


def test_same_key_replay_and_payload_mismatch(app):
    content, assignments = source()
    one = app.save(KEY, content, '新版.xlsx', '契約', assignments, 0, 'admin', 'save-1')
    again = app.save(KEY, content, '新版.xlsx', '契約', assignments, 0, 'admin', 'save-1')
    assert again.replayed and again.version_id == one.version_id
    with pytest.raises(TemplateManagementError, match='不同內容') as caught:
        app.save(KEY, content, '另一檔.xlsx', '契約', assignments, 0, 'admin', 'save-1')
    assert caught.value.category == 'idempotency_mismatch'
    assert len(app.query(KEY).versions) == 2


def test_catalog_cas_only_one_competing_save_commits(app):
    content, assignments = source()
    def save(key):
        try:
            return app.save(KEY, content, '新版.xlsx', '契約', assignments, 0, 'admin', key).version_id
        except TemplateManagementError as error:
            return error.category
    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(pool.map(save, ['first', 'second']))
    assert results.count('conflict') == 1
    assert app.query(KEY).catalog_version == 1


@pytest.mark.parametrize('change,code', [('unknown', 'contract_template_field_unknown'), ('covered', 'contract_template_merge_anchor_required'), ('missing', 'contract_template_required_fields_missing')])
def test_invalid_mappings_cannot_be_saved(app, change, code):
    content, assignments = source()
    if change == 'unknown':
        assignments['B1'] = 'raw_database_column'
    elif change == 'covered':
        assignments['C1'] = assignments.pop('B1')
    else:
        assignments = {}
    with pytest.raises(TemplateManagementError) as caught:
        app.save(KEY, content, '新版.xlsx', None, assignments, 0, 'admin', 'save')
    assert caught.value.code == code
    assert not app.store.root.exists()


def test_corrupt_workbook_and_unmapped_formula_are_closed(app):
    with pytest.raises(TemplateManagementError) as caught:
        app.inspect(KEY, b'not zip', 'bad.xlsx', None, {})
    assert caught.value.code == 'contract_template_workbook_invalid'
    content, assignments = source()
    book = load_workbook(BytesIO(content))
    book.active['F2'] = '=1+1'
    stream = BytesIO()
    book.save(stream)
    with pytest.raises(TemplateManagementError) as caught:
        app.save(KEY, stream.getvalue(), 'formula.xlsx', None, assignments, 0, 'admin', 'save')
    assert caught.value.code == 'contract_template_formula_unmapped'


def test_tampered_version_never_falls_back_to_builtin(app):
    content, assignments = source()
    saved = app.save(KEY, content, '新版.xlsx', None, assignments, 0, 'admin', 'save')
    app.activate(KEY, saved.version_id, 1, 'admin', 'activate')
    (app.store.root / KEY / saved.version_id / 'source.xlsx').write_bytes(b'tampered')
    with pytest.raises(TemplateManagementError) as caught:
        app.store.load_active(KEY)
    assert caught.value.code == 'contract_template_version_unavailable'


def test_script_text_is_data_and_unsupported_embedded_content_is_rejected(app):
    from zipfile import ZipFile, ZIP_DEFLATED
    content, _assignments = source()
    book = load_workbook(BytesIO(content))
    book.active['F1'] = '<img src=x onerror="alert(1)">'
    stream = BytesIO()
    book.save(stream)
    preview = app.inspect(KEY, stream.getvalue(), 'text.xlsx', None, {})
    assert next(cell['value'] for cell in preview.sheet.cells if cell['ref'] == 'F1') == '<img src=x onerror="alert(1)">'
    injected = BytesIO()
    with ZipFile(BytesIO(content)) as source_zip, ZipFile(injected, 'w', compression=ZIP_DEFLATED) as target:
        for entry in source_zip.infolist():
            target.writestr(entry.filename, source_zip.read(entry.filename))
        target.writestr('xl/vbaProject.bin', b'not executable in this workflow')
    with pytest.raises(TemplateManagementError) as caught:
        app.inspect(KEY, injected.getvalue(), 'macros-renamed.xlsx', None, {})
    assert caught.value.code == 'contract_template_unsupported_content'


def test_capacity_is_checked_before_openpyxl_allocates_cells(app, monkeypatch):
    from zipfile import ZipFile, ZIP_DEFLATED
    content, _assignments = source()
    oversized = BytesIO()
    with ZipFile(BytesIO(content)) as original, ZipFile(oversized, 'w', compression=ZIP_DEFLATED) as target:
        for entry in original.infolist():
            payload = original.read(entry.filename)
            if entry.filename == 'xl/worksheets/sheet1.xml':
                payload = b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData><row r="1048576"><c r="A1048576"><v>1</v></c></row></sheetData></worksheet>'
            target.writestr(entry.filename, payload)
    def unexpected_load(*_args, **_kwargs):
        raise AssertionError('unsafe workbook allocation reached')
    monkeypatch.setattr('infrastructure.file.contract_template_workbook.load_workbook', unexpected_load)
    with pytest.raises(TemplateManagementError) as caught:
        app.inspect(KEY, oversized.getvalue(), 'oversized.xlsx', None, {})
    assert caught.value.code == 'contract_template_capacity_exceeded'


def test_partial_first_save_does_not_publish_or_break_builtin(app, monkeypatch):
    content, assignments = source()
    baseline = app.store.load_active(KEY)
    write_version = app.store._put_version
    def interrupted(key, metadata, _content, _mapping, *, storage_id=None):
        directory = app.store._version_path(key, storage_id or metadata.version_id)
        directory.mkdir()
        (directory / 'source.xlsx').write_bytes(b'incomplete')
        raise OSError('simulated interrupted write')
    monkeypatch.setattr(app.store, '_put_version', interrupted)
    with pytest.raises(OSError, match='interrupted'):
        app.save(KEY, content, '新版.xlsx', None, assignments, 0, 'admin:7', 'same-command')
    assert app.query(KEY).catalog_version == 0
    assert app.query(KEY).active_version_id == 'builtin'
    assert app.read_version(KEY, 'builtin').sheet.rows == 185
    assert app.store.load_active(KEY).template_sha256 == baseline.template_sha256
    monkeypatch.setattr(app.store, '_put_version', write_version)
    saved = app.save(KEY, content, '新版.xlsx', None, assignments, 0, 'admin:7', 'same-command')
    assert saved.catalog.catalog_version == 1
    assert len(saved.catalog.versions) == 2
