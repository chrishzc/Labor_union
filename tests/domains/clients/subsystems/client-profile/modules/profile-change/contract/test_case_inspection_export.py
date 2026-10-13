"""Case export grain, scope, safe XLSX and authenticated transport contracts."""

from dataclasses import fields
from datetime import date
from decimal import Decimal
from io import BytesIO

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from api.dependencies.admin_auth import require_registry_reader
from api.dependencies.client_registry import get_client_registry_case_inspection_export_application
from api.routes.client_registry import router
from infrastructure.mysql.client_case_inspection_repository import MySqlClientCaseInspectionRepository
from subsystems.access.authentication_session import AdminPrincipal
from subsystems.client_profile.case_inspection_export import (
    AssignmentRecord, CaseInspectionData, ClientRegistryCaseInspectionExportApplication,
    LedgerRecord, ScheduleRecord,
)
from subsystems.client_profile.order_accounting_export import OrderAccountingExportQuery, OrderAccountingExportRow, XLSX_MEDIA_TYPE


def overview(case_no="00115000001"):
    values = {field.name: None for field in fields(OrderAccountingExportRow)}
    values.update(case_no=case_no, name="=1+1", is_twins=False, imported_virtual_accounts=(),
                  received_total_ntd=300, planned_start_date=date(2026, 10, 1))
    return OrderAccountingExportRow(**values)


class Repository:
    selection = None

    def query_case_inspection(self, selection):
        self.selection = selection
        case_no = overview().case_no
        return CaseInspectionData(
            overview=(overview(),),
            assignments=(AssignmentRecord(case_no, "=1+1", 7, 2, "月嫂測試", 1, 1,
                None, None, None, None, Decimal("8.5"), Decimal("8"), 350, 0, "replaced", None),),
            schedules=tuple(ScheduleRecord(case_no, "=1+1", i, 7, 2, "月嫂測試", 1,
                date(2026, 10, i), True, False, True) for i in (1, 2, 3)),
            ledger=(LedgerRecord(case_no, "=1+1", 9223372036854775807, None, "receipt", 300, date(2026, 10, 1), None, None),
                    LedgerRecord(case_no, "=1+1", 2, None, "reversal", -300, date(2026, 10, 2), None, 9223372036854775807)),
        )


def test_workbook_separates_detail_grains_keeps_reversals_and_db_text():
    repo = Repository()
    content = ClientRegistryCaseInspectionExportApplication(repo).export(OrderAccountingExportQuery(query=" test "))
    assert repo.selection.query == "test"
    wb = load_workbook(BytesIO(content))
    assert wb.sheetnames == ["案件總覽", "案件保存資料", "派案明細", "每日排班", "客戶義務", "客戶收付流水", "客戶款項分配", "月嫂應付", "月嫂付款分配"]
    assert wb["案件總覽"].max_row == 2
    assert wb["派案明細"].max_row == 2
    assert wb["每日排班"].max_row == 4
    ledger = wb["客戶收付流水"]
    assert ledger.max_row == 3
    assert ledger["F3"].value == -300
    assert ledger["F3"].data_type == "n"
    assert ledger["C2"].value == "9223372036854775807"
    assert ledger["I3"].value == "9223372036854775807"
    assert wb["案件總覽"]["B2"].value == "00115000001"
    assert wb["案件總覽"]["B2"].data_type == "s"
    assert wb["案件總覽"]["C2"].value == "=1+1"
    assert ledger["G2"].value.date() == date(2026, 10, 1)
    assert wb["派案明細"]["L2"].value == 8.5
    for sheet in wb:
        assert sheet.freeze_panes
        assert sheet.auto_filter.ref
        assert not any(cell.data_type == "f" for row in sheet for cell in row)


def test_empty_selection_produces_nine_header_only_sheets_without_detail_queries(monkeypatch):
    class NoConnection:
        def cursor(self):
            raise AssertionError("empty selection must not query detail tables")
    monkeypatch.setattr("infrastructure.mysql.client_case_inspection_repository.MySqlClientRegistryQueryRepository.query_order_accounting_rows", lambda *_: ())
    content = ClientRegistryCaseInspectionExportApplication(MySqlClientCaseInspectionRepository(NoConnection())).export(OrderAccountingExportQuery())
    wb = load_workbook(BytesIO(content))
    assert len(wb.sheetnames) == 9
    assert all(sheet.max_row == 1 for sheet in wb)


def test_repository_reads_fixed_case_scoped_joins_without_page_limit_or_status_exclusions(monkeypatch):
    statements = []
    class Cursor:
        def __enter__(self): return self
        def __exit__(self, *_): return False
        def execute(self, sql, parameters): statements.append((sql, parameters))
        def fetchall(self): return ()
    class Connection:
        def cursor(self): return Cursor()
    cases = tuple(overview(f"{i:011}") for i in range(105))
    monkeypatch.setattr("infrastructure.mysql.client_case_inspection_repository.MySqlClientRegistryQueryRepository.query_order_accounting_rows", lambda *_: cases)
    result = MySqlClientCaseInspectionRepository(Connection()).query_case_inspection(OrderAccountingExportQuery())
    assert result.overview == cases
    assert len(statements) == 8
    for sql, parameters in statements:
        assert sql.startswith("SELECT ")
        assert "JOIN clients c ON c.id=o.client_id" in sql
        assert parameters == tuple(row.case_no for row in cases)
        assert " IN (" in sql and sql.count("%s") == 105
        assert "LIMIT" not in sql and "SELECT *" not in sql and "c.case_no" not in sql
        assert all(field not in sql for field in ("source_account", "counterparty_account", "identity_card", "raw_payload", "c.phone", "c.address"))
    assert "t.id=a.settlement_detail_id" in statements[-1][0]
    assert "a.allocated_amount" in statements[-1][0] and "p.amount" not in statements[-1][0]


def test_route_requires_login_then_returns_filtered_xlsx_without_pagination(monkeypatch):
    monkeypatch.setenv("ACCESS_CONTROL_PROFILE", "production")
    repo = Repository()
    app = FastAPI()
    app.include_router(router)
    app.dependency_overrides[get_client_registry_case_inspection_export_application] = lambda: ClientRegistryCaseInspectionExportApplication(repo)
    client = TestClient(app)
    path = "/api/v1/admin/registries/clients/export/case-inspection"
    assert client.get(path).status_code == 401
    app.dependency_overrides[require_registry_reader] = lambda: AdminPrincipal(9, "reader", "Test Reader", "system_admin")
    response = client.get(path, params={"query": " test ", "requires_cooking": "false", "limit": 1, "offset": 100})
    assert response.status_code == 200
    assert response.headers["content-type"] == XLSX_MEDIA_TYPE
    assert response.headers["content-disposition"] == 'attachment; filename="client-case-inspection.xlsx"'
    assert response.headers["cache-control"] == "no-store"
    assert repo.selection == OrderAccountingExportQuery(query="test", requires_cooking=False)
    assert len(load_workbook(BytesIO(response.content)).sheetnames) == 9
    invalid = client.get(path, params={"column_filters": "not-json"})
    assert invalid.status_code == 422
