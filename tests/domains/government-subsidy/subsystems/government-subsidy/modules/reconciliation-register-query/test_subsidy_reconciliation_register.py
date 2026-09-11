from datetime import date
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook

from subsystems.government_subsidy import reconciliation_register_query as register


class FakeCursor:
    def __init__(self, rows):
        self.rows = rows
        self.executed = []

    def __enter__(self):
        return self

    def __exit__(self, exc_type, exc, traceback):
        return False

    def execute(self, sql, params):
        self.executed.append((" ".join(sql.split()), params))

    def fetchall(self):
        return self.rows


class FakeConnection:
    def __init__(self, rows):
        self.cursor_instance = FakeCursor(rows)
        self.closed = False

    def cursor(self):
        return self.cursor_instance

    def close(self):
        self.closed = True


def _order_row(**overrides):
    row = {
        "case_no": "115000002", "identity_status": "一般市民",
        "actual_start_date": date(2026, 3, 1), "actual_end_date": date(2026, 3, 20),
        "service_days": 20, "service_hours_per_day": Decimal("8"),
        "order_status": "訂單完成",
        "employer_name": "王小明", "employer_address": "台北市中正區",
        "staff_name": "月嫂甲", "survey_details": {"身分證字號": "A123456789"},
    }
    row.update(overrides)
    return row


def test_quarterly_register_includes_established_orders_without_claim_batch():
    connection = FakeConnection([
        _order_row(),
        _order_row(
            case_no="115000001", identity_status="補助市民",
            actual_start_date=date(2025, 12, 1), actual_end_date=date(2026, 1, 1),
            employer_name="陳小美", employer_address="新北市板橋區",
            staff_name="月嫂乙", survey_details='{"身分證字號": "B223456789"}',
            service_days=15, service_hours_per_day=Decimal("8"),
        ),
        _order_row(case_no="115000099", actual_end_date=date(2026, 4, 1)),
    ])
    result = register.build_quarterly_subsidy_register(2026, 1, lambda: connection)

    assert [row["\u5e02\u5e9c\u8a02\u55ae\u865f\u78bc"] for row in result["general_citizen_rows"]] == ["115000002"]
    assert [row["\u5e02\u5e9c\u8a02\u55ae\u865f\u78bc"] for row in result["subsidized_citizen_rows"]] == ["115000001"]
    assert result["general_citizen_rows"][0]["補助天數"] == Decimal("5.00")
    assert result["general_citizen_rows"][0]["補助款金額"] == Decimal("12000")
    assert result["general_citizen_rows"][0]["服務人員"] == "月嫂甲"
    assert result["general_citizen_rows"][0]["身分證字號"] == "A123456789"
    assert result["subsidized_citizen_rows"][0]["\u7c3d\u9818"] == ""
    assert connection.closed is True
    assert "INSERT" not in connection.cursor_instance.executed[0][0].upper()
    assert "c.identity_status" in connection.cursor_instance.executed[0][0]
    assert "clients.identity_status" not in connection.cursor_instance.executed[0][0]
    assert "subsidy_claim_batches" not in connection.cursor_instance.executed[0][0]
    assert "o.status IN (%s, %s, %s)" in connection.cursor_instance.executed[0][0]
    assert "COALESCE(o.actual_end_date, o.end_date)" in connection.cursor_instance.executed[0][0]
    assert connection.cursor_instance.executed[0][1] == (
        "訂單成立", "服務中", "訂單完成", "一般市民", "補助市民",
        date(2026, 1, 1), date(2026, 4, 1),
    )

    workbook = load_workbook(BytesIO(result["xlsx_bytes"]))
    worksheet = workbook["\u5206\u5b63\u6838\u92b7"]
    assert worksheet.cell(row=1, column=1).value == "\u4e00\u822c\u5e02\u6c11"
    assert worksheet.cell(row=2, column=15).value == "\u7c3d\u9818"
    assert worksheet.cell(row=3, column=7).number_format == "0.00"
    assert worksheet.cell(row=5, column=1).value == "\u88dc\u52a9\u5e02\u6c11"


def test_annual_summary_uses_established_orders_and_repairs_legacy_key():
    legacy_key = "\u8eab\u5206\u8b49\u5b57\u865f".encode("utf-8").decode("latin1")
    connection = FakeConnection([
        _order_row(
            case_no="115000010", actual_start_date="2026-07-01",
            actual_end_date="2026-07-20", employer_name="林太太",
            employer_address="桃園市", staff_name="月嫂丙",
            survey_details={legacy_key: "C123456789"},
        ),
    ])

    result = register.build_annual_subsidy_summary(2026, lambda: connection)
    row = result["general_citizen_rows"][0]
    assert row["\u8eab\u5206\u8b49\u5b57\u865f"] == "C123456789"
    assert result["subsidized_citizen_rows"] == []
    assert connection.cursor_instance.executed[0][1][-2:] == (
        date(2026, 1, 1), date(2027, 1, 1),
    )

    workbook = load_workbook(BytesIO(result["xlsx_bytes"]))
    worksheet = workbook["\u5e74\u5ea6\u7e3d\u8868"]
    values = [cell.value for cell in worksheet["A"]]
    assert "\u88dc\u52a9\u5e02\u6c11" not in values
    assert worksheet.max_column == 10


def test_operations_report_annual_rows_include_historical_established_orders_without_claim_batch():
    connection = FakeConnection([
        _order_row(
            case_no="114000003",
            actual_start_date=date(2025, 12, 20),
            actual_end_date=date(2026, 1, 8),
        ),
    ])

    result = register.build_operations_report_annual_subsidy_rows(
        2026,
        lambda: connection,
    )

    row = result["general_citizen_rows"][0]
    assert row["市府訂單號碼"] == "114000003"
    assert row["核銷月份"] == "第一季"
    assert row["核銷狀態"] == "結案"
    assert result["subsidized_citizen_rows"] == []
    sql, params = connection.cursor_instance.executed[0]
    assert "subsidy_claim_batches" not in sql
    assert "current_revision" not in sql
    assert "o.status IN (%s, %s, %s, %s, %s, %s, %s)" in sql
    assert params == (
        "訂單成立",
        "服務中",
        "訂單完成",
        "歷史訂單－未服務",
        "歷史訂單－服務中",
        "歷史訂單－服務完成",
        "歷史訂單－帳務完成",
        "一般市民",
        "補助市民",
        date(2026, 1, 1),
        date(2027, 1, 1),
    )


def test_operations_report_reconciliation_period_uses_service_end_quarter():
    assert register._operations_reconciliation_period(date(2026, 3, 31)) == (2026, "第一季")
    assert register._operations_reconciliation_period(date(2026, 4, 1)) == (2026, "第二季")
    assert register._operations_reconciliation_period(date(2026, 7, 1)) == (2026, "第三季")
    assert register._operations_reconciliation_period(date(2026, 10, 1)) == (2026, "第四季")
    assert register._operations_reconciliation_period(date(2027, 1, 1)) == (2027, "第一季")


def test_invalid_quarter_is_rejected_before_database_access(monkeypatch):
    try:
        register.build_quarterly_subsidy_register(
            2026,
            5,
            lambda: (_ for _ in ()).throw(AssertionError("database must not be opened")),
        )
    except ValueError as exc:
        assert "quarter" in str(exc)
    else:
        raise AssertionError("invalid quarter must be rejected")


def test_register_caps_subsidy_hours_at_case_total_service_hours():
    row = register._to_register_row({
        "case_no": "115000011", "identity_status": "一般市民",
        "actual_start_date": date(2026, 1, 1), "actual_end_date": date(2026, 1, 3),
        "service_days": 3, "service_hours_per_day": 9,
        "employer_name": "王小明", "employer_address": "台北市", "staff_name": "月嫂甲",
        "survey_details": {},
    })

    assert row["補助時數"] == Decimal("27")
    assert row["補助款金額"] == Decimal("8100")


def test_combined_subsidy_register_has_both_quarterly_and_annual_sheets():
    connection = FakeConnection([
        _order_row(case_no="115000001", employer_name="陳小姐", staff_name="王月嫂"),
    ])

    result = register.build_combined_subsidy_register(2026, 1, lambda: connection)
    assert result["application_year"] == 2026
    assert result["quarter"] == 1
    workbook = load_workbook(BytesIO(result["xlsx_bytes"]))
    assert workbook.sheetnames == ["季核銷", "年度總表"]
    quarter_ws = workbook["季核銷"]
    annual_ws = workbook["年度總表"]
    assert quarter_ws.max_row >= 3
    assert annual_ws.max_row >= 3
    assert "hc_case_no" not in result["quarterly_general_rows"][0]
    assert "115000001" in [str(c.value) for row in quarter_ws.iter_rows() for c in row if c.value]
