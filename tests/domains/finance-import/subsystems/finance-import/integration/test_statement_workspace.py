"""Observable archive validation, download safety and display-only warning rules."""
from io import BytesIO

from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pytest

from api.dependencies.admin_auth import require_admin
from api.dependencies.finance_import import get_accounting_workspace, get_statement_archive_application
from api.routes.finance_statement_workspace import router
from domains.finance_import.statement_archive import accounting_warning, bank_import_type, validate_manual_statement
from subsystems.finance_import.statement_archive import StatementExportRow


@pytest.mark.parametrize("fields", [
    {"交易日": "2026-10-29", "支出": "9970", "備註": "HC115001退款"},
    {"交易日": "2026-10-29", "存入": 10000, "銷帳編號": "000001"},
    {"交易日": "2026-10-29", "支出": 10, "銷帳編號": "000001", "備註": "=1+1"},
])
def test_manual_minimum_fields_preserve_missing_values(fields):
    result = validate_manual_statement(fields)
    assert result["帳號"] is None and result["計息日"] is None
    assert result["備註"] == fields.get("備註")
    assert result["銷帳編號"] == fields.get("銷帳編號")


@pytest.mark.parametrize("update", [
    {"交易日": None}, {"交易日": "2026-02-30"}, {"支出": None},
    {"支出": "0"}, {"支出": -1}, {"支出": "NaN"}, {"支出": "1.5"},
    {"支出": True}, {"存入": "1"}, {"備註": " "}, {"未知欄位": "x"},
])
def test_invalid_manual_input_is_rejected(update):
    with pytest.raises(ValueError):
        validate_manual_statement({"交易日": "2026-10-29", "支出": "9970", "備註": "HC115001退款", **update})


@pytest.mark.parametrize("difference,warning", [(0, None), (30, None), (99, None), (100, None), (101, "under_paid"), (-1, "over_paid"), (-30, "over_paid")])
def test_payment_warning_threshold_is_display_only(difference, warning):
    assert accounting_warning("payable", 10000, 10000 - difference, True) == warning
    assert accounting_warning("payable", 10000, 0, False) is None


def test_receipt_surplus_does_not_use_payment_threshold():
    assert accounting_warning("receivable", 10000, 10001, True) == "over_received"
    assert accounting_warning("receivable", 10000, 9000, True) is None


def test_opposing_differences_never_cancel_individual_warnings():
    from subsystems.client_finance.case_accounting_query import ClientAccountingItem
    from subsystems.finance_import.accounting_workspace import AccountingWorkspace
    class State:
        def case_client(self, _): return "Synthetic"
        def dismissed(self, _, *, for_update=False): return {}
    items = tuple(ClientAccountingItem(f"refund:{i}", "payable", "客戶退款", "Synthetic", 10000, actual, True, "1:0:0", None) for i, actual in enumerate((9850, 10150)))
    result = AccountingWorkspace(State(), lambda *_, **__: (items, ()), lambda *_, **__: ((), ())).query("case")
    assert result.payable_ntd - result.paid_ntd == 0
    assert result.pending_warning_count == 2
    assert {row.kind for row in result.warnings} == {"under_paid", "over_paid"}


def test_client_accounting_uses_remaining_and_effective_allocations_after_date_only_event():
    from infrastructure.mysql.client_payments_query_repository import query_case_accounting
    class Cursor:
        def __enter__(self): return self
        def __exit__(self, *_): pass
        def execute(self, sql, _):
            if "FROM client_obligations o" in sql:
                self.rows = [{"obligation_identity": "deposit:case", "obligation_type": "deposit", "direction": "receivable_from_client", "due_date": None, "status": "settled", "amount_due_ntd": 0, "projection_version": 2, "after_amount_ntd": 0, "account_version": 3, "client_name": "Synthetic", "confirmation_kind": None, "historical_event": None}]
            elif "FROM client_ledger_entries l" in sql:
                self.rows = [{"id": 1, "entry_type": "receipt", "occurred_on": "2026-10-29", "reversal_of_entry_id": None, "reversed_type": None, "obligation_identity": "deposit:case", "amount_ntd": 10000, "format_id": "sinopac"}]
            else:
                self.rows = []
        def fetchall(self): return self.rows
    class Connection:
        def cursor(self): return Cursor()
    items, transactions = query_case_accounting(Connection(), "case")
    assert items[0].expected_ntd == items[0].actual_ntd == 10000
    assert len(transactions) == 1


@pytest.mark.parametrize("bank,direction,result", [
    ("sinopac", "incoming", "receivable"), ("sinopac", "outgoing", "payable"),
    ("taishin", "incoming", "government_subsidy"), ("taishin", "outgoing", "subsidy_return"),
])
def test_bank_types_preserve_owner_classification_boundary(bank, direction, result):
    assert bank_import_type(bank, direction) == result


def test_full_export_keeps_both_bank_columns_and_text_without_formulas():
    class Archive:
        def export_rows(self):
            return (
                StatementExportRow("bank:1", "sinopac", "reconciled", ("115000001",), (("帳號", "000001"), ("備註", "=1+1"), ("更正註記", "Y"))),
                StatementExportRow("bank:2", "taishin", "reconciled", ("115000002",), (("台新原欄", "000002"), ("備註", "+cmd"))),
                StatementExportRow("manual:1", "manual", "manual", (), (("交易日", "2026-01-01"), ("支出", 10), ("備註", "備註"))),
            )
    app = FastAPI(); app.include_router(router)
    app.dependency_overrides[require_admin] = lambda: object()
    app.dependency_overrides[get_statement_archive_application] = Archive
    app.dependency_overrides[get_accounting_workspace] = lambda: object()
    with TestClient(app) as client:
        response = client.get("/statements/export")
    assert response.status_code == 200 and response.headers["X-Statement-Count"] == "3"
    workbook = load_workbook(BytesIO(response.content), data_only=False)
    assert workbook.sheetnames == ["人工新增", "永豐", "台新"]
    sheet = workbook["永豐"]
    assert sheet["A2"].value == "000001" and sheet["M2"].data_type == "s"
    assert sheet["M2"].value == "=1+1" and sheet.auto_filter.ref
    assert workbook["台新"]["A2"].value == "000002"
    assert workbook["人工新增"]["A2"].value is None


def test_export_failure_never_returns_partial_workbook_and_requires_authorization():
    class Archive:
        def export_rows(self):
            raise ValueError("sensitive raw data must not appear")
    app = FastAPI(); app.include_router(router)
    app.dependency_overrides[require_admin] = lambda: object()
    app.dependency_overrides[get_statement_archive_application] = Archive
    with TestClient(app) as client:
        response = client.get("/statements/export")
        assert response.status_code == 422
        assert "sensitive raw" not in response.text
        def deny():
            raise HTTPException(403, "forbidden")
        app.dependency_overrides[require_admin] = deny
        app.dependency_overrides[get_accounting_workspace] = lambda: object()
        for path in ("/statements/export", "/cases/115000001/accounting", "/accounting/warnings"):
            assert client.get(path).status_code == 403


def test_empty_export_explicitly_identifies_zero_rows_without_fabricating_facts():
    class Archive:
        def export_rows(self): return ()
    app = FastAPI(); app.include_router(router)
    app.dependency_overrides[require_admin] = lambda: object()
    app.dependency_overrides[get_statement_archive_application] = Archive
    with TestClient(app) as client:
        response = client.get("/statements/export")
    assert response.status_code == 200 and response.headers["X-Statement-Count"] == "0"
    workbook = load_workbook(BytesIO(response.content))
    assert workbook.sheetnames == ["無成功明細（0筆）"]
    assert workbook.active.max_row == 1
