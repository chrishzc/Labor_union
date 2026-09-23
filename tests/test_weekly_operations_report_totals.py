"""年度累計、星期一歸月、月末小計與同 candidate XLSX 的 focused contract。"""
from dataclasses import asdict, replace
from datetime import date, datetime, timedelta
from io import BytesIO
from zoneinfo import ZoneInfo

import pytest
from openpyxl import load_workbook
from pydantic import ValidationError

from api.schemas.operations_reports import WeeklyReportCaseTotalsView
from subsystems.reporting.weekly_operations_report_export import export_weekly_operations_report
from subsystems.reporting.weekly_operations_report_query import (
    MISSING_ORDER_STATUS, ORDER_STATUS_LABELS,
    SubsidyFacts, WeeklyCaseFact, WeeklyOperationsReportQuery,
)
from subsystems.reporting.weekly_report_metrics_service import WeeklyReportMetric, week_starts_between


def _case(day: str, serial: int, **changes) -> WeeklyCaseFact:
    fact = WeeklyCaseFact(
        client_id=serial, case_no=f"CASE-{serial}", created_at=datetime.fromisoformat(day),
        applicant_name=f"測試案件{serial}", identity_status="一般市民", reject_reason=None,
        district="東區", order_status="訂單成立", service_days=20, service_hours_per_day=8,
        planned_start_date=date(2026, 10, 1), planned_end_date=date(2026, 10, 28),
        seq_num=serial,
    )
    return replace(fact, **changes)


class _Facts:
    def __init__(self, cases=(), counts=None, default=0):
        self.cases = list(cases)
        self.counts = counts or {}
        self.default = default
        self.calls = []

    def list_case_facts(self, start, end):
        self.calls.append(("cases", start, end))
        return [row for row in self.cases if row.created_at is None or start <= row.created_at.date() <= end]

    def list_weekly_metrics(self, start, end):
        self.calls.append(("metrics", start, end))
        return [WeeklyReportMetric(
            week, week + timedelta(days=6),
            *self.counts.get(week.isoformat(), (self.default, self.default)),
        ) for week in week_starts_between(start, end)]

    def list_service_facts(self, start, end):
        self.calls.append(("service", start, end))
        return []

    def list_subsidy_facts(self, start, end):
        self.calls.append(("subsidy", start, end))
        return SubsidyFacts((), ())


def _query(facts, start="2026-08-24", end="2026-09-13"):
    return WeeklyOperationsReportQuery(
        facts, lambda: datetime(2026, 9, 23, 11, tzinfo=ZoneInfo("Asia/Taipei")),
    ).query(date.fromisoformat(start), date.fromisoformat(end))


def _cross_month_facts():
    return _Facts([
        _case("2026-01-06", 1),
        _case("2026-08-24", 2),
        _case("2026-08-31", 3),
        _case("2026-09-01", 4),
        _case("2026-09-06", 5),
        _case("2026-09-07", 6, order_status="洽談中"),
        _case("2026-09-14", 7),
    ], {
        "2026-01-05": (7, 9),
        "2026-08-24": (5, 10),
        "2026-08-31": (12, 34),
        "2026-09-07": (3, 0),
        "2026-09-14": (99, 99),
    })


def _workbook(report):
    return load_workbook(BytesIO(export_weekly_operations_report(report)), data_only=True)


def test_cross_month_whole_week_belongs_to_monday_month_once():
    report = _query(_cross_month_facts())
    august, september = report.monthly_subtotals
    assert [(r.year, r.month, r.application_count) for r in report.monthly_subtotals] == [
        (2026, 8, 4), (2026, 9, 1),
    ]
    assert (august.promotion_count, august.inquiry_count) == (17, 44)
    assert (september.promotion_count, september.inquiry_count) == (3, 0)
    assert len(report.case_rows) == report.summary.application_count == 5
    assert len(report.weekly_metrics) == 3
    assert [r.case_no for r in report.case_rows] == [f"CASE-{i}" for i in range(2, 7)]
    assert next(r for r in report.case_rows if r.case_no == "CASE-4").month_label == ""
    assert next(r for r in report.case_rows if r.case_no == "CASE-6").month_label == "9月"


def test_annual_totals_read_year_to_date_not_only_selected_rows():
    facts = _cross_month_facts()
    report = _query(facts)
    annual, = report.annual_totals
    assert annual.month is None and annual.year == 2026
    assert (annual.start_date, annual.end_date) == (date(2026, 1, 5), date(2026, 9, 13))
    assert annual.application_count == 6
    assert (annual.promotion_count, annual.inquiry_count) == (27, 53)
    assert report.summary.application_count == 5
    assert ("cases", date(2026, 1, 5), date(2026, 9, 13)) in facts.calls
    assert ("metrics", date(2026, 1, 5), date(2026, 9, 13)) in facts.calls
    assert [call for call in facts.calls if call[0] == "service"] == [("service", date(2026, 8, 24), date(2026, 9, 13))]
    assert [call for call in facts.calls if call[0] == "subsidy"] == [("subsidy", date(2026, 8, 24), date(2026, 9, 13))]


def test_partial_date_selection_keeps_case_filter_but_never_splits_week_metrics():
    report = _query(_cross_month_facts(), "2026-09-01", "2026-09-03")
    subtotal, = report.monthly_subtotals
    assert subtotal.month == 8
    assert (subtotal.start_date, subtotal.end_date) == (date(2026, 9, 1), date(2026, 9, 3))
    assert subtotal.application_count == 1
    assert (subtotal.promotion_count, subtotal.inquiry_count) == (12, 34)
    assert report.case_rows[0].month_label == "8月"
    assert report.case_rows[0].application_date_roc == "115/09/01"
    assert report.annual_totals[0].application_count == 4


@pytest.mark.parametrize("counts, expected", [
    ({"2026-08-24": (0, 0), "2026-08-31": (0, 0)}, (0, 0)),
    ({"2026-08-24": (5, 1), "2026-08-31": (None, 4)}, (None, 5)),
    ({"2026-08-24": (5, None), "2026-08-31": (3, 0)}, (8, None)),
    ({}, (None, None)),
])
def test_zero_and_missing_metrics_are_not_conflated(counts, expected):
    report = _query(_Facts(counts=counts, default=None), "2026-08-24", "2026-09-06")
    subtotal, = report.monthly_subtotals
    assert (subtotal.promotion_count, subtotal.inquiry_count) == expected
    assert report.annual_totals[0].promotion_count is None


def test_missing_week_entry_does_not_produce_partial_total():
    class MissingWeek(_Facts):
        def list_weekly_metrics(self, start, end):
            return [metric for metric in super().list_weekly_metrics(start, end)
                    if metric.week_start_date != date(2026, 8, 31)]

    report = _query(MissingWeek(), "2026-08-24", "2026-09-06")
    assert report.monthly_subtotals[0].promotion_count is None
    assert report.monthly_subtotals[0].inquiry_count is None


def test_cross_year_uses_monday_year_without_duplication():
    facts = _Facts([
        _case("2025-01-06", 1), _case("2025-12-29", 2),
        _case("2026-01-04", 3), _case("2026-01-05", 4),
    ], {"2025-12-29": (12, 34), "2026-01-05": (5, 6)})
    report = _query(facts, "2025-12-29", "2026-01-11")
    assert [(t.year, t.month, t.application_count) for t in report.monthly_subtotals] == [(2025, 12, 2), (2026, 1, 1)]
    assert [(t.year, t.application_count) for t in report.annual_totals] == [(2025, 3), (2026, 1)]
    assert report.annual_totals[0].end_date == date(2026, 1, 4)
    assert report.annual_totals[1].start_date == date(2026, 1, 5)
    assert [t.promotion_count for t in report.annual_totals] == [12, 5]
    sheet = _workbook(report).worksheets[0]
    assert [sheet.cell(i, 1).value for i in (4, 5)] == ["114年度累計", "115年度累計"]
    labels = [sheet.cell(i, 1).value for i in range(6, sheet.max_row + 1)]
    assert labels == [2, 3, "114年12月小計", 4, "115年1月小計"]


def test_early_january_week_remains_in_previous_december():
    report = _query(_Facts([_case("2026-01-01", 1)]), "2026-01-01", "2026-01-04")
    assert [(t.year, t.month) for t in report.monthly_subtotals] == [(2025, 12)]
    assert [(t.year, t.application_count) for t in report.annual_totals] == [(2025, 1)]


def test_undated_rows_are_preserved_but_not_assigned_to_year_or_month():
    report = _query(_Facts([_case("2026-09-01", 1, created_at=None)]))
    assert len(report.case_rows) == 1
    assert "application_date_missing" in report.case_rows[0].data_quality_codes
    assert report.annual_totals[0].application_count == 0
    assert all(t.application_count == 0 for t in report.monthly_subtotals)
    sheet = _workbook(report).worksheets[0]
    assert sheet.cell(sheet.max_row, 4).value == "日期未登錄"


def test_rejection_statistics_reuse_existing_case_flags():
    facts = _Facts([
        _case("2026-09-01", 1, reject_reason="測試資格不符", order_status="訂單取消"),
        _case("2026-09-02", 2, identity_status="補助市民", reject_reason="測試資格不符", order_status="訂單取消"),
        _case("2026-09-03", 3, identity_status=None, reject_reason="測試待分流", order_status=None),
        _case("2026-09-04", 4, order_status="審核不符合"),
    ])
    report = _query(facts, "2026-08-31", "2026-09-06")
    total, = report.monthly_subtotals
    assert (total.general_ineligible_count, total.subsidized_ineligible_count) == (1, 1)
    assert total.review_rejected_count == sum(row.review_rejected for row in report.case_rows) == 4
    assert total.cancelled_count == 0
    assert total.rejection_unpartitioned_count == report.summary.rejection_unpartitioned_count == 3


def test_unknown_rejection_partition_is_not_invented():
    report = _query(_Facts([_case("2026-09-01", 1, identity_status=None, reject_reason="測試待分流")]))
    total = report.monthly_subtotals[0]
    assert total.general_ineligible_count is None
    assert total.subsidized_ineligible_count is None
    assert total.review_rejected_count == 1


def test_xlsx_annual_first_monthly_after_last_week_and_merges_stop_before_subtotals():
    report = _query(_cross_month_facts())
    workbook = _workbook(report)
    assert workbook.sheetnames == ["週報案件受理總表", "補助案件統計表", "每周服務中說明"]
    sheet = workbook.worksheets[0]
    assert sheet.cell(4, 1).value == "115年度累計"
    assert list(sheet.values)[3][5:9] == (27, 53, 6, 6)
    assert [sheet.cell(i, 1).value for i in range(5, sheet.max_row + 1)] == [2, 3, 4, 5, "115年8月小計", 6, "115年9月小計"]
    assert list(sheet.values)[8][5:9] == (17, 44, 4, 4)
    assert list(sheet.values)[10][5:9] == (3, 0, 1, 1)
    assert sheet.cell(6, 4).value == "2026-08-31 ~ 2026-09-06"
    assert sheet.cell(6, 6).value == 12
    assert {"D6:D8", "F6:F8", "G6:G8", "A9:E9"}.issubset({str(r) for r in sheet.merged_cells.ranges})
    assert sheet.freeze_panes == "A5"
    assert sheet.cell(4, 1).comment and "2026-01-05" in sheet.cell(4, 1).comment.text


def test_empty_cases_keep_week_rows_and_monthly_zero():
    report = _query(_Facts(), "2026-08-31", "2026-09-06")
    sheet = _workbook(report).worksheets[0]
    assert sheet.cell(4, 1).value == "115年度累計"
    assert sheet.cell(5, 4).value == "2026-08-31 ~ 2026-09-06"
    assert sheet.cell(5, 6).value == 0
    assert sheet.cell(6, 1).value == "115年8月小計"
    assert sheet.cell(6, 8).value == 0


def test_xlsx_incomplete_metrics_label_is_not_zero():
    report = _query(_Facts(default=None), "2026-08-31", "2026-09-06")
    sheet = _workbook(report).worksheets[0]
    assert sheet.cell(4, 6).value == "未完整登錄"
    assert sheet.cell(5, 6).value == "未登錄"
    assert sheet.cell(6, 6).value == "未完整登錄"


def test_totals_are_strict_serializable_api_values():
    report = _query(_cross_month_facts())
    for total in (*report.annual_totals, *report.monthly_subtotals):
        view = WeeklyReportCaseTotalsView.model_validate(asdict(total))
        assert view.model_dump() == asdict(total)
        assert view.model_dump(mode="json")["start_date"] == total.start_date.isoformat()
        with pytest.raises(ValidationError):
            WeeklyReportCaseTotalsView.model_validate({**asdict(total), "unknown_total": 1})


def test_inverted_period_does_not_read_sources():
    facts = _Facts()
    with pytest.raises(ValueError, match="operations_report_date_range_invalid"):
        _query(facts, "2026-09-02", "2026-09-01")
    assert facts.calls == []


def test_every_lifecycle_status_is_counted_separately_in_both_totals_and_xlsx():
    assert ORDER_STATUS_LABELS == (
        "待補件", "洽談中", "訂單成立", "服務中", "訂單完成", "訂單取消",
        "歷史訂單－未服務", "歷史訂單－服務中", "歷史訂單－服務完成", "歷史訂單－帳務完成",
    )
    cases = [_case("2026-01-06", 1)]
    cases += [_case("2026-09-01" if i < 6 else "2026-09-07", i + 2, order_status=status)
              for i, status in enumerate(ORDER_STATUS_LABELS)]
    cases.append(_case("2026-09-07", 12, order_status=None))
    report = _query(_Facts(cases), "2026-08-31", "2026-09-13")
    annual, = report.annual_totals
    august, september = report.monthly_subtotals
    assert (annual.application_count, august.application_count, september.application_count) == (12, 6, 5)
    for i, status in enumerate(ORDER_STATUS_LABELS):
        assert annual.order_status_counts[status] == (2 if status == "訂單成立" else 1)
        assert august.order_status_counts[status] == int(i < 6)
        assert september.order_status_counts[status] == int(i >= 6)
    assert annual.order_status_counts[MISSING_ORDER_STATUS] == 1
    for total in (annual, august, september):
        assert sum(total.order_status_counts.values()) == total.application_count

    sheet = _workbook(report).worksheets[0]
    assert tuple(sheet.cell(3, col).value for col in range(13, 24)) == (*ORDER_STATUS_LABELS, MISSING_ORDER_STATUS)
    assert sheet.cell(2, 24).value == "審核不符合"
    assert sheet.cell(2, 25).value == "服務天數"
    assert sheet.max_column == 31
    total_rows = [row for row in range(4, sheet.max_row + 1)
                  if isinstance(sheet.cell(row, 1).value, str)]
    for row_number, total in zip(total_rows, (annual, august, september), strict=True):
        for col, status in enumerate((*ORDER_STATUS_LABELS, MISSING_ORDER_STATUS), 13):
            assert sheet.cell(row_number, col).value == total.order_status_counts[status]
    for row_number in range(5, sheet.max_row + 1):
        if isinstance(sheet.cell(row_number, 1).value, int):
            assert sum(sheet.cell(row_number, col).value or 0 for col in range(13, 24)) == 1


def test_cancelled_and_rejected_are_independent_not_mutually_exclusive():
    report = _query(_Facts([_case("2026-09-01", 1, order_status="訂單取消", reject_reason="測試不符")]))
    for total in (*report.annual_totals, *report.monthly_subtotals):
        if total.application_count:
            assert total.order_status_counts["訂單取消"] == 1
            assert sum(total.order_status_counts.values()) == 1
            assert total.review_rejected_count == 1


def test_missing_empty_and_additional_source_status_are_not_lost():
    report = _query(_Facts([
        _case("2026-09-01", 1, order_status=None),
        _case("2026-09-01", 2, order_status=""),
        _case("2026-09-01", 3, order_status="審核不符合"),
    ]), "2026-08-31", "2026-09-06")
    total, = report.monthly_subtotals
    assert total.order_status_counts[MISSING_ORDER_STATUS] == 2
    assert total.order_status_counts["審核不符合"] == 1
    assert sum(total.order_status_counts.values()) == 3
    assert total.review_rejected_count == 1
    sheet = _workbook(report).worksheets[0]
    assert sheet.cell(3, 24).value == "來源狀態：審核不符合"
    assert sheet.cell(2, 25).value == "審核不符合"


def test_empty_report_still_has_every_status_with_numeric_zero():
    report = _query(_Facts(), "2026-08-31", "2026-09-06")
    for total in (*report.annual_totals, *report.monthly_subtotals):
        assert tuple(total.order_status_counts) == (*ORDER_STATUS_LABELS, MISSING_ORDER_STATUS)
        assert all(value == 0 for value in total.order_status_counts.values())


@pytest.mark.parametrize("bad_count", [-1, 1.5, True, "1"])
def test_status_counts_api_rejects_noninteger_or_negative_values(bad_count):
    total = asdict(_query(_cross_month_facts()).annual_totals[0])
    total["order_status_counts"]["待補件"] = bad_count
    with pytest.raises(ValidationError):
        WeeklyReportCaseTotalsView.model_validate(total)
