"""
File: test_finance_import_disposable_mysql_e2e.py
Description: 以 disposable MySQL 驗證 Finance Import 資料完整性、投影與重試停損。
"""

from __future__ import annotations

from argparse import Namespace
import os
from pathlib import Path

import pandas as pd
import pytest

from scripts.bootstrap_disposable_mysql_schema import bootstrap
from scripts.imports.finance_statement_normalizer import normalize_workbook


DATABASE = os.getenv("LABOR_UNION_TEST_MYSQL_DATABASE")
PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEIDENTIFIED_TAISHIN_FIXTURE = (
    PROJECT_ROOT / "tests/fixtures/finance_import/taishin_deidentified_minimal.xlsx"
)
pytestmark = pytest.mark.skipif(
    not DATABASE,
    reason="requires an explicitly configured disposable lu_test_* MySQL database",
)


def test_sinopac_memo_case_isolation_on_fresh_mysql(tmp_path):
    """Requires a prebootstrapped, empty disposable schema; never resets data.

    Shared counterparties and identical amounts deliberately make account-only
    matching unsafe. Every settlement must stay inside the memo's case.
    """
    import json

    from api.dependencies.finance_import import build_finance_import_application
    from infrastructure.mysql.finance_import_owning_domain_composite import (
        MySqlFinanceImportOwningDomainComposite,
    )
    from infrastructure.mysql.mysql_adapter import get_connection
    from scripts.imports.finance_formats.sinopac import SINOPAC_HEADERS
    from shared_kernel.identities import (
        ActorContext, CorrelationId, ExpectedVersion, IdempotencyKey,
    )
    from subsystems.client_finance.virtual_account_resolution import build_client_virtual_account
    from subsystems.finance_import.import_workflow import FinanceImportApplyRequest
    from subsystems.finance_import.ingestion import ingest_finance_workbook

    assert DATABASE and DATABASE.startswith("lu_test_")
    connection = get_connection()
    try:
        with connection.cursor() as cursor:
            cursor.execute("SELECT DATABASE() AS database_name")
            assert cursor.fetchone()["database_name"] == DATABASE
            for table in ("orders", "clients", "staff", "finance_import_rows"):
                cursor.execute(f"SELECT COUNT(*) AS count FROM {table}")
                assert cursor.fetchone()["count"] == 0, "requires fresh empty fixture database"
        _seed_sinopac_memo_cases(connection)
        before = _sinopac_memo_readback(connection)
    finally:
        connection.close()

    client_account, staff_account = "9" * 16, "8" * 16
    scenarios = (
        ("refund", "HC115701退款", client_account, 1200, 0),
        ("staff-hc", "HC115702服務人員", staff_account, 12000, 0),
        ("staff-a", "A115703服務人員", staff_account, 14000, 0),
        ("passbook-only", "", staff_account, 12000, 0),
        ("multiple-codes", "HC115702 HC115704", staff_account, 12000, 0),
        ("unknown-case", "HC115799退款", client_account, 1200, 0),
        ("wrong-client", "HC115706退款", client_account, 1200, 0),
        ("seven-digits", "HC1157020", staff_account, 12000, 0),
        ("receipt-source", "", "", 0, 2000),
        ("wrong-case-amount", "HC115701服務人員", staff_account, 12000, 0),
    )
    rows = []
    for index, (label, memo, account, debit, credit) in enumerate(scenarios, 1):
        row = dict.fromkeys(SINOPAC_HEADERS)
        row.update({
            "帳號": "10000001", "交易日": f"2026-10-29 09:00:{index:02d}",
            "計息日": "2026-10-29", "入帳日": "2026-10-29", "摘要": "轉帳",
            "幣別": "TWD", "支出": debit, "存入": credit, "餘額": 100000 - index,
            "交易參考編號": account, "備註": memo, "票據號碼": label,
            "存摺備註": "HC115704服務人員",
        })
        if credit:
            row["銷帳編號"] = build_client_virtual_account("115000705")
            row["更正註記"] = build_client_virtual_account("115000704")
            row["交易參考編號"] = build_client_virtual_account("115000706")
        rows.append(row)
    workbook = tmp_path / "sinopac-synthetic-memo-cases.xlsx"
    pd.DataFrame(rows, columns=SINOPAC_HEADERS).to_excel(workbook, index=False)
    actor = ActorContext("memo-mysql-test")
    ingest_arguments = (str(workbook), IdempotencyKey("memo-mysql-ingest"), actor)
    receipt = ingest_finance_workbook(
        *ingest_arguments, connection_factory=get_connection, normalizer=normalize_workbook,
    )
    assert receipt.source_row_count == receipt.canonical_created_count == 10
    assert receipt.source_warning_count == 0

    connection = get_connection()
    try:
        with connection.cursor() as cursor:
            cursor.execute(
                "SELECT id,memo,debit,credit,bank_references,raw_payload "
                "FROM finance_import_rows ORDER BY id"
            )
            immutable_before = cursor.fetchall()
        application = build_finance_import_application(
            connection, MySqlFinanceImportOwningDomainComposite(connection),
        )
        preview = application.preview_batch(receipt.batch_identity, CorrelationId("memo-preview"))
        assert preview.apply_allowed
        assert len(preview.rows) == 10
        assert {row.target_identities for row in preview.dispatchable_rows} == {
            ("refund:115000701",), ("staff:115000702",),
            ("staff:115000703",), ("deposit:115000705",),
        }
        request = FinanceImportApplyRequest(
            receipt.batch_identity, ExpectedVersion(preview.batch_version), preview.fingerprint,
            IdempotencyKey("memo-mysql-apply"), actor, "synthetic memo case isolation",
            CorrelationId("memo-apply"),
        )
        from infrastructure.mysql.finance_statement_archive_repository import MySqlStatementArchiveRepository
        from subsystems.finance_import.import_workflow import FinanceImportWorkflowError
        # Ingestion and Preview alone are never archive success evidence.
        archive = MySqlStatementArchiveRepository(connection)
        assert archive.export_rows() == ()
        snapshot = _workspace_financial_snapshot(connection)

        class FailAfterFormalWrite(MySqlFinanceImportOwningDomainComposite):
            def post(self, candidate):
                super().post(candidate)
                raise RuntimeError("synthetic rollback after child write")

        failing_application = build_finance_import_application(connection, FailAfterFormalWrite(connection))
        with pytest.raises(FinanceImportWorkflowError):
            failing_application.apply_batch(request)
        assert archive.export_rows() == ()
        assert _workspace_financial_snapshot(connection) == snapshot
        applied = application.apply_batch(request)
        assert applied.reconciled_count == 4
        after = _sinopac_memo_readback(connection)
        assert after["client_entries"] == [
            {"case_no": "115000701", "entry_type": "refund", "amount_ntd": 1200},
            {"case_no": "115000705", "entry_type": "receipt", "amount_ntd": 2000},
        ]
        assert after["client_allocations"] == [
            {"obligation_identity": "deposit:115000705", "amount_ntd": 2000},
            {"obligation_identity": "refund:115000701", "amount_ntd": 1200},
        ]
        assert after["staff_allocations"] == [
            {"case_no": "115000702", "obligation_identity": "staff:115000702", "allocated_amount_ntd": 12000},
            {"case_no": "115000703", "obligation_identity": "staff:115000703", "allocated_amount_ntd": 14000},
        ]
        assert after["staff_projections"] == [
            {"obligation_identity": "staff:115000702", "net_paid_ntd": 12000, "balance_ntd": 0, "status": "completed"},
            {"obligation_identity": "staff:115000703", "net_paid_ntd": 14000, "balance_ntd": 0, "status": "completed"},
        ]
        target_client = {"refund:115000701", "deposit:115000705"}
        decoys = lambda state: [
            row for row in state["client_obligations"]
            if row["obligation_identity"] not in target_client
        ]
        assert decoys(after) == decoys(before)
        assert after["staff_obligations"] == before["staff_obligations"]
        for row in after["client_obligations"]:
            if row["obligation_identity"] in target_client:
                assert row["amount_due_ntd"] == 0 and row["status"] == "settled"
        assert after["case_count"] == before["case_count"] == 6
        with connection.cursor() as cursor:
            cursor.execute(
                "SELECT bank_row.id FROM finance_import_rows bank_row "
                "LEFT JOIN client_ledger_entries client_entry ON client_entry.finance_import_row_id=bank_row.id "
                "LEFT JOIN staff_payout_events staff_event ON staff_event.finance_import_row_id=bank_row.id "
                "WHERE client_entry.id IS NULL AND staff_event.id IS NULL ORDER BY bank_row.id"
            )
            unposted_ids = [row["id"] for row in cursor.fetchall()]
            assert unposted_ids == [immutable_before[i]["id"] for i in (3, 4, 5, 6, 7, 9)]
            cursor.execute(
                "SELECT id,memo,debit,credit,bank_references,raw_payload "
                "FROM finance_import_rows ORDER BY id"
            )
            assert cursor.fetchall() == immutable_before
        connection.rollback()  # End read snapshot before checking committed replay.
        assert application.apply_batch(request) == applied
    finally:
        connection.close()

    assert ingest_finance_workbook(
        *ingest_arguments, connection_factory=get_connection, normalizer=normalize_workbook,
    ) == receipt
    connection = get_connection()
    try:
        assert _sinopac_memo_readback(connection) == after
        with connection.cursor() as cursor:
            cursor.execute("SELECT COUNT(*) AS count FROM finance_import_rows")
            assert cursor.fetchone()["count"] == 10
            cursor.execute("SELECT COUNT(*) AS count FROM staff_payout_events")
            assert cursor.fetchone()["count"] == 2
    finally:
        connection.close()
    report = {
        "status": "PASS", "database": DATABASE,
        "new_cases": [f"11500070{i}" for i in range(1, 7)],
        "source_rows": 10, "settled_rows": 4, "unposted_rows": 6,
        "checks": {name: "PASS" for name in (
            "refund_memo_case", "staff_hc_memo_case", "staff_a_memo_case",
            "missing_memo_no_passbook_fallback", "multiple_codes_pending",
            "unknown_case_pending", "wrong_counterparty_pending", "seven_digits_pending",
            "wrong_case_amount_pending", "receipt_reads_cancellation_column",
            "decoy_obligations_unchanged", "immutable_bank_source", "idempotent_replay",
        )},
        "before": before, "after": after,
        "external_provider_effects": 0, "production_writes": 0,
        "batch_identity": receipt.batch_identity,
    }
    (tmp_path / "memo-case-receipt.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2, default=str), encoding="utf-8",
    )


def _seed_sinopac_memo_cases(connection):
    """Create synthetic root facts only; all settlement goes through real owners."""
    import hashlib

    with connection.cursor() as cursor:
        cursor.execute("INSERT INTO clients(case_no,name) VALUES ('115000701','Memo Test Client')")
        client_id = cursor.lastrowid
        cursor.execute("INSERT INTO clients(case_no,name) VALUES ('115000706','Memo Test Other Client')")
        other_client_id = cursor.lastrowid
        cursor.execute("INSERT INTO staff(name) VALUES ('Memo Test Staff')")
        staff_id = cursor.lastrowid
        cursor.execute(
            "INSERT INTO staff_bank_accounts(staff_id,bank_code,account_no,is_primary) VALUES (%s,'999',%s,1)",
            (staff_id, "8" * 16),
        )
        for index in range(1, 7):
            case_no = f"11500070{index}"
            if index in (2, 3, 4, 5):
                cursor.execute(
                    "INSERT INTO clients(case_no,name) VALUES (%s,'Memo Test Case Record')",
                    (case_no,),
                )
            cursor.execute(
                "INSERT INTO orders(case_no,client_id,status) VALUES (%s,%s,'訂單完成')",
                (case_no, other_client_id if index == 6 else client_id),
            )
            if index in (1, 2, 4, 5, 6):
                kind, direction, amount = (
                    ("deposit", "receivable_from_client", 2000) if index == 5
                    else ("refund", "payable_to_client", 1200)
                )
                identity = f"{kind}:{case_no}"
                cursor.execute(
                    "INSERT INTO client_obligation_events(obligation_identity,case_no,obligation_type,direction,"
                    "event_type,before_amount_ntd,after_amount_ntd,after_due_date,source_event_identity,"
                    "expected_account_version,idempotency_key,actor,reason) "
                    "VALUES (%s,%s,%s,%s,'established',0,%s,'2026-10-29',%s,0,%s,'memo-test','synthetic root')",
                    (identity, case_no, kind, direction, amount, identity, identity),
                )
                event_id = cursor.lastrowid
                cursor.execute(
                    "INSERT INTO client_obligations(obligation_identity,case_no,obligation_type,direction,"
                    "amount_due_ntd,due_date,status,current_event_id,projection_version) "
                    "VALUES (%s,%s,%s,%s,%s,'2026-10-29','open',%s,0)",
                    (identity, case_no, kind, direction, amount, event_id),
                )
                if kind == "refund":
                    cursor.execute(
                        "INSERT INTO client_refund_recipient_snapshots "
                        "(refund_obligation_identity,case_no,bank_code,bank_account,source_kind) "
                        "VALUES (%s,%s,'999',%s,'test-fixture')",
                        (identity, case_no, ("7" if index == 6 else "9") * 16),
                    )
            if index <= 4:
                identity = f"staff:{case_no}"
                amount = {1: 9000, 2: 12000, 3: 14000, 4: 12000}[index]
                cursor.execute(
                    "INSERT INTO case_staff_assignments(case_no,staff_id,assignment_sequence,status) "
                    "VALUES (%s,%s,1,'completed')", (case_no, staff_id),
                )
                assignment_id = cursor.lastrowid
                cursor.execute("INSERT INTO payroll_case_accounts(case_no,aggregate_version) VALUES (%s,1)", (case_no,))
                cursor.execute(
                    "INSERT INTO staff_obligation_events(obligation_identity,assignment_id,case_no,staff_id,"
                    "obligation_kind,direction,event_type,before_amount_ntd,after_amount_ntd,due_date,"
                    "payroll_fingerprint,expected_payroll_version,resulting_payroll_version,idempotency_key,actor,reason) "
                    "VALUES (%s,%s,%s,%s,'service_pay','payable_to_staff','established',0,%s,'2026-10-29',%s,0,1,%s,'memo-test','synthetic root')",
                    (identity, assignment_id, case_no, staff_id, amount, hashlib.sha256(identity.encode()).hexdigest(), identity),
                )
                event_id = cursor.lastrowid
                cursor.execute(
                    "INSERT INTO staff_obligations(obligation_identity,assignment_id,case_no,staff_id,"
                    "obligation_kind,direction,amount_due_ntd,due_date,status,current_event_id,payroll_version) "
                    "VALUES (%s,%s,%s,%s,'service_pay','payable_to_staff',%s,'2026-10-29','open',%s,1)",
                    (identity, assignment_id, case_no, staff_id, amount, event_id),
                )
    connection.commit()


def _sinopac_memo_readback(connection):
    queries = {
        "client_obligations": "SELECT obligation_identity,case_no,amount_due_ntd,status,current_event_id,projection_version FROM client_obligations ORDER BY obligation_identity",
        "staff_obligations": "SELECT obligation_identity,case_no,amount_due_ntd,status,current_event_id,payroll_version FROM staff_obligations ORDER BY obligation_identity",
        "client_entries": "SELECT case_no,entry_type,amount_ntd FROM client_ledger_entries ORDER BY case_no",
        "client_allocations": "SELECT obligation_identity,amount_ntd FROM client_ledger_obligation_allocations ORDER BY obligation_identity",
        "staff_allocations": "SELECT obligation.case_no,link.obligation_identity,link.allocated_amount_ntd FROM staff_payout_obligation_links link JOIN staff_obligations obligation ON obligation.obligation_identity=link.obligation_identity ORDER BY obligation.case_no",
        "staff_projections": "SELECT obligation_identity,net_paid_ntd,balance_ntd,status FROM staff_payable_projections ORDER BY obligation_identity",
    }
    with connection.cursor() as cursor:
        result = {}
        for key, sql in queries.items():
            cursor.execute(sql)
            result[key] = list(cursor.fetchall())
        cursor.execute("SELECT COUNT(*) AS count FROM orders")
        result["case_count"] = cursor.fetchone()["count"]
    return result


def test_statement_archive_and_accounting_workspace_mysql(tmp_path):
    """Real ingestion/settlement plus authenticated archive and case readback."""
    from io import BytesIO
    import json

    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from openpyxl import load_workbook
    from api.dependencies.admin_auth import require_admin
    from api.routes.finance_import import router
    from infrastructure.mysql.mysql_adapter import get_connection
    from subsystems.access.authentication_session import AdminPrincipal

    test_sinopac_memo_case_isolation_on_fresh_mysql(tmp_path)
    app = FastAPI(); app.include_router(router)
    app.dependency_overrides[require_admin] = lambda: AdminPrincipal(None, "workspace-test-a", "Synthetic operator", "system_admin")

    def post(client, path, body, key):
        response = client.post("/api/v1/finance-import" + path, json=body, headers={"Idempotency-Key": key})
        assert response.status_code == 200, response.text
        return response.json()["data"]

    def export(client):
        response = client.get("/api/v1/finance-import/statements/export")
        assert response.status_code == 200, response.text
        return int(response.headers["X-Statement-Count"]), load_workbook(BytesIO(response.content), data_only=False)

    with TestClient(app) as client:
        count, workbook = export(client)
        assert count == 4 and workbook["永豐"].max_row == 5
        for case_no, expected, paid in (("115000701", 10200, 1200), ("115000702", 13200, 12000), ("115000703", 14000, 14000), ("115000704", 13200, 0)):
            response = client.get(f"/api/v1/finance-import/cases/{case_no}/accounting")
            assert response.status_code == 200, response.text
            data = response.json()["data"]
            assert data["case_no"] == case_no and data["payable_ntd"] == expected and data["paid_ntd"] == paid
            assert all(case_no in item["identity"] for item in data["client_items"] + data["staff_items"])
        assert client.get("/api/v1/finance-import/cases/not-a-case/accounting").status_code == 404
        connection = get_connection()
        try:
            unchanged = _sinopac_memo_readback(connection)
        finally:
            connection.close()
        body = {"fields": {"交易日": "2026-10-29", "支出": 1200, "備註": "HC115701退款"}}
        manual = post(client, "/statements/manual", body, "workspace-manual-1")
        assert manual["saved_count"] == 1
        assert post(client, "/statements/manual", body, "workspace-manual-1") == manual
        mismatch = client.post("/api/v1/finance-import/statements/manual", json={"fields": {**body["fields"], "支出": 1201}}, headers={"Idempotency-Key": "workspace-manual-1"})
        assert mismatch.status_code == 409
        manual2 = post(client, "/statements/manual", {"fields": {"交易日": "2025-01-01", "存入": 200, "銷帳編號": "000001", "備註": "=1+1"}}, "workspace-manual-2")
        assert manual2["saved_count"] == 1
        assert post(client, "/statements/manual", body, "workspace-manual-distinct")["saved_count"] == 1
        assert post(client, "/statements/manual", {**body, "duplicate_of": manual["identity"]}, "workspace-manual-duplicate")["excluded_count"] == 1
        assert export(client)[0] == 7
        invalids = ({"交易日": "bad", "支出": 1, "備註": "x"}, {"交易日": "2026-10-29", "支出": 1, "存入": 2, "備註": "x"}, {"交易日": "2026-10-29", "存入": 1})
        for index, fields in enumerate(invalids):
            assert client.post("/api/v1/finance-import/statements/manual", json={"fields": fields}, headers={"Idempotency-Key": f"invalid-{index}"}).status_code == 422
        sheet = workbook["永豐"]
        headers = [cell.value for cell in sheet[1]]
        bank_id = next(row[headers.index("明細識別")] for row in sheet.iter_rows(min_row=2, values_only=True) if row[headers.index("票據號碼")] == "refund")
        duplicate = {"candidate_identity": bank_id, "kept_identity": manual["identity"]}
        assert post(client, "/statements/duplicates", duplicate, "workspace-bank-duplicate")["excluded_count"] == 1
        assert post(client, "/statements/duplicates", duplicate, "workspace-bank-duplicate")["excluded_count"] == 1
        count, workbook = export(client)
        assert count == 6
        manual_sheet = workbook["人工新增"]
        assert any(cell.value == "000001" for row in manual_sheet for cell in row)
        formula = next(cell for row in manual_sheet for cell in row if cell.value == "=1+1")
        assert formula.data_type == "s"
        connection = get_connection()
        try:
            assert _sinopac_memo_readback(connection) == unchanged
            with connection.cursor() as cursor:
                cursor.execute("SELECT COUNT(*) AS count FROM finance_manual_statements")
                assert cursor.fetchone()["count"] == 3
                cursor.execute("SELECT COUNT(*) AS count FROM finance_import_rows")
                assert cursor.fetchone()["count"] == 10
                cursor.execute("SHOW COLUMNS FROM finance_statement_exclusions")
                assert {row["Field"] for row in cursor.fetchall()} == {"finance_import_row_id", "kept_identity", "actor", "created_at"}
        finally:
            connection.close()
        _verify_case_accounting_multi_case_allocation(tmp_path, client)
        count = export(client)[0]
    (tmp_path / "six-features-receipt.json").write_text(json.dumps({
        "database": DATABASE, "archive_and_case_api": "PASS", "original_import": "PASS",
        "manual_saved": 3, "excluded_raw_saved": 0, "export_count": count,
        "owner_facts_unchanged_by_archive": "PASS", "case_isolation": "PASS",
        "multi_case_bank_amount_not_double_counted": "PASS",
    }, ensure_ascii=False, indent=2), encoding="utf-8")


def _verify_case_accounting_multi_case_allocation(tmp_path, client):
    from domains.staff_payables.reconciliation import StaffPayoutEventType
    from infrastructure.mysql.mysql_adapter import get_connection
    from infrastructure.mysql.staff_payout_repository import MySqlStaffPayoutRepository, StaffPayoutMySqlUnitOfWork
    from scripts.imports.finance_formats.sinopac import SINOPAC_HEADERS
    from shared_kernel.identities import ActorContext, CorrelationId, ExpectedVersion, IdempotencyKey
    from subsystems.finance_import.ingestion import ingest_finance_workbook
    from subsystems.staff_payables.payout_reconciliation import StaffPayoutReconciliationWorkflow, StaffPayoutSelection, StaffPayoutApplyRequest
    row = dict.fromkeys(SINOPAC_HEADERS)
    row.update({"帳號": "10000001", "交易日": "2026-10-29 10:00:00", "幣別": "TWD", "支出": 21000, "存入": 0, "餘額": 70000, "備註": "HC115701服務人員", "交易參考編號": "8" * 16, "票據號碼": "explicit-multi-case"})
    path = tmp_path / "explicit-multi-case.xlsx"
    pd.DataFrame([row], columns=SINOPAC_HEADERS).to_excel(path, index=False)
    ingest_finance_workbook(str(path), IdempotencyKey("workspace-multi-case-ingest"), ActorContext("workspace-test"), connection_factory=get_connection, normalizer=normalize_workbook)
    connection = get_connection()
    try:
        with connection.cursor() as cursor:
            cursor.execute("SELECT id FROM finance_import_rows WHERE memo='HC115701服務人員' AND debit=21000")
            row_id = cursor.fetchone()["id"]
        selection = StaffPayoutSelection(StaffPayoutEventType.PAYOUT, (str(row_id),), ("staff:115000701", "staff:115000704"))
        repository = MySqlStaffPayoutRepository(connection)
        workflow = StaffPayoutReconciliationWorkflow(repository, lambda: StaffPayoutMySqlUnitOfWork(connection))
        preview = workflow.preview(selection, CorrelationId("workspace-multi-case"))
        request = StaffPayoutApplyRequest(selection, ExpectedVersion(preview.staff_payables_version), ExpectedVersion(preview.bank_facts_version), preview.fingerprint, IdempotencyKey("workspace-multi-case-post"), ActorContext("workspace-test"), "explicit owner multi-case allocation", CorrelationId("workspace-multi-case"))
        repository.bind_apply_request(request)
        try:
            workflow.apply(request)
        finally:
            repository.clear_apply_request()
    finally:
        connection.close()
    for case_no, expected in (("115000701", 9000), ("115000704", 12000)):
        result = client.get(f"/api/v1/finance-import/cases/{case_no}/accounting")
        assert result.status_code == 200, result.text
        data = result.json()["data"]
        assert sum(item["actual_ntd"] for item in data["staff_items"]) == expected
        assert sum(item["amount_ntd"] for item in data["transactions"] if item["kind"] == "payout") == expected


def test_accounting_warning_boundaries_and_persistent_dismissal_mysql(tmp_path):
    """Use new case 707 and real difference/overage writers, never synthetic ledgers."""
    from dataclasses import asdict
    from io import BytesIO
    import json

    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from api.dependencies.admin_auth import require_admin
    from api.dependencies.finance_import import build_finance_import_application
    from api.routes.finance_import import router
    from domains.client_finance.reconciliation import PaymentStage
    from domains.staff_payables.reconciliation import StaffPayoutDifferenceMode, StaffPayoutEventType
    from infrastructure.mysql.client_receipt_reconciliation_repository import MySqlClientReceiptReconciliationRepository, ClientReceiptMySqlUnitOfWork
    from infrastructure.mysql.client_refund_reversal_repository import MySqlClientRefundReversalRepository, ClientRefundReversalMySqlUnitOfWork
    from domains.client_finance.client_refund_reversal import ClientFinanceCorrectionType
    from subsystems.client_finance.client_refund_reversal_workflow import ClientRefundReversalWorkflow, ClientRefundReversalSelection, ClientRefundReversalApplyRequest
    from infrastructure.mysql.finance_import_owning_domain_composite import MySqlFinanceImportOwningDomainComposite
    from infrastructure.mysql.mysql_adapter import get_connection
    from infrastructure.mysql.staff_payout_repository import MySqlStaffPayoutRepository, StaffPayoutMySqlUnitOfWork
    from scripts.imports.finance_formats.sinopac import SINOPAC_HEADERS
    from shared_kernel.identities import ActorContext, CorrelationId, ExpectedVersion, IdempotencyKey
    from subsystems.access.authentication_session import AdminPrincipal
    from subsystems.client_finance.reconciliation_workflow import ClientReconciliationWorkflow, ClientReconciliationApplyRequest, ReconciliationSelection
    from subsystems.finance_import.ingestion import ingest_finance_workbook
    from subsystems.finance_import.import_workflow import FinanceImportApplyRequest
    from subsystems.staff_payables.payout_reconciliation import StaffPayoutReconciliationWorkflow, StaffPayoutApplyRequest, StaffPayoutSelection
    from subsystems.client_finance.virtual_account_resolution import build_client_virtual_account

    connection = get_connection()
    try:
        _seed_workspace_difference_case(connection)
    finally:
        connection.close()
    amounts = [9970, 9901, 9900, 9899, 10001]
    rows = []
    scenarios = [(f"staff-{index}", amount, 0, str(index) * 16) for index, amount in enumerate(amounts, 1)]
    scenarios.extend([("refund-30", 9970, 0, "6" * 16), ("receipt-overage", 0, 10030, "")])
    for index, (label, debit, credit, account) in enumerate(scenarios, 1):
        row = dict.fromkeys(SINOPAC_HEADERS)
        row.update({"帳號": "10000001", "交易日": f"2026-10-29 11:00:{index:02d}", "計息日": "2026-10-29", "入帳日": "2026-10-29", "幣別": "TWD", "支出": debit, "存入": credit, "餘額": 200000 - index, "摘要": "轉帳", "備註": "HC115707退款" if label.startswith("refund") else "A115707服務人員", "交易參考編號": account, "票據號碼": label})
        if credit:
            row["銷帳編號"] = build_client_virtual_account("115000707")
        rows.append(row)
    workbook = tmp_path / "warning-boundaries.xlsx"
    pd.DataFrame(rows, columns=SINOPAC_HEADERS).to_excel(workbook, index=False)
    receipt = ingest_finance_workbook(str(workbook), IdempotencyKey("warning-ingest"), ActorContext("warning-test"), connection_factory=get_connection, normalizer=normalize_workbook)
    connection = get_connection()
    try:
        with connection.cursor() as cursor:
            cursor.execute("SELECT id,JSON_UNQUOTE(JSON_EXTRACT(bank_references,'$.\"票據號碼\"')) AS label FROM finance_import_rows WHERE memo LIKE '%%115707%%'")
            ids = {row["label"]: row["id"] for row in cursor.fetchall()}
        application = build_finance_import_application(connection, MySqlFinanceImportOwningDomainComposite(connection))
        preview = application.preview_batch(receipt.batch_identity, CorrelationId("warning-preview"))
        applied = application.apply_batch(FinanceImportApplyRequest(receipt.batch_identity, ExpectedVersion(preview.batch_version), preview.fingerprint, IdempotencyKey("warning-batch-apply"), ActorContext("warning-test"), "synthetic refund underpayment", CorrelationId("warning-batch")))
        assert applied.reconciled_count == 0
        refund_repository = MySqlClientRefundReversalRepository(connection)
        refund_workflow = ClientRefundReversalWorkflow(refund_repository, lambda: ClientRefundReversalMySqlUnitOfWork(connection))
        refund_selection = ClientRefundReversalSelection("115000707", ClientFinanceCorrectionType.REFUND, bank_fact_identities=(str(ids["refund-30"]),), obligation_identities=("workspace-refund",), allow_partial_refund_recovery=True)
        refund_preview = refund_workflow.preview(refund_selection, CorrelationId("warning-partial-refund"))
        refund_request = ClientRefundReversalApplyRequest(refund_selection, ExpectedVersion(refund_preview.account_version), refund_preview.fingerprint, IdempotencyKey("warning-partial-refund"), ActorContext("warning-test"), "explicit synthetic partial refund", CorrelationId("warning-partial-refund"))
        refund_repository.bind_apply_request(refund_request)
        try:
            refund_workflow.apply(refund_request)
        finally:
            refund_repository.clear_apply_request()
        for index, amount in enumerate(amounts, 1):
            repository = MySqlStaffPayoutRepository(connection)
            workflow = StaffPayoutReconciliationWorkflow(repository, lambda: StaffPayoutMySqlUnitOfWork(connection))
            selection = StaffPayoutSelection(StaffPayoutEventType.PAYOUT, (str(ids[f"staff-{index}"]),), (f"workspace-staff:{index}",), difference_mode=StaffPayoutDifferenceMode.UNDERPAYMENT if amount < 10000 else StaffPayoutDifferenceMode.OVERPAYMENT)
            preview = workflow.preview(selection, CorrelationId(f"warning-staff-{index}"))
            request = StaffPayoutApplyRequest(selection, ExpectedVersion(preview.staff_payables_version), ExpectedVersion(preview.bank_facts_version), preview.fingerprint, IdempotencyKey(f"warning-staff-{index}"), ActorContext("warning-test"), "explicit synthetic difference", CorrelationId(f"warning-staff-{index}"))
            repository.bind_apply_request(request)
            try:
                workflow.apply(request)
            finally:
                repository.clear_apply_request()
        repository = MySqlClientReceiptReconciliationRepository(connection)
        workflow = ClientReconciliationWorkflow(repository, lambda: ClientReceiptMySqlUnitOfWork(connection))
        selection = ReconciliationSelection("115000707", PaymentStage.DEPOSIT, (str(ids["receipt-overage"]),), ("workspace-deposit",), True)
        preview = workflow.preview(selection)
        request = ClientReconciliationApplyRequest(selection, ExpectedVersion(preview.account_version), preview.fingerprint, IdempotencyKey("warning-receipt-overage"), ActorContext("warning-test"), "explicit synthetic overage", CorrelationId("warning-receipt-overage"))
        repository.bind_apply_request(request)
        try:
            workflow.apply(request)
        finally:
            repository.clear_apply_request()
        before = _workspace_financial_snapshot(connection)
    finally:
        connection.close()
    app = FastAPI(); app.include_router(router)
    app.dependency_overrides[require_admin] = lambda: AdminPrincipal(None, "warning-user-a", "Synthetic A", "system_admin")
    path = "/api/v1/finance-import/cases/115000707/accounting"
    with TestClient(app) as client:
        response = client.get(path)
        assert response.status_code == 200, response.text
        data = response.json()["data"]
        differences = {row["identity"]: row["expected_ntd"] - row["actual_ntd"] for row in data["staff_items"]}
        assert differences == {f"workspace-staff:{index}": 10000 - amount for index, amount in enumerate(amounts, 1)}
        assert data["received_ntd"] == 10030
        assert data["pending_warning_count"] == 3
        assert {(row["kind"], row["difference_ntd"]) for row in data["warnings"]} == {("under_paid", 101), ("over_paid", -1), ("over_received", -30)}
        target = next(row for row in data["warnings"] if row["kind"] == "under_paid")
        body = {key: value for key, value in target.items() if key not in {"dismissed", "dismissed_by", "dismissed_at"}}
        result = client.post(path + "/dismiss", json=body)
        assert result.status_code == 200, result.text
        assert result.json()["data"]["pending_warning_count"] == 2
        assert client.post(path + "/dismiss", json=body).status_code == 200
        app.dependency_overrides[require_admin] = lambda: AdminPrincipal(None, "warning-user-b", "Synthetic B", "system_admin")
        assert client.get(path).json()["data"]["pending_warning_count"] == 2
        dismissed = next(row for row in client.get(path).json()["data"]["warnings"] if row["kind"] == "under_paid")
        assert dismissed["dismissed"] and dismissed["dismissed_by"] == "warning-user-a"
        stale = {**body, "source_version": "0:0:0:0"}
        assert client.post(path + "/dismiss", json=stale).status_code == 409
        all_warnings = client.get("/api/v1/finance-import/accounting/warnings")
        assert all_warnings.status_code == 200 and any(row["case_no"] == "115000707" for row in all_warnings.json()["data"])
    connection = get_connection()
    try:
        assert _workspace_financial_snapshot(connection) == before
        with connection.cursor() as cursor:
            cursor.execute("SELECT COUNT(*) AS count FROM finance_accounting_dismissals WHERE case_no='115000707'")
            assert cursor.fetchone()["count"] == 1
            cursor.execute("SELECT e.id FROM staff_payout_events e JOIN staff_payout_obligation_links l ON l.payout_event_id=e.id WHERE l.obligation_identity='workspace-staff:4' AND e.event_type='payout'")
            source_id = cursor.fetchone()["id"]
        repository = MySqlStaffPayoutRepository(connection)
        workflow = StaffPayoutReconciliationWorkflow(repository, lambda: StaffPayoutMySqlUnitOfWork(connection))
        selection = StaffPayoutSelection(StaffPayoutEventType.REVERSAL, (), ("workspace-staff:4",), reopen_fact_identity=f"reversal:source:{source_id}:on:2026-10-30")
        preview = workflow.preview(selection, CorrelationId("warning-new-anomaly"))
        request = StaffPayoutApplyRequest(selection, ExpectedVersion(preview.staff_payables_version), ExpectedVersion(preview.bank_facts_version), preview.fingerprint, IdempotencyKey("warning-new-anomaly"), ActorContext("warning-test"), "synthetic reversal causes a new difference", CorrelationId("warning-new-anomaly"))
        repository.bind_apply_request(request)
        try:
            workflow.apply(request)
        finally:
            repository.clear_apply_request()
    finally:
        connection.close()
    with TestClient(app) as client:
        fresh = client.get(path).json()["data"]
        new_warning = next(row for row in fresh["warnings"] if row["item_identity"] == "workspace-staff:4")
        assert new_warning["difference_ntd"] == 10000 and not new_warning["dismissed"]
        assert fresh["pending_warning_count"] == 3
        assert client.post(path + "/dismiss", json=body).status_code == 409
        assert next(row for row in fresh["transactions"] if row["kind"] == "reversal")["amount_ntd"] == -9899
    (tmp_path / "warning-receipt.json").write_text(json.dumps({"database": DATABASE, "new_case": "115000707", "differences": differences, "warning_types": ["under_paid", "over_paid", "over_received"], "persistent_cross_user_dismissal": "PASS", "new_anomaly_after_reversal": "PASS", "stale_conflict": "PASS", "accounting_unchanged_by_dismissal": "PASS"}, indent=2), encoding="utf-8")


def _seed_workspace_difference_case(connection):
    import hashlib
    with connection.cursor() as cursor:
        cursor.execute("SELECT COUNT(*) AS count FROM orders WHERE case_no='115000707'")
        assert cursor.fetchone()["count"] == 0, "requires a new synthetic case"
        cursor.execute("INSERT INTO clients(case_no,name) VALUES ('115000707','Workspace Difference Client')")
        cursor.execute("INSERT INTO orders(case_no,client_id,status) VALUES ('115000707',%s,'訂單完成')", (cursor.lastrowid,))
        cursor.execute("INSERT INTO beclass_records(query_no,refund_bank_code,refund_account_no) VALUES ('115000707','999',%s)", ("6" * 16,))
        for kind, direction in (("deposit", "receivable_from_client"), ("refund", "payable_to_client")):
            identity = f"workspace-{kind}"
            cursor.execute("INSERT INTO client_obligation_events(obligation_identity,case_no,obligation_type,direction,event_type,before_amount_ntd,after_amount_ntd,after_due_date,source_event_identity,expected_account_version,idempotency_key,actor,reason) VALUES (%s,'115000707',%s,%s,'established',0,10000,'2026-10-29',%s,0,%s,'workspace-test','synthetic root')", (identity, kind, direction, identity, identity))
            cursor.execute("INSERT INTO client_obligations(obligation_identity,case_no,obligation_type,direction,amount_due_ntd,due_date,status,current_event_id,projection_version) VALUES (%s,'115000707',%s,%s,10000,'2026-10-29','open',%s,0)", (identity, kind, direction, cursor.lastrowid))
            if kind == "refund":
                cursor.execute("INSERT INTO client_refund_recipient_snapshots(refund_obligation_identity,case_no,bank_code,bank_account,source_kind) VALUES (%s,'115000707','999',%s,'test-fixture')", (identity, "6" * 16))
        cursor.execute("INSERT INTO payroll_case_accounts(case_no,aggregate_version) VALUES ('115000707',1)")
        for index in range(1, 6):
            cursor.execute("INSERT INTO staff(name) VALUES (%s)", (f"Workspace Staff {index}",))
            staff_id = cursor.lastrowid
            cursor.execute("INSERT INTO staff_bank_accounts(staff_id,bank_code,account_no,is_primary) VALUES (%s,'999',%s,1)", (staff_id, str(index) * 16))
            cursor.execute("INSERT INTO case_staff_assignments(case_no,staff_id,assignment_sequence,status) VALUES ('115000707',%s,%s,'completed')", (staff_id, index))
            assignment_id = cursor.lastrowid; identity = f"workspace-staff:{index}"
            cursor.execute("INSERT INTO staff_obligation_events(obligation_identity,assignment_id,case_no,staff_id,obligation_kind,direction,event_type,before_amount_ntd,after_amount_ntd,due_date,payroll_fingerprint,expected_payroll_version,resulting_payroll_version,idempotency_key,actor,reason) VALUES (%s,%s,'115000707',%s,'service_pay','payable_to_staff','established',0,10000,'2026-10-29',%s,0,1,%s,'workspace-test','synthetic root')", (identity, assignment_id, staff_id, hashlib.sha256(identity.encode()).hexdigest(), identity))
            cursor.execute("INSERT INTO staff_obligations(obligation_identity,assignment_id,case_no,staff_id,obligation_kind,direction,amount_due_ntd,due_date,status,current_event_id,payroll_version) VALUES (%s,%s,'115000707',%s,'service_pay','payable_to_staff',10000,'2026-10-29','open',%s,1)", (identity, assignment_id, staff_id, cursor.lastrowid))
    connection.commit()


def _workspace_financial_snapshot(connection):
    tables = ("finance_import_rows", "client_ledger_entries", "client_ledger_obligation_allocations", "client_obligations", "staff_payout_events", "staff_payout_obligation_links", "staff_payable_projections", "orders")
    result = {}
    with connection.cursor() as cursor:
        for table in tables:
            cursor.execute(f"SELECT * FROM {table}")
            result[table] = cursor.fetchall()
    return result


def test_customer_refund_overpayment_shared_warning_mysql(tmp_path):
    """Over-refund cash includes its owner recovery source, never only allocations."""
    import json
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from api.dependencies.admin_auth import require_admin
    from api.routes.finance_import import router
    from domains.client_finance.client_refund_reversal import ClientFinanceCorrectionType
    from infrastructure.mysql.client_refund_reversal_repository import MySqlClientRefundReversalRepository, ClientRefundReversalMySqlUnitOfWork
    from infrastructure.mysql.mysql_adapter import get_connection
    from scripts.imports.finance_formats.sinopac import SINOPAC_HEADERS
    from shared_kernel.identities import ActorContext, CorrelationId, ExpectedVersion, IdempotencyKey
    from subsystems.access.authentication_session import AdminPrincipal
    from subsystems.client_finance.client_refund_reversal_workflow import ClientRefundReversalWorkflow, ClientRefundReversalSelection, ClientRefundReversalApplyRequest
    from subsystems.finance_import.ingestion import ingest_finance_workbook

    account = "6" * 15 + "1"
    connection = get_connection()
    try:
        with connection.cursor() as cursor:
            cursor.execute("SELECT COUNT(*) AS count FROM orders WHERE case_no='115000708'")
            assert cursor.fetchone()["count"] == 0
            cursor.execute("INSERT INTO clients(case_no,name) VALUES ('115000708','Synthetic Over Refund Client')")
            cursor.execute("INSERT INTO orders(case_no,client_id,status) VALUES ('115000708',%s,'訂單完成')", (cursor.lastrowid,))
            cursor.execute("INSERT INTO client_obligation_events(obligation_identity,case_no,obligation_type,direction,event_type,before_amount_ntd,after_amount_ntd,after_due_date,source_event_identity,expected_account_version,idempotency_key,actor,reason) VALUES ('workspace-over-refund','115000708','refund','payable_to_client','established',0,10000,'2026-10-29','workspace-over-refund-root',0,'workspace-over-refund-root','workspace-test','synthetic root')")
            cursor.execute("INSERT INTO client_obligations(obligation_identity,case_no,obligation_type,direction,amount_due_ntd,due_date,status,current_event_id,projection_version) VALUES ('workspace-over-refund','115000708','refund','payable_to_client',10000,'2026-10-29','open',%s,0)", (cursor.lastrowid,))
            cursor.execute("INSERT INTO client_refund_recipient_snapshots(refund_obligation_identity,case_no,bank_code,bank_account,source_kind) VALUES ('workspace-over-refund','115000708','999',%s,'test-fixture')", (account,))
        connection.commit()
    finally:
        connection.close()
    row = dict.fromkeys(SINOPAC_HEADERS)
    row.update({"帳號": "10000001", "交易日": "2026-10-29 12:00:00", "幣別": "TWD", "支出": 10001, "存入": 0, "餘額": 80000, "備註": "HC115708退款", "交易參考編號": account})
    path = tmp_path / "client-over-refund.xlsx"
    pd.DataFrame([row], columns=SINOPAC_HEADERS).to_excel(path, index=False)
    ingest_finance_workbook(str(path), IdempotencyKey("workspace-over-refund-ingest"), ActorContext("workspace-test"), connection_factory=get_connection, normalizer=normalize_workbook)
    connection = get_connection()
    try:
        with connection.cursor() as cursor:
            cursor.execute("SELECT id FROM finance_import_rows WHERE memo='HC115708退款'")
            row_id = cursor.fetchone()["id"]
        repository = MySqlClientRefundReversalRepository(connection)
        workflow = ClientRefundReversalWorkflow(repository, lambda: ClientRefundReversalMySqlUnitOfWork(connection))
        selection = ClientRefundReversalSelection("115000708", ClientFinanceCorrectionType.REFUND_OVERAGE, bank_fact_identities=(str(row_id),), obligation_identities=("workspace-over-refund",))
        preview = workflow.preview(selection, CorrelationId("workspace-over-refund"))
        request = ClientRefundReversalApplyRequest(selection, ExpectedVersion(preview.account_version), preview.fingerprint, IdempotencyKey("workspace-over-refund-post"), ActorContext("workspace-test"), "explicit synthetic over refund", CorrelationId("workspace-over-refund"))
        repository.bind_apply_request(request)
        try:
            workflow.apply(request)
        finally:
            repository.clear_apply_request()
        before = _workspace_financial_snapshot(connection)
    finally:
        connection.close()
    app = FastAPI(); app.include_router(router)
    app.dependency_overrides[require_admin] = lambda: AdminPrincipal(None, "over-refund-test", "Synthetic Operator", "system_admin")
    with TestClient(app) as client:
        response = client.get("/api/v1/finance-import/cases/115000708/accounting")
        assert response.status_code == 200, response.text
        data = response.json()["data"]
        assert (data["payable_ntd"], data["paid_ntd"]) == (10000, 10001)
        assert data["pending_warning_count"] == 1
        assert (data["warnings"][0]["kind"], data["warnings"][0]["difference_ntd"]) == ("over_paid", -1)
        warning = {key: value for key, value in data["warnings"][0].items() if key not in {"dismissed", "dismissed_by", "dismissed_at"}}
        assert client.post("/api/v1/finance-import/cases/115000708/accounting/dismiss", json=warning).status_code == 200
        assert client.get("/api/v1/finance-import/cases/115000708/accounting").json()["data"]["pending_warning_count"] == 0
        assert client.get("/api/v1/finance-import/accounting/warnings").status_code == 200
    connection = get_connection()
    try:
        assert _workspace_financial_snapshot(connection) == before
    finally:
        connection.close()
    (tmp_path / "client-over-refund-receipt.json").write_text(json.dumps({"database": DATABASE, "new_case": "115000708", "expected_ntd": 10000, "actual_ntd": 10001, "over_paid_warning": "PASS", "persistent_dismissal": "PASS", "accounting_unchanged": "PASS"}, indent=2), encoding="utf-8")


def test_multi_obligation_staff_overpayment_is_shared_without_cross_case_cash_mysql(tmp_path):
    """Unallocated surplus has an explicit source, never a guessed case allocation."""
    import hashlib
    import json
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from api.dependencies.admin_auth import require_admin
    from api.routes.finance_import import router
    from domains.staff_payables.reconciliation import StaffPayoutEventType, StaffPayoutDifferenceMode
    from infrastructure.mysql.mysql_adapter import get_connection
    from infrastructure.mysql.staff_payout_repository import MySqlStaffPayoutRepository, StaffPayoutMySqlUnitOfWork
    from scripts.imports.finance_formats.sinopac import SINOPAC_HEADERS
    from shared_kernel.identities import ActorContext, CorrelationId, ExpectedVersion, IdempotencyKey
    from subsystems.access.authentication_session import AdminPrincipal
    from subsystems.finance_import.ingestion import ingest_finance_workbook
    from subsystems.staff_payables.payout_reconciliation import StaffPayoutReconciliationWorkflow, StaffPayoutSelection, StaffPayoutApplyRequest

    account = "5" * 15 + "1"
    connection = get_connection()
    try:
        with connection.cursor() as cursor:
            cursor.execute("INSERT INTO staff(name) VALUES ('Synthetic Multi Item Staff')")
            staff_id = cursor.lastrowid
            cursor.execute("INSERT INTO staff_bank_accounts(staff_id,bank_code,account_no,is_primary) VALUES (%s,'999',%s,1)", (staff_id, account))
            for suffix, identities in ((709, ("multi-same-a", "multi-same-b")), (710, ("multi-cross-a",)), (711, ("multi-cross-b",))):
                case_no = f"115000{suffix}"
                cursor.execute("SELECT COUNT(*) AS count FROM orders WHERE case_no=%s", (case_no,))
                assert cursor.fetchone()["count"] == 0
                cursor.execute("INSERT INTO clients(case_no,name) VALUES (%s,'Synthetic Multi Item Client')", (case_no,))
                cursor.execute("INSERT INTO orders(case_no,client_id,status) VALUES (%s,%s,'訂單完成')", (case_no, cursor.lastrowid))
                cursor.execute("INSERT INTO payroll_case_accounts(case_no,aggregate_version) VALUES (%s,1)", (case_no,))
                cursor.execute("INSERT INTO case_staff_assignments(case_no,staff_id,assignment_sequence,status) VALUES (%s,%s,1,'completed')", (case_no, staff_id))
                assignment_id = cursor.lastrowid
                for identity in identities:
                    cursor.execute("INSERT INTO staff_obligation_events(obligation_identity,assignment_id,case_no,staff_id,obligation_kind,direction,event_type,before_amount_ntd,after_amount_ntd,due_date,payroll_fingerprint,expected_payroll_version,resulting_payroll_version,idempotency_key,actor,reason) VALUES (%s,%s,%s,%s,'service_pay','payable_to_staff','established',0,5000,'2026-10-29',%s,0,1,%s,'workspace-test','synthetic root')", (identity, assignment_id, case_no, staff_id, hashlib.sha256(identity.encode()).hexdigest(), identity))
                    cursor.execute("INSERT INTO staff_obligations(obligation_identity,assignment_id,case_no,staff_id,obligation_kind,direction,amount_due_ntd,due_date,status,current_event_id,payroll_version) VALUES (%s,%s,%s,%s,'service_pay','payable_to_staff',5000,'2026-10-29','open',%s,1)", (identity, assignment_id, case_no, staff_id, cursor.lastrowid))
        connection.commit()
    finally:
        connection.close()
    rows = []
    for suffix in (709, 710):
        row = dict.fromkeys(SINOPAC_HEADERS)
        row.update({"帳號": "10000001", "交易日": f"2026-10-29 13:00:{suffix - 700:02d}", "幣別": "TWD", "支出": 10001, "存入": 0, "餘額": 80000 - suffix, "備註": f"HC115{suffix}服務人員", "交易參考編號": account})
        rows.append(row)
    path = tmp_path / "multi-obligation-overpayment.xlsx"
    pd.DataFrame(rows, columns=SINOPAC_HEADERS).to_excel(path, index=False)
    ingest_finance_workbook(str(path), IdempotencyKey("multi-overpayment-ingest"), ActorContext("workspace-test"), connection_factory=get_connection, normalizer=normalize_workbook)
    connection = get_connection()
    try:
        for suffix, identities in ((709, ("multi-same-a", "multi-same-b")), (710, ("multi-cross-a", "multi-cross-b"))):
            with connection.cursor() as cursor:
                cursor.execute("SELECT id FROM finance_import_rows WHERE memo=%s", (f"HC115{suffix}服務人員",))
                row_id = cursor.fetchone()["id"]
            repository = MySqlStaffPayoutRepository(connection)
            workflow = StaffPayoutReconciliationWorkflow(repository, lambda: StaffPayoutMySqlUnitOfWork(connection))
            selection = StaffPayoutSelection(StaffPayoutEventType.PAYOUT, (str(row_id),), identities, difference_mode=StaffPayoutDifferenceMode.OVERPAYMENT)
            preview = workflow.preview(selection, CorrelationId(f"multi-overpayment-{suffix}"))
            request = StaffPayoutApplyRequest(selection, ExpectedVersion(preview.staff_payables_version), ExpectedVersion(preview.bank_facts_version), preview.fingerprint, IdempotencyKey(f"multi-overpayment-{suffix}"), ActorContext("workspace-test"), "explicit synthetic multi-obligation excess", CorrelationId(f"multi-overpayment-{suffix}"))
            repository.bind_apply_request(request)
            try:
                workflow.apply(request)
            finally:
                repository.clear_apply_request()
        before = _workspace_financial_snapshot(connection)
    finally:
        connection.close()
    app = FastAPI(); app.include_router(router)
    app.dependency_overrides[require_admin] = lambda: AdminPrincipal(None, "multi-over-test", "Synthetic Operator", "system_admin")
    with TestClient(app) as client:
        for suffix, expected, actual in ((709, 10000, 10001), (710, 5000, 5000), (711, 5000, 5000)):
            response = client.get(f"/api/v1/finance-import/cases/115000{suffix}/accounting")
            assert response.status_code == 200, response.text
            data = response.json()["data"]
            assert (data["payable_ntd"], data["paid_ntd"]) == (expected, actual)
            assert data["pending_warning_count"] == 1
        shared = client.get("/api/v1/finance-import/accounting/warnings").json()["data"]
        identities = [warning["item_identity"] for case in shared for warning in case["warnings"] if warning["item_identity"].startswith("staff-overpayment:")]
        assert len(identities) == len(set(identities)) == 2
        data = client.get("/api/v1/finance-import/cases/115000710/accounting").json()["data"]
        body = {key: value for key, value in data["warnings"][0].items() if key not in {"dismissed", "dismissed_by", "dismissed_at"}}
        response = client.post("/api/v1/finance-import/cases/115000710/accounting/dismiss", json=body)
        assert response.status_code == 200, response.text
        assert client.get("/api/v1/finance-import/cases/115000711/accounting").json()["data"]["pending_warning_count"] == 0
    connection = get_connection()
    try:
        assert _workspace_financial_snapshot(connection) == before
    finally:
        connection.close()
    (tmp_path / "multi-overpayment-receipt.json").write_text(json.dumps({"database": DATABASE, "new_cases": ["115000709", "115000710", "115000711"], "same_case_extra_cash": "PASS", "cross_case_no_guessed_allocation": "PASS", "shared_warning_no_duplicates": "PASS", "shared_dismissal": "PASS", "accounting_unchanged": "PASS"}, indent=2), encoding="utf-8")


def test_taishin_period_receipt_and_subsidy_refund_are_in_full_export_mysql(tmp_path):
    """Both current Taishin owners qualify; no legacy transaction row is required."""
    import json
    import runpy
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from api.dependencies import admin_auth
    from api.dependencies.finance_import import build_finance_import_application
    from api.routes import finance_import, government_subsidy_periods, order_intake_terms_bootstrap
    from infrastructure.mysql.finance_import_owning_domain_composite import MySqlFinanceImportOwningDomainComposite
    from infrastructure.mysql.finance_statement_archive_repository import MySqlStatementArchiveRepository
    from infrastructure.mysql.government_subsidy_period_repository import MySqlGovernmentSubsidyPeriodRepository
    from infrastructure.mysql.mysql_adapter import get_connection
    from infrastructure.mysql.unit_of_work import MySqlUnitOfWork
    from shared_kernel.identities import ActorContext, CorrelationId, ExpectedVersion, IdempotencyKey
    from subsystems.access.authentication_session import AdminPrincipal
    from subsystems.finance_import.ingestion import ingest_finance_workbook
    from subsystems.finance_import.import_workflow import FinanceImportApplyRequest
    from domains.government_subsidy.period_reconciliation import SubsidyPeriod
    from subsystems.government_subsidy.period_reconciliation import PeriodReconciliationApplication

    # Reuse the established Government Subsidy owner's intake/source fixture.
    support = runpy.run_path('tests/domains/government-subsidy/subsystems/government-subsidy/integration/test_period_reconciliation_mysql.py')
    app = FastAPI()
    for module in (finance_import, government_subsidy_periods, order_intake_terms_bootstrap):
        app.include_router(module.router)
    principal = AdminPrincipal(1, 'taishin-archive-test', 'Synthetic Operator', 'system_admin')
    for guard in (admin_auth.require_admin, admin_auth.require_system_admin, admin_auth.require_persisted_admin):
        app.dependency_overrides[guard] = lambda: principal
    with TestClient(app, headers={'X-Correlation-ID': 'taishin-archive-test'}) as client:
        case_no = support['seed_case'](client, 2071, 1, days=20)
    connection = get_connection()
    try:
        owner = PeriodReconciliationApplication(MySqlGovernmentSubsidyPeriodRepository(connection), lambda: MySqlUnitOfWork(connection))
        period = SubsidyPeriod('month', 2071, 1)
        view = owner.query(period)
        assert view.can_reconcile and view.total_ntd > 0
        amount = view.total_ntd
    finally:
        connection.close()
    headers = ['序號', '交易日期', '交易時間', '帳務日期', '摘要', '支出金額', '存入金額', '帳戶餘額', '備註']
    path = tmp_path / 'taishin-current-period.xlsx'
    pd.DataFrame([['說明'], headers, ['0001', '2071/02/15', '09:00:00', '2071/02/15', '政府補助', '', str(amount), '80000', '新竹市政府補助入款']]).to_excel(path, sheet_name='交易明細', index=False, header=False)
    ingest_finance_workbook(str(path), IdempotencyKey('taishin-current-period'), ActorContext('workspace-test'), connection_factory=get_connection, normalizer=normalize_workbook)
    connection = get_connection()
    try:
        with connection.cursor() as cursor:
            cursor.execute("SELECT id FROM finance_import_rows WHERE transaction_date='2071-02-15' AND direction='incoming'")
            bank_id = cursor.fetchone()['id']
        archive = MySqlStatementArchiveRepository(connection)
        assert f'bank:{bank_id}' not in {row.identity for row in archive.export_rows()}
        owner = PeriodReconciliationApplication(MySqlGovernmentSubsidyPeriodRepository(connection), lambda: MySqlUnitOfWork(connection))
        owner.reconcile(period, bank_id, view.version, 'taishin-period-post', 'workspace-test')
        rows = archive.export_rows()
        statement = next(row for row in rows if row.identity == f'bank:{bank_id}')
        assert statement.bank == 'taishin' and statement.cases == (case_no,)
        with connection.cursor() as cursor:
            cursor.execute('SELECT COUNT(*) AS count FROM government_subsidy_transactions WHERE finance_import_row_id=%s', (bank_id,))
            assert cursor.fetchone()['count'] == 0
            cursor.execute("SELECT obligation_identity,amount_due_ntd FROM client_obligations WHERE case_no=%s AND obligation_type='subsidy_return'", (case_no,))
            refund = cursor.fetchone()
            account = '7' * 15 + '2'
            cursor.execute("INSERT INTO client_refund_recipient_snapshots(refund_obligation_identity,case_no,bank_code,bank_account,source_kind) VALUES (%s,%s,'999',%s,'test-fixture')", (refund['obligation_identity'], case_no, account))
        connection.commit()
    finally:
        connection.close()
    path = tmp_path / 'taishin-subsidy-refund.xlsx'
    pd.DataFrame([['說明'], headers, ['0002', '2071/02/16', '09:00:00', '2071/02/16', '補助退款', str(refund['amount_due_ntd']), '', '60000', '補助退款 ' + account]]).to_excel(path, sheet_name='交易明細', index=False, header=False)
    ingestion = ingest_finance_workbook(str(path), IdempotencyKey('taishin-current-refund'), ActorContext('workspace-test'), connection_factory=get_connection, normalizer=normalize_workbook)
    connection = get_connection()
    try:
        application = build_finance_import_application(connection, MySqlFinanceImportOwningDomainComposite(connection))
        preview = application.preview_batch(ingestion.batch_identity, CorrelationId('taishin-current-refund'))
        receipt = application.apply_batch(FinanceImportApplyRequest(ingestion.batch_identity, ExpectedVersion(preview.batch_version), preview.fingerprint, IdempotencyKey('taishin-current-refund-post'), ActorContext('workspace-test'), 'synthetic subsidy refund', CorrelationId('taishin-current-refund')))
        assert receipt.reconciled_count == 1
        with connection.cursor() as cursor:
            cursor.execute("SELECT finance_import_row_id FROM client_ledger_entries WHERE case_no=%s AND entry_type='subsidy_return'", (case_no,))
            refund_bank_id = cursor.fetchone()['finance_import_row_id']
        rows = MySqlStatementArchiveRepository(connection).export_rows()
        assert {f'bank:{bank_id}', f'bank:{refund_bank_id}'} <= {row.identity for row in rows}
    finally:
        connection.close()
    (tmp_path / 'taishin-archive-receipt.json').write_text(json.dumps({'database': DATABASE, 'new_case': case_no, 'period_receipt_export': 'PASS', 'subsidy_refund_export': 'PASS', 'legacy_transaction_required': False}, indent=2), encoding='utf-8')


def _arguments() -> Namespace:
    return Namespace(
        host=os.environ["LABOR_UNION_TEST_MYSQL_HOST"],
        port=int(os.environ["LABOR_UNION_TEST_MYSQL_PORT"]),
        user=os.environ["LABOR_UNION_TEST_MYSQL_USER"],
        password=os.environ["LABOR_UNION_TEST_MYSQL_PASSWORD"],
        database=DATABASE,
        confirm_database=DATABASE,
    )


@pytest.fixture(autouse=True)
def _use_explicit_disposable_database(monkeypatch):
    """Keep every legacy adapter import on this test's isolated database."""
    database_settings = {
        "DB_HOST": os.environ["LABOR_UNION_TEST_MYSQL_HOST"],
        "DB_PORT": os.environ["LABOR_UNION_TEST_MYSQL_PORT"],
        "DB_USER": os.environ["LABOR_UNION_TEST_MYSQL_USER"],
        "DB_PASSWORD": os.environ["LABOR_UNION_TEST_MYSQL_PASSWORD"],
        "DB_DATABASE": DATABASE,
    }
    for setting_name, setting_value in database_settings.items():
        monkeypatch.setenv(setting_name, setting_value)

    from infrastructure.mysql import mysql_adapter

    monkeypatch.setattr(
        mysql_adapter,
        "DB_CONFIG",
        {
            "host": database_settings["DB_HOST"],
            "port": int(database_settings["DB_PORT"]),
            "user": database_settings["DB_USER"],
            "password": database_settings["DB_PASSWORD"],
            "database": database_settings["DB_DATABASE"],
            "charset": "utf8mb4",
        },
    )


def test_deidentified_taishin_fixture_becomes_root_fact_and_unresolved_reprocess_blocks():
    bootstrap(_arguments())
    workbook = DEIDENTIFIED_TAISHIN_FIXTURE
    assert workbook.is_file()

    from infrastructure.mysql.mysql_adapter import get_connection
    from subsystems.finance_import.ingestion import ingest_finance_workbook
    from shared_kernel.identities import ActorContext, CorrelationId, IdempotencyKey

    receipt = ingest_finance_workbook(
        str(workbook),
        IdempotencyKey("lu-test-unresolved-historical-reprocess"),
        ActorContext("lu-test-runner"),
        connection_factory=get_connection,
        normalizer=normalize_workbook,
    )

    assert receipt.source_row_count == receipt.canonical_created_count == 1
    connection = get_connection()
    try:
        with connection.cursor() as cursor:
            cursor.execute(
                "SELECT classification_type FROM finance_import_classification_events"
            )
            assert cursor.fetchone()["classification_type"] == "non_business_review"
    finally:
        connection.close()

    from infrastructure.mysql.finance_import_owning_domain_composite import MySqlFinanceImportOwningDomainComposite
    from infrastructure.mysql.finance_import_repository import FinanceImportMySqlUnitOfWork
    from infrastructure.mysql.historical_reprocess_repository import MySqlHistoricalReprocessRepository
    from subsystems.finance_import.historical_reprocess_workflow import HistoricalReprocessWorkflow, HistoricalReprocessWorkflowError

    connection = get_connection()
    try:
        workflow = HistoricalReprocessWorkflow(
            MySqlHistoricalReprocessRepository(connection),
            MySqlFinanceImportOwningDomainComposite(connection),
            lambda: FinanceImportMySqlUnitOfWork(connection),
        )
        with pytest.raises(HistoricalReprocessWorkflowError) as error:
            workflow.preview(receipt.batch_identity, CorrelationId("lu-test-preview"))
    finally:
        connection.close()

    assert error.value.error.code == "reprocess_owner_not_resolved"


def test_mixed_finance_workbook_keeps_valid_row_and_acknowledges_owner_review(tmp_path):
    bootstrap(_arguments())
    workbook = tmp_path / "taishin-mixed-source-review.xlsx"
    pd.DataFrame(
        [
            ["說明"],
            ["序號", "交易日期", "交易時間", "帳務日期", "摘要", "支出金額", "存入金額", "帳戶餘額", "備註"],
            ["0001", "2026/08/04", "09:08:07", "2026/08/04", "轉帳", "300", "", "9000", "正常列"],
            ["0002", "2026/08/05", "09:09:07", "2026/08/05", "轉帳", "300.5", "", "8699.5", "敏感備註 0912345678"],
        ]
    ).to_excel(workbook, sheet_name="交易明細", index=False, header=False)

    from infrastructure.mysql.mysql_adapter import get_connection
    from infrastructure.mysql.anomaly_runtime import build_anomaly_runtime
    from shared_kernel.identities import ActorContext, IdempotencyKey
    from subsystems.finance_import.finance_import_anomaly_consumer import (
        consume_finance_import_anomaly_events,
    )
    from subsystems.finance_import.ingestion import ingest_finance_workbook

    receipt = ingest_finance_workbook(
        str(workbook),
        IdempotencyKey("lu-test-finance-source-review"),
        ActorContext("lu-test-runner"),
        connection_factory=get_connection,
        normalizer=normalize_workbook,
    )
    replay = ingest_finance_workbook(
        str(workbook),
        IdempotencyKey("lu-test-finance-source-review"),
        ActorContext("lu-test-runner"),
        connection_factory=get_connection,
        normalizer=normalize_workbook,
    )

    assert receipt == replay
    assert receipt.source_row_count == 2
    assert receipt.canonical_created_count == 1
    assert receipt.source_warning_count == 1
    assert receipt.source_warning_created_count == 1

    connection = get_connection()
    try:
        projected = consume_finance_import_anomaly_events(
            connection, runtime=build_anomaly_runtime()
        )
        assert projected.failed_count == 0
        with connection.cursor() as cursor:
            cursor.execute(
                "SELECT COUNT(*) AS count FROM finance_import_rows"
            )
            assert cursor.fetchone() == {"count": 1}
            cursor.execute(
                "SELECT source_identity,issue_codes "
                "FROM finance_import_source_reviews"
            )
            stored_review = cursor.fetchone()
            assert stored_review["source_identity"].startswith("finance-taishin-row-")
            assert "invalid:transaction_amount" in stored_review["issue_codes"]
            assert "0912345678" not in str(stored_review)
            assert "敏感備註" not in str(stored_review)
            cursor.execute(
                "SELECT COUNT(*) AS count FROM finance_import_source_review_outbox "
                "WHERE published_at IS NOT NULL"
            )
            assert cursor.fetchone() == {"count": 1}
            cursor.execute(
                "SELECT COUNT(*) AS count FROM import_warning_occurrences "
                "WHERE owning_lane='finance_import'"
            )
            assert cursor.fetchone() == {"count": 0}
    finally:
        connection.close()


def test_historical_owner_selection_posts_once_without_mutating_bank_root_fact(tmp_path):
    bootstrap(_arguments())
    intake_receipt = _ingest_unresolved_taishin_outflow(tmp_path)
    _seed_open_refund_obligation()

    from api.dependencies.finance_import import HistoricalReprocessApplication
    from infrastructure.mysql.finance_import_owning_domain_composite import (
        MySqlFinanceImportOwningDomainComposite,
    )
    from infrastructure.mysql.finance_import_repository import FinanceImportMySqlUnitOfWork
    from infrastructure.mysql.historical_reprocess_repository import MySqlHistoricalReprocessRepository
    from infrastructure.mysql.mysql_adapter import get_connection
    from shared_kernel.identities import (
        ActorContext,
        CorrelationId,
        ExpectedVersion,
        IdempotencyKey,
    )
    from subsystems.finance_import.historical_reprocess_workflow import (
        HistoricalOwnerSelection,
        HistoricalReprocessApplyRequest,
        HistoricalReprocessWorkflow,
    )

    connection = get_connection()
    try:
        with connection.cursor() as cursor:
            cursor.execute(
                "SELECT bank_references FROM finance_import_rows WHERE id=1"
            )
            original_bank_references = cursor.fetchone()["bank_references"]
        posting_port = MySqlFinanceImportOwningDomainComposite(connection)
        application = HistoricalReprocessApplication(
            HistoricalReprocessWorkflow(
                MySqlHistoricalReprocessRepository(connection),
                posting_port,
                lambda: FinanceImportMySqlUnitOfWork(connection),
            ),
            posting_port,
        )
        selection = HistoricalOwnerSelection(
            "finance-import-row:1",
            "C-1",
            "refund:C-1",
            "reviewed original bank statement and refund request",
            ("bank-statement:line-3", "refund-request:C-1"),
        )
        preview = application.preview(
            intake_receipt.batch_identity,
            CorrelationId("historical-owner-selection-preview"),
            (selection,),
        )
        request = HistoricalReprocessApplyRequest(
            intake_receipt.batch_identity,
            ExpectedVersion(preview.batch_version),
            preview.fingerprint,
            IdempotencyKey("historical-owner-selection-apply"),
            ActorContext("lu-test-runner"),
            "apply reviewed historical refund owner",
            CorrelationId("historical-owner-selection-apply"),
            (selection,),
        )

        receipt = application.apply(request)
        assert application.apply(request) == receipt

        with connection.cursor() as cursor:
            cursor.execute(
                "SELECT bank_references FROM finance_import_rows WHERE id=1"
            )
            assert cursor.fetchone()["bank_references"] == original_bank_references
            cursor.execute(
                "SELECT case_no,obligation_identity,source_canonical_fact_version,"
                "resulting_canonical_fact_version,obligation_projection_version "
                "FROM historical_owner_selection_events"
            )
            assert cursor.fetchone() == {
                "case_no": "C-1",
                "obligation_identity": "refund:C-1",
                "source_canonical_fact_version": 0,
                "resulting_canonical_fact_version": 1,
                "obligation_projection_version": 0,
            }
            cursor.execute("SELECT actor FROM finance_import_reprocess_runs")
            assert cursor.fetchone() == {"actor": "lu-test-runner"}
            cursor.execute("SELECT entry_type,amount_ntd FROM client_ledger_entries")
            assert cursor.fetchone() == {"entry_type": "refund", "amount_ntd": 300}
            cursor.execute(
                "SELECT amount_due_ntd,status FROM client_obligations "
                "WHERE obligation_identity='refund:C-1'"
            )
            assert cursor.fetchone() == {"amount_due_ntd": 0, "status": "settled"}
            cursor.execute(
                "SELECT COUNT(*) AS count FROM finance_import_historical_reprocess_receipts"
            )
            assert cursor.fetchone() == {"count": 1}
    finally:
        connection.close()


def test_manual_refund_correction_posts_ledger_allocation_and_resolves_anomaly(tmp_path):
    bootstrap(_arguments())
    _seed_open_refund_obligation()
    _ingest_unresolved_taishin_outflow(tmp_path)
    _deliver_finance_import_outbox()

    from domains.finance_import.correction import FinanceImportCorrectionSelection
    from domains.finance_import.planning import FinanceClassificationType
    from api.dependencies.finance_import import build_finance_import_application
    from infrastructure.mysql.finance_import_owning_domain_composite import MySqlFinanceImportOwningDomainComposite
    from infrastructure.mysql.mysql_adapter import get_connection
    from shared_kernel.identities import ActorContext, CorrelationId, ExpectedVersion, IdempotencyKey
    from subsystems.finance_import.correction_workflow import FinanceImportCorrectionApplyRequest

    connection = get_connection()
    try:
        application = build_finance_import_application(
            connection,
            MySqlFinanceImportOwningDomainComposite(connection),
        )
        selection = FinanceImportCorrectionSelection(
            "finance-import-row:1",
            FinanceClassificationType.CLIENT_REFUND,
            ("refund:C-1",),
            "bank statement and customer refund request were reviewed",
            ("bank-statement:line-3", "customer-refund-request:C-1"),
        )
        preview = application.preview_correction(selection, CorrelationId("lu-test-correction-preview"))
        receipt = application.correct_and_post(
            FinanceImportCorrectionApplyRequest(
                selection,
                ExpectedVersion(preview.batch_version),
                ExpectedVersion(preview.canonical_fact_version),
                ExpectedVersion(preview.alert_version),
                preview.fingerprint,
                IdempotencyKey("lu-test-manual-refund-correction"),
                ActorContext("lu-test-runner"),
                CorrelationId("lu-test-correction-apply"),
            )
        )
        with connection.cursor() as cursor:
            cursor.execute("SELECT entry_type,amount_ntd FROM client_ledger_entries")
            assert cursor.fetchone() == {"entry_type": "refund", "amount_ntd": 300}
            cursor.execute("SELECT amount_due_ntd,status FROM client_obligations WHERE obligation_identity='refund:C-1'")
            assert cursor.fetchone() == {"amount_due_ntd": 0, "status": "settled"}
    finally:
        connection.close()

    _deliver_finance_import_outbox()
    assert receipt.ledger_entry_count == receipt.allocation_count == 1
    _assert_manual_review_stays_out_of_anomaly_center()


def test_durable_correction_worker_posts_manual_refund_once(tmp_path):
    """A correction survives API-process loss because the worker owns Apply."""
    bootstrap(_arguments())
    _seed_open_refund_obligation()
    _ingest_unresolved_taishin_outflow(tmp_path)
    _deliver_finance_import_outbox()

    from api.dependencies.finance_import import build_finance_import_application
    from api.routes.finance_import import _correction_apply_job_command
    from domains.finance_import.correction import FinanceImportCorrectionSelection
    from domains.finance_import.planning import FinanceClassificationType
    from infrastructure.mysql.background_job_repository import BackgroundJobRepository
    from infrastructure.mysql.finance_import_owning_domain_composite import MySqlFinanceImportOwningDomainComposite
    from infrastructure.mysql.mysql_adapter import get_connection
    from shared_kernel.identities import CorrelationId
    from api.dependencies.durable_job_handlers import default_job_handlers
    from subsystems.jobs.durable_job_worker import DurableJobWorker
    from subsystems.jobs.command_application import DurableJobCommandApplication

    connection = get_connection()
    try:
        application = build_finance_import_application(
            connection,
            MySqlFinanceImportOwningDomainComposite(connection),
        )
        selection = FinanceImportCorrectionSelection(
            "finance-import-row:1",
            FinanceClassificationType.CLIENT_REFUND,
            ("refund:C-1",),
            "durable manual refund verified",
            ("bank-statement:line-3",),
        )
        preview = application.preview_correction(
            selection,
            CorrelationId("durable-correction-preview"),
        )
        request = _correction_request(
            selection,
            preview,
            "durable-correction-apply",
            "durable-correction-apply",
        )
        repository = BackgroundJobRepository(connection)
        job_id = "durable-correction-job"
        assert DurableJobCommandApplication(repository, connection).enqueue(
            _correction_apply_job_command(job_id, request)
        ).job_id == job_id
    finally:
        connection.close()

    worker_connection = get_connection()
    try:
        worker = DurableJobWorker(
            BackgroundJobRepository(worker_connection),
            worker_connection,
            default_job_handlers(),
            "durable-correction-worker",
            retry_delay_seconds=0,
        )
        assert worker.recover_and_run_once() is True
        assert worker.recover_and_run_once() is False
        stored = BackgroundJobRepository(worker_connection).get_job(job_id)
        assert stored is not None
        assert stored.status == "succeeded", stored.error_payload
        assert stored.receipt_payload == {
            "kind": "success",
            "result_reference": "finance_import_correction:finance-import-row:1",
            "schema_version": 1,
        }
        with worker_connection.cursor() as cursor:
            cursor.execute("SELECT COUNT(*) AS count FROM client_ledger_entries")
            assert cursor.fetchone() == {"count": 1}
            cursor.execute(
                "SELECT COUNT(*) AS count FROM finance_import_correction_receipts"
            )
            assert cursor.fetchone() == {"count": 1}
    finally:
        worker_connection.close()


def test_real_taishin_refund_return_reopens_original_refund_once(tmp_path):
    bootstrap(_arguments())
    _seed_open_refund_obligation()
    _post_manual_refund(tmp_path)
    _ingest_unresolved_taishin_inflow(tmp_path)
    _deliver_finance_import_outbox()

    from api.dependencies.finance_import import build_finance_import_application
    from domains.finance_import.correction import FinanceImportCorrectionSelection
    from domains.finance_import.planning import FinanceClassificationType
    from infrastructure.mysql.finance_import_owning_domain_composite import MySqlFinanceImportOwningDomainComposite
    from infrastructure.mysql.mysql_adapter import get_connection
    from shared_kernel.identities import ActorContext, CorrelationId, ExpectedVersion, IdempotencyKey
    from subsystems.finance_import.correction_workflow import FinanceImportCorrectionApplyRequest

    connection = get_connection()
    try:
        with connection.cursor() as cursor:
            cursor.execute("SELECT id FROM client_ledger_entries WHERE entry_type='refund'")
            refund_ledger_entry_id = str(cursor.fetchone()["id"])
        application = build_finance_import_application(connection, MySqlFinanceImportOwningDomainComposite(connection))
        selection = FinanceImportCorrectionSelection(
            "finance-import-row:2", FinanceClassificationType.CLIENT_REFUND_RETURN,
            ("refund:C-1",), "returned customer refund was verified",
            ("bank-statement:return-line", "refund-ledger:" + refund_ledger_entry_id),
            refund_ledger_entry_id,
        )
        preview = application.preview_correction(selection, CorrelationId("lu-test-refund-return-preview"))
        request = FinanceImportCorrectionApplyRequest(
            selection, ExpectedVersion(preview.batch_version), ExpectedVersion(preview.canonical_fact_version),
            ExpectedVersion(preview.alert_version), preview.fingerprint, IdempotencyKey("lu-test-refund-return"),
            ActorContext("lu-test-runner"), CorrelationId("lu-test-refund-return-apply"),
        )
        receipt = application.correct_and_post(request)
        assert application.correct_and_post(request) == receipt
        with connection.cursor() as cursor:
            cursor.execute("SELECT entry_type,amount_ntd FROM client_ledger_entries ORDER BY id")
            assert cursor.fetchall() == [
                {"entry_type": "refund", "amount_ntd": 300},
                {"entry_type": "refund_reversal", "amount_ntd": 300},
            ]
            cursor.execute("SELECT amount_due_ntd,status FROM client_obligations WHERE obligation_identity='refund:C-1'")
            assert cursor.fetchone() == {"amount_due_ntd": 300, "status": "open"}
    finally:
        connection.close()


def test_mismatched_refund_return_remains_manual_review_without_partial_writes(tmp_path):
    bootstrap(_arguments())
    _seed_open_refund_obligation()
    _post_manual_refund(tmp_path)
    _ingest_unresolved_taishin_inflow(tmp_path, amount=299)
    _deliver_finance_import_outbox()

    from api.dependencies.finance_import import build_finance_import_application
    from domains.finance_import.correction import FinanceImportCorrectionSelection
    from domains.finance_import.planning import FinanceClassificationType
    from infrastructure.mysql.finance_import_owning_domain_composite import MySqlFinanceImportOwningDomainComposite
    from infrastructure.mysql.mysql_adapter import get_connection
    from shared_kernel.identities import CorrelationId
    from subsystems.finance_import.correction_workflow import FinanceImportCorrectionWorkflowError

    connection = get_connection()
    try:
        with connection.cursor() as cursor:
            cursor.execute("SELECT id FROM client_ledger_entries WHERE entry_type='refund'")
            ledger_id = str(cursor.fetchone()["id"])
        application = build_finance_import_application(connection, MySqlFinanceImportOwningDomainComposite(connection))
        selection = FinanceImportCorrectionSelection(
            "finance-import-row:2", FinanceClassificationType.CLIENT_REFUND_RETURN,
            ("refund:C-1",), "returned amount does not equal original refund",
            ("bank-statement:return-line", "refund-ledger:" + ledger_id), ledger_id,
        )
        with pytest.raises(FinanceImportCorrectionWorkflowError) as error:
            application.preview_correction(selection, CorrelationId("lu-test-refund-return-mismatch"))
        assert error.value.error.code == "allocation_not_exact"
        with connection.cursor() as cursor:
            cursor.execute("SELECT entry_type FROM client_ledger_entries ORDER BY id")
            assert cursor.fetchall() == [{"entry_type": "refund"}]
            cursor.execute("SELECT reconciliation_status FROM finance_import_rows WHERE id=2")
            assert cursor.fetchone() == {"reconciliation_status": "pending"}
            cursor.execute("SELECT predicate_active,workflow_status FROM anomaly_current_alerts WHERE definition_code='finance_import_manual_review'")
            assert not cursor.fetchall()
    finally:
        connection.close()


def test_g12_failed_correction_rolls_back_then_retries_exactly_once(tmp_path):
    bootstrap(_arguments())
    _seed_open_refund_obligation()
    _ingest_unresolved_taishin_outflow(tmp_path)
    _deliver_finance_import_outbox()

    from api.dependencies.finance_import import build_finance_import_application
    from domains.finance_import.correction import FinanceImportCorrectionSelection
    from domains.finance_import.planning import FinanceClassificationType
    from infrastructure.mysql.finance_import_owning_domain_composite import MySqlFinanceImportOwningDomainComposite
    from infrastructure.mysql.mysql_adapter import get_connection
    from shared_kernel.identities import ActorContext, CorrelationId, ExpectedVersion, IdempotencyKey
    from subsystems.finance_import.correction_workflow import FinanceImportCorrectionApplyRequest, FinanceImportCorrectionWorkflowError

    selection = FinanceImportCorrectionSelection(
        "finance-import-row:1", FinanceClassificationType.CLIENT_REFUND,
        ("refund:C-1",), "verified customer refund", ("bank-statement:line-3",),
    )
    connection = get_connection()
    try:
        failed_application = build_finance_import_application(connection, _FailingPostingPort())
        preview = failed_application.preview_correction(selection, CorrelationId("g12-failed-preview"))
        failed_request = _correction_request(selection, preview, "g12-retry-key", "g12-failed-apply")
        with pytest.raises(FinanceImportCorrectionWorkflowError) as error:
            failed_application.correct_and_post(failed_request)
        assert error.value.error.code == "transaction_failed"
        _assert_g12_rollback_state(connection)

        application = build_finance_import_application(connection, MySqlFinanceImportOwningDomainComposite(connection))
        retry_preview = application.preview_correction(selection, CorrelationId("g12-retry-preview"))
        receipt = application.correct_and_post(
            _correction_request(selection, retry_preview, "g12-retry-key", "g12-retry-apply")
        )
        assert application.correct_and_post(
            _correction_request(selection, retry_preview, "g12-retry-key", "g12-retry-apply")
        ) == receipt
        with connection.cursor() as cursor:
            cursor.execute("SELECT COUNT(*) AS count FROM client_ledger_entries")
            assert cursor.fetchone() == {"count": 1}
            cursor.execute("SELECT COUNT(*) AS count FROM finance_import_correction_receipts")
            assert cursor.fetchone() == {"count": 1}
            cursor.execute("SELECT COUNT(*) AS count FROM finance_import_classification_events")
            assert cursor.fetchone() == {"count": 2}
    finally:
        connection.close()


@pytest.mark.parametrize(
    "failure_point",
    (
        "append_manual_classification",
        "post",
        "append_reconciliation_receipt",
        "append_alert_resolved_event",
        "append_outbox",
        "advance_batch_version",
        "save_correction_receipt",
    ),
)
def test_g08_each_finance_correction_persistence_failure_rolls_back(tmp_path, failure_point):
    """Every correction persistence seam shares the same outer MySQL rollback."""
    bootstrap(_arguments())
    _seed_open_refund_obligation()
    _ingest_unresolved_taishin_outflow(tmp_path)
    _deliver_finance_import_outbox()

    from domains.finance_import.correction import FinanceImportCorrectionSelection
    from domains.finance_import.planning import FinanceClassificationType
    from infrastructure.mysql.finance_import_owning_domain_composite import MySqlFinanceImportOwningDomainComposite
    from infrastructure.mysql.finance_import_repository import FinanceImportMySqlUnitOfWork, MySqlFinanceImportRepository
    from infrastructure.mysql.mysql_adapter import get_connection
    from shared_kernel.identities import CorrelationId
    from subsystems.finance_import.correction_workflow import FinanceImportCorrectionWorkflow, FinanceImportCorrectionWorkflowError

    selection = FinanceImportCorrectionSelection(
        "finance-import-row:1", FinanceClassificationType.CLIENT_REFUND,
        ("refund:C-1",), "G08 persistence fault", ("bank-statement:line-3",),
    )
    connection = get_connection()
    try:
        repository = _FailAtCorrectionPersistence(
            MySqlFinanceImportRepository(connection), failure_point,
        )
        composite = MySqlFinanceImportOwningDomainComposite(connection)
        posting_port = _FailAfterPosting(composite) if failure_point == "post" else composite
        workflow = FinanceImportCorrectionWorkflow(
            repository,
            posting_port,
            lambda: FinanceImportMySqlUnitOfWork(connection),
        )
        preview = workflow.preview(selection, CorrelationId(f"g08-{failure_point}-preview"))
        request = _correction_request(
            selection,
            preview,
            f"g08-{failure_point}",
            f"g08-{failure_point}-apply",
        )
        composite.bind_request(request)
        try:
            with pytest.raises(FinanceImportCorrectionWorkflowError) as error:
                workflow.correct_and_post(request)
        finally:
            composite.clear_request()
        assert error.value.error.code == "transaction_failed"
        _assert_g08_no_partial_correction_commit(connection)
    finally:
        connection.close()


def test_g11_ordinary_finance_review_acknowledges_owner_outbox_without_legacy_projection(tmp_path):
    bootstrap(_arguments())
    receipt = _ingest_unresolved_taishin_outflow(tmp_path)
    replay = _ingest_unresolved_taishin_outflow(tmp_path)
    assert replay == receipt
    _deliver_finance_import_outbox()

    from infrastructure.mysql.mysql_adapter import get_connection

    connection = get_connection()
    try:
        with connection.cursor() as cursor:
            cursor.execute(
                "SELECT COUNT(*) AS count FROM anomaly_current_alerts "
                "WHERE definition_code='finance_import_manual_review'"
            )
            assert cursor.fetchone() == {"count": 0}
            cursor.execute("SELECT COUNT(*) AS count FROM finance_import_outbox")
            assert cursor.fetchone() == {"count": 1}
            cursor.execute("SELECT COUNT(*) AS count FROM finance_import_apply_receipts")
            assert cursor.fetchone() == {"count": 0}
            cursor.execute(
                "SELECT COUNT(*) AS count FROM import_warning_occurrences "
                "WHERE owning_lane='finance_import'"
            )
            assert cursor.fetchone() == {"count": 0}
    finally:
        connection.close()


def test_finance_final_dispatch_acknowledges_without_legacy_warning_mutation(tmp_path):
    bootstrap(_arguments())
    _ingest_unresolved_taishin_outflow(tmp_path)
    _deliver_finance_import_outbox()

    from infrastructure.mysql.mysql_adapter import get_connection

    connection = get_connection()
    try:
        with connection.cursor() as cursor:
            cursor.execute(
                "INSERT INTO finance_import_outbox "
                "(batch_id,intent_key,intent_type,payload_snapshot) "
                "VALUES (1,'lu-test-final-dispatch','dispatch_completed',%s)",
                (
                    '{"batch_identity":"finance-import-batch:1",'
                    '"results":[{"row_identity":"finance-import-row:1",'
                    '"outcome":"existing"}]}',
                ),
            )
        connection.commit()
    finally:
        connection.close()

    _deliver_finance_import_outbox()
    _deliver_finance_import_outbox()
    from infrastructure.mysql.mysql_adapter import get_connection

    connection = get_connection()
    try:
        with connection.cursor() as cursor:
            cursor.execute(
                "SELECT status FROM finance_import_outbox "
                "WHERE intent_key='lu-test-final-dispatch'"
            )
            assert cursor.fetchone() == {"status": "delivered"}
            cursor.execute(
                "SELECT COUNT(*) AS count FROM import_warning_occurrences "
                "WHERE owning_lane='finance_import'"
            )
            assert cursor.fetchone() == {"count": 0}
    finally:
        connection.close()


def test_real_taishin_subsidy_payout_advances_then_recovers_after_government_allocation(tmp_path):
    bootstrap(_arguments())
    _seed_open_subsidy_return_with_claim_link()
    receipt = _ingest_unresolved_taishin_outflow(tmp_path, amount=6000)
    _deliver_finance_import_outbox()

    from api.dependencies.finance_import import build_finance_import_application
    from domains.finance_import.correction import FinanceImportCorrectionSelection
    from domains.finance_import.planning import FinanceClassificationType
    from infrastructure.mysql.finance_import_owning_domain_composite import MySqlFinanceImportOwningDomainComposite
    from infrastructure.mysql.mysql_adapter import get_connection
    from shared_kernel.identities import ActorContext, CorrelationId, ExpectedVersion, IdempotencyKey
    from subsystems.finance_import.correction_workflow import FinanceImportCorrectionApplyRequest

    connection = get_connection()
    try:
        application = build_finance_import_application(
            connection,
            MySqlFinanceImportOwningDomainComposite(connection),
        )
        selection = FinanceImportCorrectionSelection(
            "finance-import-row:1",
            FinanceClassificationType.CLIENT_SUBSIDY_RETURN,
            ("subsidy:C-ADV",),
            "verified subsidy claim has not received a government allocation",
            ("bank-statement:line-3", "subsidy-claim-item:C-ADV"),
        )
        preview = application.preview_correction(
            selection,
            CorrelationId("lu-test-subsidy-advance-preview"),
        )
        posted = application.correct_and_post(
            FinanceImportCorrectionApplyRequest(
                selection,
                ExpectedVersion(preview.batch_version),
                ExpectedVersion(preview.canonical_fact_version),
                ExpectedVersion(preview.alert_version),
                preview.fingerprint,
                IdempotencyKey("lu-test-subsidy-advance-correction"),
                ActorContext("lu-test-runner"),
                CorrelationId("lu-test-subsidy-advance-apply"),
            )
        )
        with connection.cursor() as cursor:
            cursor.execute("SELECT entry_type,amount_ntd FROM client_ledger_entries")
            assert cursor.fetchone() == {"entry_type": "subsidy_advance", "amount_ntd": 6000}
            cursor.execute("SELECT amount_due_ntd,status FROM client_obligations WHERE obligation_identity='subsidy:C-ADV'")
            assert cursor.fetchone() == {"amount_due_ntd": 0, "status": "settled"}
    finally:
        connection.close()

    assert posted.ledger_entry_count == posted.allocation_count == 1
    _append_full_government_receipt_allocation()
    _deliver_subsidy_advance_recovery()
    _assert_subsidy_advance_was_recovered_without_second_payout()


def _seed_open_subsidy_return_with_claim_link() -> None:
    from infrastructure.mysql.mysql_adapter import get_connection

    connection = get_connection()
    try:
        with connection.cursor() as cursor:
            cursor.execute("INSERT INTO clients(case_no,name) VALUES ('C-ADV','Advance Client')")
            client_id = int(cursor.lastrowid)
            cursor.execute("INSERT INTO orders(case_no,client_id,status,actual_end_date) VALUES ('C-ADV',%s,'訂單完成','2026-01-31')", (client_id,))
            cursor.execute("INSERT INTO staff(name) VALUES ('Advance Staff')")
            staff_id = int(cursor.lastrowid)
            cursor.execute("INSERT INTO case_staff_assignments(case_no,staff_id,assignment_sequence,status) VALUES ('C-ADV',%s,1,'completed')", (staff_id,))
            assignment_id = int(cursor.lastrowid)
            cursor.execute("INSERT INTO client_obligation_events(obligation_identity,case_no,obligation_type,direction,event_type,before_amount_ntd,after_amount_ntd,before_due_date,after_due_date,source_event_identity,source_obligation_identity,expected_account_version,idempotency_key,actor,reason) VALUES ('subsidy:C-ADV','C-ADV','subsidy_return','payable_to_client','established',0,6000,NULL,'2026-03-15','lu-test-subsidy-root',NULL,0,'lu-test-subsidy-root','lu-test-runner','fixture root fact')")
            event_id = int(cursor.lastrowid)
            cursor.execute("INSERT INTO client_obligations(obligation_identity,case_no,obligation_type,direction,source_obligation_identity,amount_due_ntd,due_date,status,current_event_id,projection_version) VALUES ('subsidy:C-ADV','C-ADV','subsidy_return','payable_to_client',NULL,6000,'2026-03-15','open',%s,0)", (event_id,))
            cursor.execute("INSERT INTO subsidy_claim_batches(application_year,quarter,revision,status,requested_amount,approved_amount,paid_amount,submitted_at,approved_at) VALUES (2026,1,1,'approved',6000,6000,0,CURRENT_TIMESTAMP,CURRENT_TIMESTAMP)")
            batch_id = int(cursor.lastrowid)
            cursor.execute("INSERT INTO subsidy_claim_batch_items(batch_id,case_no,assignment_id,staff_id,requested_amount,approved_amount) VALUES (%s,'C-ADV',%s,%s,6000,6000)", (batch_id, assignment_id, staff_id))
            claim_item_id = int(cursor.lastrowid)
            cursor.execute("INSERT INTO client_subsidy_return_claim_item_links(obligation_identity,claim_item_id,entitled_amount_ntd) VALUES ('subsidy:C-ADV',%s,6000)", (claim_item_id,))
            cursor.execute("INSERT INTO government_subsidy_batch_accounts(batch_id,aggregate_version,requested_total_ntd,approved_total_ntd,net_allocated_ntd,outstanding_ntd,status) VALUES (%s,1,6000,6000,0,6000,'approved')", (batch_id,))
            cursor.execute(
                "INSERT INTO client_refund_recipient_snapshots "
                "(refund_obligation_identity,case_no,bank_code,bank_account,source_kind) "
                "VALUES ('subsidy:C-ADV','C-ADV','synthetic-bank',%s,'test-fixture')",
                ("9" * 16,),
            )
        connection.commit()
    finally:
        connection.close()


def _append_full_government_receipt_allocation() -> None:
    from infrastructure.mysql.mysql_adapter import get_connection

    connection = get_connection()
    try:
        with connection.cursor() as cursor:
            cursor.execute("SELECT id FROM subsidy_claim_batches")
            batch_id = int(cursor.fetchone()["id"])
            cursor.execute("SELECT id FROM subsidy_claim_batch_items")
            claim_item_id = int(cursor.fetchone()["id"])
            incoming_row_id = _insert_finance_row(cursor, "b", "incoming", "2026-04-01", credit=6000)
            cursor.execute("INSERT INTO government_subsidy_transactions(claim_batch_id,finance_import_row_id,transaction_type,transaction_status,amount,occurred_at,external_reference) VALUES (%s,%s,'receipt','succeeded',6000,'2026-04-01','lu-test-government-receipt')", (batch_id, incoming_row_id))
            transaction_id = int(cursor.lastrowid)
            cursor.execute("INSERT INTO government_subsidy_allocations(transaction_id,claim_batch_id,claim_item_id,allocation_type,allocated_amount) VALUES (%s,%s,%s,'receipt',6000)", (transaction_id, batch_id, claim_item_id))
            cursor.execute("UPDATE government_subsidy_batch_accounts SET aggregate_version=2,net_allocated_ntd=6000,outstanding_ntd=0,status='paid' WHERE batch_id=%s", (batch_id,))
            cursor.execute("INSERT INTO government_subsidy_projection_events(batch_id,transaction_id,before_status,after_status,before_net_allocated_ntd,after_net_allocated_ntd,outstanding_ntd,expected_batch_version,resulting_batch_version,preview_fingerprint,actor,reason,idempotency_key) VALUES (%s,%s,'approved','paid',0,6000,0,1,2,%s,'lu-test-runner','fixture receipt','lu-test-projection')", (batch_id, transaction_id, "c" * 64))
            projection_event_id = int(cursor.lastrowid)
            cursor.execute("INSERT INTO government_subsidy_outbox(batch_id,transaction_id,projection_event_id,intent_key,intent_type,payload_snapshot) VALUES (%s,%s,%s,'lu-test-subsidy-allocation','government_subsidy_receipt_allocated',%s)", (batch_id, transaction_id, projection_event_id, '{"transaction_id":' + str(transaction_id) + ',"allocations":[{"claim_item_id":' + str(claim_item_id) + ',"case_no":"C-ADV","amount_ntd":6000}]}'))
        connection.commit()
    finally:
        connection.close()


def _deliver_subsidy_advance_recovery() -> None:
    from infrastructure.mysql.mysql_adapter import get_connection
    from infrastructure.mysql.subsidy_advance_recovery_repository import MySqlSubsidyAdvanceRecoveryRepository
    from subsystems.government_subsidy.subsidy_advance_outbox_consumer import (
        consume_government_subsidy_advance_events,
    )

    connection = get_connection()
    try:
        assert consume_government_subsidy_advance_events(
            connection, MySqlSubsidyAdvanceRecoveryRepository
        ) == (1, 0)
    finally:
        connection.close()


def _assert_subsidy_advance_was_recovered_without_second_payout() -> None:
    from infrastructure.mysql.mysql_adapter import get_connection
    from infrastructure.mysql.subsidy_advance_recovery_repository import MySqlSubsidyAdvanceRecoveryRepository

    connection = get_connection()
    try:
        with connection.cursor() as cursor:
            cursor.execute("SELECT entry_type,amount_ntd FROM client_ledger_entries ORDER BY id")
            assert cursor.fetchall() == [{"entry_type": "subsidy_advance", "amount_ntd": 6000}]
            cursor.execute("SELECT recovered_amount_ntd FROM client_subsidy_advance_recoveries")
            assert cursor.fetchone() == {"recovered_amount_ntd": 6000}
    finally:
        connection.close()


def test_government_receipt_recovers_existing_subsidy_advance_without_second_client_payout():
    bootstrap(_arguments())
    _seed_subsidy_advance_recovery_facts()

    from infrastructure.mysql.mysql_adapter import get_connection
    from infrastructure.mysql.subsidy_advance_recovery_repository import (
        MySqlSubsidyAdvanceRecoveryRepository,
    )
    from subsystems.government_subsidy.subsidy_advance_outbox_consumer import (
        consume_government_subsidy_advance_events,
    )

    connection = get_connection()
    try:
        assert consume_government_subsidy_advance_events(
            connection, MySqlSubsidyAdvanceRecoveryRepository
        ) == (1, 0)
        with connection.cursor() as cursor:
            cursor.execute("SELECT entry_type,amount_ntd FROM client_ledger_entries ORDER BY id")
            assert cursor.fetchall() == [{"entry_type": "subsidy_advance", "amount_ntd": 6000}]
            cursor.execute("SELECT recovered_amount_ntd FROM client_subsidy_advance_recoveries")
            assert cursor.fetchone() == {"recovered_amount_ntd": 6000}
            cursor.execute("SELECT status FROM government_subsidy_outbox")
            assert cursor.fetchone() == {"status": "delivered"}
    finally:
        connection.close()


def test_partial_government_subsidy_allocation_creates_review_without_auto_netting():
    bootstrap(_arguments())
    _seed_subsidy_advance_recovery_facts(government_allocation_amount=5000)

    from infrastructure.mysql.mysql_adapter import get_connection
    from infrastructure.mysql.subsidy_advance_recovery_repository import MySqlSubsidyAdvanceRecoveryRepository
    from subsystems.government_subsidy.subsidy_advance_outbox_consumer import (
        consume_government_subsidy_advance_events,
    )

    connection = get_connection()
    try:
        assert consume_government_subsidy_advance_events(
            connection, MySqlSubsidyAdvanceRecoveryRepository
        ) == (1, 0)
        with connection.cursor() as cursor:
            cursor.execute("SELECT COUNT(*) AS count FROM client_subsidy_advance_recoveries")
            assert cursor.fetchone() == {"count": 0}
            cursor.execute("SELECT intent_type,payload_snapshot FROM client_finance_outbox")
            anomaly = cursor.fetchone()
            assert anomaly["intent_type"] == "anomaly_review_required"
            assert "subsidy_advance_settlement_ambiguous" in anomaly["payload_snapshot"]
            cursor.execute("SELECT entry_type,amount_ntd FROM client_ledger_entries ORDER BY id")
            assert cursor.fetchall() == [{"entry_type": "subsidy_advance", "amount_ntd": 6000}]
    finally:
        connection.close()


def _seed_subsidy_advance_recovery_facts(*, government_allocation_amount=6000) -> None:
    from infrastructure.mysql.mysql_adapter import get_connection

    connection = get_connection()
    try:
        with connection.cursor() as cursor:
            cursor.execute("INSERT INTO clients(case_no,name) VALUES ('C-ADV','Advance Client')")
            client_id = int(cursor.lastrowid)
            cursor.execute("INSERT INTO orders(case_no,client_id,status,actual_end_date) VALUES ('C-ADV',%s,'訂單完成','2026-01-31')", (client_id,))
            cursor.execute("INSERT INTO staff(name) VALUES ('Advance Staff')")
            staff_id = int(cursor.lastrowid)
            cursor.execute("INSERT INTO case_staff_assignments(case_no,staff_id,assignment_sequence,status) VALUES ('C-ADV',%s,1,'completed')", (staff_id,))
            assignment_id = int(cursor.lastrowid)
            cursor.execute("INSERT INTO client_obligation_events(obligation_identity,case_no,obligation_type,direction,event_type,before_amount_ntd,after_amount_ntd,before_due_date,after_due_date,source_event_identity,source_obligation_identity,expected_account_version,idempotency_key,actor,reason) VALUES ('subsidy:C-ADV','C-ADV','subsidy_return','payable_to_client','established',0,6000,NULL,'2026-03-15','lu-test-subsidy-root',NULL,0,'lu-test-subsidy-root','lu-test-runner','fixture root fact')")
            event_id = int(cursor.lastrowid)
            cursor.execute("INSERT INTO client_obligations(obligation_identity,case_no,obligation_type,direction,source_obligation_identity,amount_due_ntd,due_date,status,current_event_id,projection_version) VALUES ('subsidy:C-ADV','C-ADV','subsidy_return','payable_to_client',NULL,0,'2026-03-15','settled',%s,1)", (event_id,))
            outgoing_row_id = _insert_finance_row(cursor, "a", "outgoing", "2026-03-15", debit=6000)
            cursor.execute("INSERT INTO client_ledger_entries(case_no,finance_import_row_id,entry_type,amount_ntd,occurred_on,reconciliation_reference,idempotency_key,actor,reason) VALUES ('C-ADV',%s,'subsidy_advance',6000,'2026-03-15','lu-test-subsidy-advance','lu-test-subsidy-advance','lu-test-runner','fixture advance')", (outgoing_row_id,))
            advance_entry_id = int(cursor.lastrowid)
            cursor.execute("INSERT INTO client_ledger_obligation_allocations(ledger_entry_id,obligation_identity,amount_ntd,allocation_ordinal) VALUES (%s,'subsidy:C-ADV',6000,1)", (advance_entry_id,))
            cursor.execute("INSERT INTO subsidy_claim_batches(application_year,quarter,revision,status,requested_amount,approved_amount,paid_amount,submitted_at,approved_at) VALUES (2026,1,1,'approved',6000,6000,0,CURRENT_TIMESTAMP,CURRENT_TIMESTAMP)")
            batch_id = int(cursor.lastrowid)
            cursor.execute("INSERT INTO subsidy_claim_batch_items(batch_id,case_no,assignment_id,staff_id,requested_amount,approved_amount) VALUES (%s,'C-ADV',%s,%s,6000,6000)", (batch_id, assignment_id, staff_id))
            claim_item_id = int(cursor.lastrowid)
            cursor.execute("INSERT INTO client_subsidy_return_claim_item_links(obligation_identity,claim_item_id,entitled_amount_ntd) VALUES ('subsidy:C-ADV',%s,6000)", (claim_item_id,))
            incoming_row_id = _insert_finance_row(cursor, "b", "incoming", "2026-04-01", credit=6000)
            cursor.execute("INSERT INTO government_subsidy_transactions(claim_batch_id,finance_import_row_id,transaction_type,transaction_status,amount,occurred_at,external_reference) VALUES (%s,%s,'receipt','succeeded',%s,'2026-04-01','lu-test-government-receipt')", (batch_id, incoming_row_id, government_allocation_amount))
            transaction_id = int(cursor.lastrowid)
            cursor.execute("INSERT INTO government_subsidy_allocations(transaction_id,claim_batch_id,claim_item_id,allocation_type,allocated_amount) VALUES (%s,%s,%s,'receipt',%s)", (transaction_id, batch_id, claim_item_id, government_allocation_amount))
            cursor.execute("INSERT INTO government_subsidy_batch_accounts(batch_id,aggregate_version,requested_total_ntd,approved_total_ntd,net_allocated_ntd,outstanding_ntd,status) VALUES (%s,1,6000,6000,%s,%s,%s)", (batch_id, government_allocation_amount, 6000 - government_allocation_amount, "paid" if government_allocation_amount == 6000 else "partially_paid"))
            cursor.execute("INSERT INTO government_subsidy_projection_events(batch_id,transaction_id,before_status,after_status,before_net_allocated_ntd,after_net_allocated_ntd,outstanding_ntd,expected_batch_version,resulting_batch_version,preview_fingerprint,actor,reason,idempotency_key) VALUES (%s,%s,'approved',%s,0,%s,%s,0,1,%s,'lu-test-runner','fixture receipt','lu-test-projection')", (batch_id, transaction_id, "paid" if government_allocation_amount == 6000 else "partially_paid", government_allocation_amount, 6000 - government_allocation_amount, "c" * 64))
            projection_event_id = int(cursor.lastrowid)
            cursor.execute("INSERT INTO government_subsidy_outbox(batch_id,transaction_id,projection_event_id,intent_key,intent_type,payload_snapshot) VALUES (%s,%s,%s,'lu-test-subsidy-allocation','government_subsidy_receipt_allocated',%s)", (batch_id, transaction_id, projection_event_id, '{"transaction_id":' + str(transaction_id) + ',"allocations":[{"claim_item_id":' + str(claim_item_id) + ',"case_no":"C-ADV","amount_ntd":' + str(government_allocation_amount) + '}]}'))
        connection.commit()
    finally:
        connection.close()


def _insert_finance_row(cursor, suffix: str, direction: str, transaction_date: str, *, debit=0, credit=0) -> int:
    cursor.execute("INSERT INTO finance_import_rows(dedup_fingerprint,format_id,transaction_date,debit,credit,direction,bank_references,warnings,raw_payload) VALUES (%s,'taishin',%s,%s,%s,%s,JSON_OBJECT(),JSON_ARRAY(),JSON_OBJECT())", (suffix * 64, transaction_date, debit or None, credit or None, direction))
    return int(cursor.lastrowid)


def _seed_open_refund_obligation() -> None:
    from infrastructure.mysql.mysql_adapter import get_connection

    connection = get_connection()
    try:
        with connection.cursor() as cursor:
            synthetic_account = "9" * 16
            cursor.execute("INSERT INTO clients(case_no,name) VALUES ('C-1','E2E Client')")
            client_id = int(cursor.lastrowid)
            cursor.execute(
                "INSERT INTO beclass_records(query_no,refund_bank_code,refund_account_no) "
                "VALUES ('C-1','synthetic-bank',%s)",
                (synthetic_account,),
            )
            cursor.execute("INSERT INTO orders(case_no,client_id,status) VALUES ('C-1',%s,'訂單完成')", (client_id,))
            cursor.execute("INSERT INTO client_obligation_events(obligation_identity,case_no,obligation_type,direction,event_type,before_amount_ntd,after_amount_ntd,before_due_date,after_due_date,source_event_identity,source_obligation_identity,expected_account_version,idempotency_key,actor,reason) VALUES ('refund:C-1','C-1','refund','payable_to_client','established',0,300,NULL,'2026-08-15','lu-test-refund-root',NULL,0,'lu-test-refund-root','lu-test-runner','fixture root fact')")
            event_id = int(cursor.lastrowid)
            cursor.execute("INSERT INTO client_obligations(obligation_identity,case_no,obligation_type,direction,source_obligation_identity,amount_due_ntd,due_date,status,current_event_id,projection_version) VALUES ('refund:C-1','C-1','refund','payable_to_client',NULL,300,'2026-08-15','open',%s,0)", (event_id,))
            cursor.execute(
                "INSERT INTO client_refund_recipient_snapshots "
                "(refund_obligation_identity,case_no,bank_code,bank_account,source_kind) "
                "VALUES ('refund:C-1','C-1','synthetic-bank',%s,'test-fixture')",
                (synthetic_account,),
            )
        connection.commit()
    finally:
        connection.close()


def _ingest_unresolved_taishin_outflow(tmp_path, *, amount=300):
    workbook = tmp_path / "taishin-manual-refund.xlsx"
    if not workbook.exists():
        pd.DataFrame(
            [["說明"], ["序號", "交易日期", "交易時間", "帳務日期", "摘要", "支出金額", "存入金額", "帳戶餘額", "備註"], ["0001", "2026/08/04", "09:08:07", "2026/08/04", "轉帳", str(amount), "", "9000", "客戶退款 " + ("9" * 16)]]
        ).to_excel(workbook, sheet_name="交易明細", index=False, header=False)
    from infrastructure.mysql.mysql_adapter import get_connection
    from shared_kernel.identities import ActorContext, IdempotencyKey
    from subsystems.finance_import.ingestion import ingest_finance_workbook
    return ingest_finance_workbook(str(workbook), IdempotencyKey("lu-test-manual-refund-ingest"), ActorContext("lu-test-runner"), connection_factory=get_connection, normalizer=normalize_workbook)


def _post_manual_refund(tmp_path):
    receipt = _ingest_unresolved_taishin_outflow(tmp_path)
    _deliver_finance_import_outbox()
    from api.dependencies.finance_import import build_finance_import_application
    from domains.finance_import.correction import FinanceImportCorrectionSelection
    from domains.finance_import.planning import FinanceClassificationType
    from infrastructure.mysql.finance_import_owning_domain_composite import MySqlFinanceImportOwningDomainComposite
    from infrastructure.mysql.mysql_adapter import get_connection
    from shared_kernel.identities import ActorContext, CorrelationId, ExpectedVersion, IdempotencyKey
    from subsystems.finance_import.correction_workflow import FinanceImportCorrectionApplyRequest
    connection = get_connection()
    try:
        application = build_finance_import_application(connection, MySqlFinanceImportOwningDomainComposite(connection))
        selection = FinanceImportCorrectionSelection("finance-import-row:1", FinanceClassificationType.CLIENT_REFUND, ("refund:C-1",), "refund verified", ("bank-statement:line-3",))
        preview = application.preview_correction(selection, CorrelationId("lu-test-refund-preview"))
        application.correct_and_post(FinanceImportCorrectionApplyRequest(selection, ExpectedVersion(preview.batch_version), ExpectedVersion(preview.canonical_fact_version), ExpectedVersion(preview.alert_version), preview.fingerprint, IdempotencyKey("lu-test-refund-post"), ActorContext("lu-test-runner"), CorrelationId("lu-test-refund-apply")))
    finally:
        connection.close()
    return receipt


def _ingest_unresolved_taishin_inflow(tmp_path, *, amount=300):
    workbook = tmp_path / "taishin-refund-return.xlsx"
    pd.DataFrame([["說明"], ["序號", "交易日期", "交易時間", "帳務日期", "摘要", "支出金額", "存入金額", "帳戶餘額", "備註"], ["0002", "2026/08/05", "09:08:07", "2026/08/05", "退款退回", "", str(amount), str(9000 + amount), "客戶退款退回"]]).to_excel(workbook, sheet_name="交易明細", index=False, header=False)
    from infrastructure.mysql.mysql_adapter import get_connection
    from shared_kernel.identities import ActorContext, IdempotencyKey
    from subsystems.finance_import.ingestion import ingest_finance_workbook
    return ingest_finance_workbook(str(workbook), IdempotencyKey("lu-test-refund-return-ingest"), ActorContext("lu-test-runner"), connection_factory=get_connection, normalizer=normalize_workbook)


def _correction_request(selection, preview, idempotency_key, correlation_id):
    from shared_kernel.identities import ActorContext, CorrelationId, ExpectedVersion, IdempotencyKey
    from subsystems.finance_import.correction_workflow import FinanceImportCorrectionApplyRequest
    return FinanceImportCorrectionApplyRequest(
        selection, ExpectedVersion(preview.batch_version), ExpectedVersion(preview.canonical_fact_version),
        ExpectedVersion(preview.alert_version), preview.fingerprint, IdempotencyKey(idempotency_key),
        ActorContext("admin_user_id:1"), CorrelationId(correlation_id),
    )


def _assert_g12_rollback_state(connection):
    with connection.cursor() as cursor:
        cursor.execute("SELECT COUNT(*) AS count FROM client_ledger_entries")
        assert cursor.fetchone() == {"count": 0}
        cursor.execute("SELECT COUNT(*) AS count FROM client_ledger_obligation_allocations")
        assert cursor.fetchone() == {"count": 0}
        cursor.execute("SELECT COUNT(*) AS count FROM finance_import_correction_receipts")
        assert cursor.fetchone() == {"count": 0}
        cursor.execute("SELECT COUNT(*) AS count FROM finance_import_classification_events")
        assert cursor.fetchone() == {"count": 1}
        cursor.execute("SELECT predicate_active,workflow_status FROM anomaly_current_alerts WHERE definition_code='finance_import_manual_review'")
        # Ordinary classification belongs to Finance; it must not create an Anomalies alert.
        assert cursor.fetchone() is None


def _assert_g08_no_partial_correction_commit(connection):
    _assert_g12_rollback_state(connection)
    with connection.cursor() as cursor:
        cursor.execute(
            "SELECT COUNT(*) AS count FROM finance_import_outbox "
            "WHERE intent_type='manual_correction_completed'"
        )
        assert cursor.fetchone() == {"count": 0}
        cursor.execute(
            "SELECT batch_version FROM finance_import_batch_contracts"
        )
        assert cursor.fetchone() == {"batch_version": 0}


class _FailAtCorrectionPersistence:
    def __init__(self, repository, failure_point):
        self._repository = repository
        self._failure_point = failure_point

    def __getattr__(self, name):
        if name == self._failure_point:
            return self._raise_persistence_failure
        return getattr(self._repository, name)

    def _raise_persistence_failure(self, *_args, **_kwargs):
        raise RuntimeError(f"injected failure at {self._failure_point}")


class _FailAfterPosting:
    def __init__(self, posting_port):
        self._posting_port = posting_port

    def post(self, candidate):
        self._posting_port.post(candidate)
        raise RuntimeError("injected failure after owning Domain posting")


class _FailingPostingPort:
    def bind_request(self, request):
        del request

    def clear_request(self):
        return None

    def resolve(self, candidate):
        return candidate

    def post(self, candidate):
        del candidate
        raise RuntimeError("injected owning-domain failure")


def _deliver_finance_import_outbox() -> None:
    from infrastructure.mysql.anomaly_runtime import build_anomaly_runtime
    from infrastructure.mysql.mysql_adapter import get_connection
    from subsystems.finance_import.finance_import_anomaly_consumer import consume_finance_import_anomaly_events
    connection = get_connection()
    try:
        result = consume_finance_import_anomaly_events(
            connection, runtime=build_anomaly_runtime()
        )
        assert result.failed_count == 0
    finally:
        connection.close()


def _assert_manual_review_stays_out_of_anomaly_center() -> None:
    from infrastructure.mysql.mysql_adapter import get_connection
    connection = get_connection()
    try:
        with connection.cursor() as cursor:
            cursor.execute("SELECT predicate_active,workflow_status FROM anomaly_current_alerts WHERE definition_code='finance_import_manual_review'")
            # Ordinary classification belongs to Finance; it must not create an Anomalies alert.
            assert cursor.fetchone() is None
    finally:
        connection.close()


def _assert_manual_review_warning_remains_open() -> None:
    from infrastructure.mysql.mysql_adapter import get_connection

    connection = get_connection()
    try:
        with connection.cursor() as cursor:
            cursor.execute(
                "SELECT task.tracking_status,task.tracking_version "
                "FROM import_warning_current_tasks task "
                "JOIN import_warning_occurrences occurrence "
                "ON occurrence.id=task.occurrence_id "
                "WHERE occurrence.owning_lane='finance_import' "
                "AND occurrence.logical_code='FINANCE-ROW-001'"
            )
            assert cursor.fetchone() == {
                "tracking_status": "open",
                "tracking_version": 1,
            }
            cursor.execute(
                "SELECT event.action,event.actor_kind,event.reason_code "
                "FROM import_warning_tracking_events event "
                "JOIN import_warning_occurrences occurrence "
                "ON occurrence.id=event.occurrence_id "
                "WHERE occurrence.owning_lane='finance_import' "
                "ORDER BY event.resulting_version"
            )
            assert cursor.fetchall() == [
                {
                    "action": "opened",
                    "actor_kind": "system",
                    "reason_code": "source_review_opened",
                }
            ]
            cursor.execute(
                "SELECT COUNT(*) AS count FROM import_warning_tracking_receipts "
                "WHERE occurrence_id=(SELECT id FROM import_warning_occurrences "
                "WHERE owning_lane='finance_import' AND logical_code='FINANCE-ROW-001')"
            )
            assert cursor.fetchone() == {"count": 0}
            cursor.execute(
                "SELECT COUNT(*) AS count FROM import_warning_tracking_outbox outbox "
                "JOIN import_warning_tracking_events event "
                "ON event.id=outbox.tracking_event_id "
                "WHERE event.action='auto_resolved'"
            )
            assert cursor.fetchone() == {"count": 0}
    finally:
        connection.close()
