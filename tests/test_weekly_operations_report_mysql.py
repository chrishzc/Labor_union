"""真實 MySQL、Session、Reporting dependencies 與 XLSX 的週報驗證。

只建立明確指定且尚不存在的 localhost lu_test_* database；不覆寫既有資料庫。
此處直接 seed 合成 root facts 以覆蓋查詢狀態，並非訂單轉態流程驗收。
"""
from argparse import Namespace
from contextlib import closing
from datetime import date, datetime
from io import BytesIO
import os
import secrets

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from api.dependencies.admin_auth import admin_auth_is_enabled
from api.exception_handlers import CorrelationBoundaryMiddleware, install_typed_error_handlers
from api.routes import operations_reports
from domains.orders.lifecycle import OrderLifecycleStatus
from infrastructure.mysql.mysql_adapter import DB_CONFIG, get_connection
from scripts.bootstrap_disposable_mysql_schema import bootstrap
from subsystems.access.authentication_session import (
    authenticate_local_developer_root, bootstrap_root_admin, revoke_admin_session,
)
from subsystems.reporting.weekly_report_metrics_service import WeeklyReportMetricsService, week_starts_between


DATABASE = os.getenv("LABOR_UNION_TEST_MYSQL_DATABASE", "")
pytestmark = [
    pytest.mark.integration,
    pytest.mark.skipif(not DATABASE, reason="requires explicit disposable MySQL database"),
]
BASE = "/api/v1/operations-reports/weekly"
PARAMS = {"start_date": "2026-08-31", "end_date": "2026-09-13"}
V4 = {**PARAMS, "schema_version": "operations-report.v4"}


@pytest.fixture(scope="module")
def api():
    # The target is shared by real dependency factories, not patched connections.
    if (os.getenv("APP_ENV") != "test"
            or DB_CONFIG["host"] not in {"localhost", "127.0.0.1"}
            or DB_CONFIG["database"] != DATABASE):
        raise RuntimeError("weekly report integration requires the explicit local test target")
    bootstrap(Namespace(
        host=DB_CONFIG["host"], port=DB_CONFIG["port"], user=DB_CONFIG["user"],
        password=DB_CONFIG["password"], database=DATABASE, confirm_database=DATABASE,
    ))
    assert admin_auth_is_enabled()
    password = secrets.token_urlsafe(32)
    bootstrap_root_admin(
        connection_factory=get_connection, username="weekly_report_test",
        password=password, display_name="週報測試管理員",
    )
    session = authenticate_local_developer_root(
        "weekly_report_test", password, connection_factory=get_connection,
    )
    if session is None:
        raise AssertionError("test root session could not be created")
    token = session[0]

    statuses = list(OrderLifecycleStatus)
    rows = [(1, datetime(2026, 1, 6, 9), OrderLifecycleStatus.ESTABLISHED)]
    rows.extend((index + 2, datetime(2026, 8, 31, 9) if index < 5
                 else datetime(2026, 9, 7, 9), status)
                for index, status in enumerate(statuses))
    rows.append((12, datetime(2026, 9, 7, 10), None))
    with closing(get_connection()) as connection:
        with connection.cursor() as cursor:
            for serial, created_at, status in rows:
                case_no = f"115000{serial:03d}"
                cursor.execute(
                    """INSERT INTO clients
                       (seq_num, case_no, created_at, name, identity_status, city, address, reject_reason)
                       VALUES (%s,%s,%s,%s,%s,%s,%s,%s)""",
                    (serial, case_no, created_at, f"週報合成案例{serial}", "一般市民",
                     "新竹市", "新竹市東區測試地址", "測試審核不符" if status == OrderLifecycleStatus.CANCELLED else None),
                )
                client_id = cursor.lastrowid
                if status is not None:
                    # Missing service dates are supported partial reporting facts.
                    cursor.execute(
                        """INSERT INTO orders (case_no,client_id,status,service_days,service_hours_per_day)
                           VALUES (%s,%s,%s,20,8)""",
                        (case_no, client_id, status.value),
                    )
        connection.commit()
        metrics = WeeklyReportMetricsService(connection)
        counts = {date(2026, 1, 5): (7, 9), date(2026, 8, 31): (12, 34), date(2026, 9, 7): (5, 0)}
        for monday in week_starts_between(date(2026, 1, 5), date(2026, 9, 13)):
            metrics.save_metric(monday, *counts.get(monday, (0, 0)))

    app = FastAPI()
    app.include_router(operations_reports.router)
    app.add_middleware(CorrelationBoundaryMiddleware)
    install_typed_error_handlers(app)
    assert not app.dependency_overrides
    with TestClient(app) as client:
        client.headers["Authorization"] = f"Bearer {token}"
        yield client, lambda: authenticate_local_developer_root(
            "weekly_report_test", password, connection_factory=get_connection,
        )
    revoke_admin_session(token, connection_factory=get_connection)


def _data(client, params=V4):
    response = client.get(BASE, params=params)
    assert response.status_code == 200, response.text
    return response.json()["data"]


def test_real_authentication_required_and_revoked_session_rejected(api):
    client, create_session = api
    with TestClient(client.app) as anonymous:
        assert anonymous.get(BASE, params=V4).status_code == 401
        assert anonymous.get(BASE, params=V4, headers={"Authorization": "Bearer invalid-test-session"}).status_code == 401
        session = create_session()
        if session is None:
            raise AssertionError("test root session could not be created")
        token = session[0]
        assert revoke_admin_session(token, connection_factory=get_connection)
        assert anonymous.get(BASE, params=V4, headers={"Authorization": f"Bearer {token}"}).status_code == 401
    assert _data(client)["schema_version"] == "operations-report.v4"


def test_real_mysql_counts_and_both_wire_versions(api):
    client, _ = api
    current = _data(client)
    legacy = _data(client, PARAMS)
    assert legacy["schema_version"] == "operations-report.v3"
    assert "annual_totals" not in legacy and "monthly_subtotals" not in legacy
    assert legacy["case_rows"] == current["case_rows"]
    assert current["summary"]["application_count"] == len(current["case_rows"]) == 11
    annual, = current["annual_totals"]
    august, september = current["monthly_subtotals"]
    assert (annual["application_count"], august["application_count"], september["application_count"]) == (12, 5, 6)
    assert [(t["month"], t["promotion_count"], t["inquiry_count"]) for t in (august, september)] == [(8, 12, 34), (9, 5, 0)]
    assert (annual["promotion_count"], annual["inquiry_count"]) == (24, 43)
    for status in OrderLifecycleStatus:
        assert annual["order_status_counts"][status.value] == (2 if status == OrderLifecycleStatus.ESTABLISHED else 1)
    for total in (annual, august, september):
        assert sum(total["order_status_counts"].values()) == total["application_count"]
    assert september["order_status_counts"]["無訂單／狀態缺值"] == 1
    assert september["order_status_counts"]["訂單取消"] == september["review_rejected_count"] == 1


def test_real_metric_writer_reload_preserves_null_and_zero(api):
    client, _ = api
    try:
        response = client.put(f"{BASE}/metrics/2026-08-31", json={"promotion_count": None, "inquiry_count": 0})
        assert response.status_code == 200, response.text
        data = _data(client)
        assert data["monthly_subtotals"][0]["promotion_count"] is None
        assert data["monthly_subtotals"][0]["inquiry_count"] == 0
        assert data["annual_totals"][0]["promotion_count"] is None
        assert data["annual_totals"][0]["inquiry_count"] == 9
    finally:
        response = client.put(f"{BASE}/metrics/2026-08-31", json={"promotion_count": 12, "inquiry_count": 34})
        assert response.status_code == 200, response.text


def test_real_export_matches_api_totals_without_mutating_business_rows(api):
    client, _ = api
    def business_rows():
        with closing(get_connection()) as connection:
            with connection.cursor() as cursor:
                result = []
                for table, key in (("clients", "id"), ("orders", "case_no"), ("weekly_report_metrics", "week_start_date")):
                    cursor.execute(f"SELECT * FROM {table} ORDER BY {key}")
                    result.append(cursor.fetchall())
                return result
    before = business_rows()
    data = _data(client)
    response = client.get(f"{BASE}/export", params=PARAMS)
    assert response.status_code == 200, response.text[:200]
    assert response.headers["X-Operations-Report-Version"] == "operations-report.v4"
    workbook = load_workbook(BytesIO(response.content), data_only=True)
    assert len(workbook.sheetnames) == 3
    sheet = workbook["週報案件受理總表"]
    annual, = data["annual_totals"]
    assert sheet.cell(4, 1).value == "115年度累計"
    assert sheet.cell(4, 8).value == annual["application_count"]
    for column, status in enumerate((*[s.value for s in OrderLifecycleStatus], "無訂單／狀態缺值"), 13):
        assert sheet.cell(3, column).value == status
        assert sheet.cell(4, column).value == annual["order_status_counts"][status]
    labels = [(row, sheet.cell(row, 1).value) for row in range(5, sheet.max_row + 1)]
    august_row = next(row for row, label in labels if label == "115年8月小計")
    september_row = next(row for row, label in labels if label == "115年9月小計")
    assert august_row == 10 and september_row == 17
    assert sheet.cell(august_row, 8).value == 5
    assert sheet.cell(september_row, 8).value == 6
    assert business_rows() == before


def test_real_empty_week_has_unknown_metrics_not_fabricated_zero(api):
    client, _ = api
    data = _data(client, {**V4, "end_date": "2026-09-20"})
    assert data["weekly_metrics"][-1]["promotion_count"] is None
    assert data["weekly_metrics"][-1]["inquiry_count"] is None
    assert data["monthly_subtotals"][-1]["promotion_count"] is None
    assert data["annual_totals"][0]["inquiry_count"] is None
